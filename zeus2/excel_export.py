from __future__ import annotations

import os
import shutil
import uuid
from copy import deepcopy
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Color, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .aging import aging_for_ticket, report_sort_key
from .excel_import import (
    ManagedWorkbook,
    WorkbookValidationError,
    read_closed,
    read_pendings,
    validate_closed_against_index,
)
from .reconcile import import_pendings
from .store import ZeusStore
from .tickets import (
    CLOSED_SCHEMA_DRIFT_COLUMNS,
    LOCAL_COLUMNS,
    PENDING_COLUMNS,
    UPSTREAM_COLUMNS,
    ticket_cell_value,
    workflow_code,
)
from .utils import (
    atomic_write_json,
    atomic_write_text,
    iso_now,
    parse_date,
    parse_datetime,
    sha256_file,
)


class WorkbookPublicationError(RuntimeError):
    pass


NAVY = "17365D"
BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
WHITE = "FFFFFF"
TEXT = "1F2937"
RED = "F4CCCC"
YELLOW = "FFF2CC"
AMBER = "FCE5CD"
GREY = "E7E6E6"
GREEN = "D9EAD3"

REPORT_COLUMNS = [
    "SRNo",
    "Problem Summary",
    "Done?",
    "Planned Date",
    "Planned Days",
    "Ticket Age",
    "ResolveBy",
    "Resolve Days",
    "Email Inactivity",
    "Latest Direction",
    "Received",
    "Sent",
    "Current Handler",
    "Customer Severity",
    "Status",
]


def _to_excel_value(column: str, value: Any) -> Any:
    if value is None or value == "":
        return None
    if column == "SRNo":
        return str(value)
    if column == "Planned Date":
        parsed = parse_date(value)
        return parsed if parsed else value
    if column in {"Report Date", "ResolveBy", "Resolve By Suspend"}:
        parsed = parse_datetime(value)
        if parsed is not None and parsed.tzinfo is not None:
            parsed = parsed.astimezone().replace(tzinfo=None)
        return parsed or value
    return value


def _color_from_snapshot(value: dict[str, Any] | None) -> Color:
    if not value or not value.get("type"):
        return Color()
    color_type = value.get("type")
    kwargs: dict[str, Any] = {"type": color_type, "tint": value.get("tint", 0.0)}
    if color_type == "rgb":
        kwargs["rgb"] = value.get("value")
    elif color_type == "indexed":
        kwargs["indexed"] = value.get("value")
    elif color_type == "theme":
        kwargs["theme"] = value.get("value")
    elif color_type == "auto":
        kwargs["auto"] = value.get("value")
    return Color(**kwargs)


def _apply_snapshot(cell: Any, snapshot: dict[str, Any] | None) -> None:
    if not snapshot:
        return
    fill = snapshot.get("fill", {})
    cell.fill = PatternFill(
        fill_type=fill.get("fill_type"),
        fgColor=_color_from_snapshot(fill.get("fg_color")),
        bgColor=_color_from_snapshot(fill.get("bg_color")),
    )
    font = snapshot.get("font", {})
    cell.font = Font(
        name=font.get("name"),
        sz=font.get("size"),
        bold=font.get("bold", False),
        italic=font.get("italic", False),
        underline=font.get("underline"),
        strike=font.get("strike", False),
        color=_color_from_snapshot(font.get("color")),
        vertAlign=font.get("vertAlign"),
    )


def _header_style(cell: Any) -> None:
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.font = Font(color=WHITE, bold=True)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _set_column_widths(worksheet: Any, headers: list[str]) -> None:
    widths = {
        "SRNo": 12,
        "Problem Summary": 58,
        "Planned Date": 14,
        "Site": 10,
        "Cloud": 10,
        "Model": 18,
        "Device": 22,
        "Slot": 18,
        "Part": 16,
        "BOM": 16,
        "Old SN": 24,
        "New SN": 24,
        "RelatedSR": 16,
        "Notes": 36,
        "Report Date": 20,
        "Customer Contact": 22,
        "Customer Severity": 18,
        "Product": 30,
        "Current Handler": 32,
        "Status": 24,
        "ResolveBy": 20,
        "Resolve By Suspend": 20,
        "Spare": 10,
        "Done?": 10,
    }
    for index, header in enumerate(headers, start=1):
        worksheet.column_dimensions[worksheet.cell(1, index).column_letter].width = widths.get(header, 16)


def _write_ticket_sheet(
    workbook: Workbook,
    tickets: Iterable[dict[str, Any]],
    *,
    header_order: list[str] | None = None,
    sheet_name: str = "Pendings",
) -> Any:
    worksheet = workbook.active
    worksheet.title = sheet_name
    headers = list(header_order or PENDING_COLUMNS)
    worksheet.append(headers)
    for cell in worksheet[1]:
        _header_style(cell)
    worksheet.row_dimensions[1].height = 30
    ordered = sorted(tickets, key=lambda ticket: int(ticket["ticket_id"]), reverse=True)
    for ticket in ordered:
        row = [_to_excel_value(column, ticket_cell_value(ticket, column)) for column in headers]
        worksheet.append(row)
        row_number = worksheet.max_row
        styles = ticket.get("local", {}).get("presentation", {}).get("cell_styles", {})
        for column_number, header in enumerate(headers, start=1):
            cell = worksheet.cell(row_number, column_number)
            _apply_snapshot(cell, styles.get(header))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if header in {"Planned Date"} and isinstance(cell.value, date):
                cell.number_format = "yyyy-mm-dd"
            elif header in {"Report Date", "ResolveBy", "Resolve By Suspend"} and isinstance(cell.value, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm:ss"
            if header == "SRNo":
                cell.number_format = "@"
                url = ticket.get("upstream", {}).get("sr_url")
                if url:
                    cell.hyperlink = url
                    if header not in styles:
                        cell.font = Font(color="0563C1", underline="single")
        worksheet.row_dimensions[row_number].height = 33
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:{worksheet.cell(1, len(headers)).coordinate}"
    _set_column_widths(worksheet, headers)
    worksheet.sheet_view.showGridLines = False
    done_column = headers.index("Done?") + 1
    validation = DataValidation(type="list", formula1='"Y,N,P,?"', allow_blank=True)
    validation.error = "Use Y, N, P, or ?. Other values normalize to N at the next startup."
    validation.errorTitle = "Zeus workflow code"
    validation.showErrorMessage = True
    worksheet.add_data_validation(validation)
    validation.add(
        f"{worksheet.cell(2, done_column).coordinate}:"
        f"{worksheet.cell(max(2, worksheet.max_row), done_column).coordinate}"
    )
    return worksheet


def _fill_for_color(name: str | None) -> PatternFill:
    mapping = {"red": RED, "yellow": YELLOW, "amber": AMBER, "grey": GREY, "green": GREEN}
    return PatternFill("solid", fgColor=mapping[name]) if name in mapping else PatternFill()


def _report_row(ticket: dict[str, Any], config: dict[str, Any]) -> tuple[list[Any], dict[str, str | None]]:
    facts = aging_for_ticket(ticket, config)
    upstream = ticket.get("upstream", {}).get("fields", {})
    local = ticket.get("local", {}).get("fields", {})
    email = ticket.get("email", {})
    values = [
        ticket["ticket_id"],
        upstream.get("Problem Summary"),
        workflow_code(ticket),
        facts.planned_label,
        facts.planned_days,
        facts.ticket_age_days,
        facts.resolve_label,
        facts.resolve_days,
        facts.communication_label,
        email.get("last_direction") or ("No email found" if not email.get("last_activity_at") else None),
        int(email.get("total_received") or 0),
        int(email.get("total_sent") or 0),
        upstream.get("Current Handler"),
        upstream.get("Customer Severity"),
        upstream.get("Status"),
    ]
    colors = {
        "Planned Date": facts.planned_color,
        "Planned Days": facts.planned_color,
        "Ticket Age": facts.ticket_age_color,
        "ResolveBy": facts.resolve_color,
        "Resolve Days": facts.resolve_color,
        "Email Inactivity": facts.communication_color,
        "Latest Direction": "grey" if facts.communication_inactivity_days is None else None,
    }
    return values, colors


def _write_report_sheet(
    workbook: Workbook,
    active: list[dict[str, Any]],
    closure_pending: list[dict[str, Any]],
    config: dict[str, Any],
) -> Any:
    worksheet = workbook.create_sheet("Report")
    worksheet.sheet_view.showGridLines = False
    worksheet["A1"] = "Zeus ticket report"
    worksheet["A1"].font = Font(size=18, bold=True, color=WHITE)
    worksheet["A1"].fill = PatternFill("solid", fgColor=NAVY)
    worksheet.merge_cells("A1:O1")
    counts = {code: 0 for code in ("Y", "N", "P", "?")}
    unplanned = overdue = no_email = 0
    for ticket in active:
        counts[workflow_code(ticket)] += 1
        facts = aging_for_ticket(ticket, config)
        unplanned += facts.planned_state == "unplanned"
        overdue += facts.planned_state == "overdue"
        no_email += facts.communication_inactivity_days is None
    summaries = [
        ("Active", len(active)),
        ("Y", counts["Y"]),
        ("N", counts["N"]),
        ("P", counts["P"]),
        ("?", counts["?"]),
        ("Overdue", overdue),
        ("Unplanned", unplanned),
        ("No email", no_email),
        ("Pending closure", len(closure_pending)),
    ]
    for column, (label, value) in enumerate(summaries, start=1):
        worksheet.cell(3, column, label)
        worksheet.cell(4, column, value)
        worksheet.cell(3, column).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        worksheet.cell(3, column).font = Font(bold=True, color=TEXT)
        worksheet.cell(4, column).font = Font(bold=True, size=14, color=BLUE)
        worksheet.cell(3, column).alignment = Alignment(horizontal="center")
        worksheet.cell(4, column).alignment = Alignment(horizontal="center")

    start_row = 7
    worksheet.cell(start_row, 1, "Active tickets")
    worksheet.cell(start_row, 1).font = Font(bold=True, size=13, color=BLUE)
    header_row = start_row + 1
    for column, header in enumerate(REPORT_COLUMNS, start=1):
        worksheet.cell(header_row, column, header)
        _header_style(worksheet.cell(header_row, column))
    row_number = header_row
    for ticket in sorted(active, key=lambda item: report_sort_key(item, config)):
        values, colors = _report_row(ticket, config)
        row_number += 1
        for column, value in enumerate(values, start=1):
            cell = worksheet.cell(row_number, column, value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            header = REPORT_COLUMNS[column - 1]
            color = colors.get(header)
            if color:
                cell.fill = _fill_for_color(color)
        worksheet.row_dimensions[row_number].height = 32

    active_end_row = max(header_row, row_number)
    worksheet.auto_filter.ref = (
        f"A{header_row}:{get_column_letter(len(REPORT_COLUMNS))}{active_end_row}"
    )

    closure_title_row = row_number + 3
    worksheet.cell(closure_title_row, 1, "Pending closure")
    worksheet.cell(closure_title_row, 1).font = Font(bold=True, size=13, color=BLUE)
    closure_header_row = closure_title_row + 1
    for column, header in enumerate(REPORT_COLUMNS, start=1):
        worksheet.cell(closure_header_row, column, header)
        _header_style(worksheet.cell(closure_header_row, column))
    row_number = closure_header_row
    for ticket in sorted(closure_pending, key=lambda item: int(item["ticket_id"]), reverse=True):
        values, colors = _report_row(ticket, config)
        row_number += 1
        for column, value in enumerate(values, start=1):
            cell = worksheet.cell(row_number, column, value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            color = colors.get(REPORT_COLUMNS[column - 1])
            if color:
                cell.fill = _fill_for_color(color)
    widths = [12, 58, 9, 14, 14, 12, 14, 13, 18, 17, 11, 10, 32, 18, 25]
    for column, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(column)].width = width
    worksheet.freeze_panes = f"A{header_row + 1}"
    return worksheet


def _save_pending_temp(
    store: ZeusStore,
    path: Path,
    active: list[dict[str, Any]],
    closure_pending: list[dict[str, Any]],
    *,
    header_order: list[str],
) -> None:
    workbook = Workbook()
    _write_ticket_sheet(workbook, active, header_order=header_order)
    _write_report_sheet(workbook, active, closure_pending, store.config)
    workbook.save(path)
    workbook.close()


def _append_closed_rows(
    store: ZeusStore,
    path: Path,
    closure_pending: list[dict[str, Any]],
    *,
    existing_path: Path | None,
    header_order: list[str],
) -> set[str]:
    if existing_path is not None:
        shutil.copy2(existing_path, path)
        workbook = load_workbook(path, data_only=False, read_only=False, keep_links=True)
        worksheet = workbook.worksheets[0]
        existing = read_closed(existing_path)
        existing_ids = set(existing.records)
    else:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Closed"
        worksheet.append(header_order)
        for cell in worksheet[1]:
            _header_style(cell)
        existing_ids = set()
    for ticket in sorted(closure_pending, key=lambda item: int(item["ticket_id"]), reverse=True):
        if ticket["ticket_id"] in existing_ids:
            continue
        values = [_to_excel_value(column, ticket_cell_value(ticket, column)) for column in header_order]
        worksheet.append(values)
        row_number = worksheet.max_row
        styles = ticket.get("local", {}).get("presentation", {}).get("cell_styles", {})
        for column_number, header in enumerate(header_order, start=1):
            cell = worksheet.cell(row_number, column_number)
            _apply_snapshot(cell, styles.get(header))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if header == "SRNo":
                cell.number_format = "@"
                url = ticket.get("upstream", {}).get("sr_url")
                if url:
                    cell.hyperlink = url
            if header == "Planned Date" and isinstance(cell.value, date):
                cell.number_format = "yyyy-mm-dd"
            elif header in {"Report Date", "ResolveBy", "Resolve By Suspend"} and isinstance(cell.value, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm:ss"
        existing_ids.add(ticket["ticket_id"])
    if existing_path is None:
        _set_column_widths(worksheet, header_order)
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = f"A1:{worksheet.cell(1, len(header_order)).coordinate}"
        worksheet.sheet_view.showGridLines = False
    workbook.save(path)
    workbook.close()
    return existing_ids


def _backup_directory(workbook_directory: Path) -> Path:
    return workbook_directory / "Zeus Backups"


def list_pendings_backups(workbook_directory: Path) -> list[Path]:
    directory = _backup_directory(workbook_directory)
    if not directory.exists():
        return []
    return sorted(
        directory.glob("Pendings_*.xlsx"),
        key=lambda path: path.stat().st_mtime,
        reverse=True,
    )


def _create_paired_backups(
    workbook_directory: Path,
    pending_path: Path,
    closed_path: Path,
    *,
    retention: int,
) -> dict[str, str | None]:
    directory = _backup_directory(workbook_directory)
    directory.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    result: dict[str, str | None] = {"pendings": None, "closed": None}
    if pending_path.exists():
        destination = directory / f"Pendings_{timestamp}.xlsx"
        shutil.copy2(pending_path, destination)
        result["pendings"] = str(destination)
    if closed_path.exists():
        destination = directory / f"Closed_{timestamp}.xlsx"
        shutil.copy2(closed_path, destination)
        result["closed"] = str(destination)
    for prefix in ("Pendings_", "Closed_"):
        files = sorted(
            directory.glob(f"{prefix}*.xlsx"),
            key=lambda path: path.stat().st_mtime,
            reverse=True,
        )
        for old in files[retention:]:
            old.unlink(missing_ok=True)
    return result


def _replace_pair(
    pending_temp: Path,
    closed_temp: Path,
    pending_target: Path,
    closed_target: Path,
) -> None:
    token = uuid.uuid4().hex
    pending_rollback = pending_target.parent / f".zeus-rollback-pendings-{token}.xlsx"
    closed_rollback = closed_target.parent / f".zeus-rollback-closed-{token}.xlsx"
    had_pending = pending_target.exists()
    had_closed = closed_target.exists()
    if had_pending:
        shutil.copy2(pending_target, pending_rollback)
    if had_closed:
        shutil.copy2(closed_target, closed_rollback)
    pending_replaced = closed_replaced = False
    try:
        os.replace(pending_temp, pending_target)
        pending_replaced = True
        os.replace(closed_temp, closed_target)
        closed_replaced = True
    except Exception:
        if pending_replaced:
            if had_pending:
                os.replace(pending_rollback, pending_target)
            else:
                pending_target.unlink(missing_ok=True)
        if closed_replaced:
            if had_closed:
                os.replace(closed_rollback, closed_target)
            else:
                closed_target.unlink(missing_ok=True)
        raise
    finally:
        pending_rollback.unlink(missing_ok=True)
        closed_rollback.unlink(missing_ok=True)
        pending_temp.unlink(missing_ok=True)
        closed_temp.unlink(missing_ok=True)


def _snapshot_for_tickets(
    tickets: Iterable[dict[str, Any]],
    *,
    pending_path: Path,
    header_order: list[str],
) -> dict[str, Any]:
    tickets = list(tickets)
    return {
        "filename": pending_path.name,
        "sha256": sha256_file(pending_path),
        "captured_at": iso_now(),
        "header_order": header_order,
        "ticket_ids": sorted((ticket["ticket_id"] for ticket in tickets), reverse=True),
        "protected_values": {
            ticket["ticket_id"]: deepcopy(ticket.get("upstream", {}).get("fields", {}))
            for ticket in tickets
        },
    }


def _recreated_pendings_headers(store: ZeusStore) -> list[str]:
    snapshot = store.state().get("publication_state", {}).get("pendings_snapshot")
    saved = snapshot.get("header_order") if isinstance(snapshot, dict) else None
    if (
        isinstance(saved, list)
        and len(saved) == len(PENDING_COLUMNS)
        and set(saved) == set(PENDING_COLUMNS)
    ):
        return [str(value) for value in saved]
    return list(PENDING_COLUMNS)


def _finalize_pendings_recreation(
    store: ZeusStore,
    journal: dict[str, Any],
) -> None:
    summary = {
        "path": journal["path"],
        "rows": journal["rows"],
        "source": "markdown_database",
        "sha256": journal["new_hash"],
    }
    with store.transaction(
        "pendings-recreation-finalize",
        summary,
        backup=False,
    ) as staging:
        state = store.state(staging)
        state.setdefault("publication_state", {})["pendings_snapshot"] = journal[
            "pendings_snapshot"
        ]
        pendings_state = state.setdefault("pendings_state", {})
        pendings_state["last_error"] = None
        pendings_state["last_recreated_at"] = journal["recreated_at"]
        pendings_state["last_recreated_sha256"] = journal["new_hash"]
        (staging / "pendings_recreation_journal.json").unlink(missing_ok=True)
        atomic_write_json(staging / "state.json", state)


def recover_pendings_recreation(
    store: ZeusStore,
    workbook_directory: Path,
) -> dict[str, Any] | None:
    """Finish or safely discard an interrupted database-driven recreation."""

    journal_path = store.current / "pendings_recreation_journal.json"
    if not journal_path.exists():
        return None
    import json

    journal = json.loads(journal_path.read_text(encoding="utf-8"))
    target = workbook_directory.expanduser().resolve() / "Pendings.xlsx"
    actual = sha256_file(target) if target.is_file() else None
    if actual == journal.get("new_hash"):
        _finalize_pendings_recreation(store, journal)
        return {
            "recovered": "finalized",
            "created": True,
            "path": str(target),
            "rows": int(journal.get("rows") or 0),
            "source": "markdown_database",
            "sha256": journal.get("new_hash"),
            "restoredBackup": False,
            "notice": journal.get("notice"),
        }
    if actual is None:
        with store.transaction(
            "pendings-recreation-abort",
            {"reason": "target_not_replaced"},
            backup=False,
        ) as staging:
            (staging / "pendings_recreation_journal.json").unlink(missing_ok=True)
        return {"recovered": "aborted", "created": False}
    raise WorkbookPublicationError(
        "An interrupted Pendings recreation found a different Pendings.xlsx. "
        "Zeus preserved that file; query it after confirming its contents."
    )


def recreate_pendings_from_database(
    store: ZeusStore,
    workbook_directory: Path,
) -> dict[str, Any]:
    """Create a missing Pendings workbook from current Markdown records only.

    This is intentionally not operational publication: it does not create or
    change Closed.xlsx, finalize closure-pending tickets, or restore a backup.
    """

    store.ensure_layout()
    workbook_directory = workbook_directory.expanduser().resolve()
    workbook_directory.mkdir(parents=True, exist_ok=True)
    recovered = recover_pendings_recreation(store, workbook_directory)
    target = workbook_directory / "Pendings.xlsx"
    if target.is_file():
        if recovered and recovered.get("created"):
            return recovered
        return {
            "created": False,
            "path": str(target),
            "source": "existing_workbook",
        }

    tickets = list(store.iter_tickets())
    active = [
        ticket
        for ticket in tickets
        if ticket.get("lifecycle", {}).get("status") == "active"
    ]
    closure_pending = [ticket for ticket in tickets if ticket not in active]
    header_order = _recreated_pendings_headers(store)
    token = uuid.uuid4().hex
    temporary = workbook_directory / f".zeus-recreate-pendings-{token}.xlsx"
    workbook = Workbook()
    try:
        _write_ticket_sheet(workbook, tickets, header_order=header_order)
        _write_report_sheet(workbook, active, closure_pending, store.config)
        workbook.save(temporary)
    finally:
        workbook.close()

    try:
        generated = read_pendings(temporary)
        expected_ids = {ticket["ticket_id"] for ticket in tickets}
        if set(generated.records) != expected_ids:
            raise WorkbookPublicationError(
                "Recreated Pendings.xlsx failed ticket-ID verification"
            )
        recreated_at = iso_now()
        new_hash = generated.sha256
        notice = (
            "Pendings.xlsx was missing and Zeus recreated it from "
            f"{len(tickets)} Markdown ticket record(s). No backup was restored."
        )
        snapshot = _snapshot_for_tickets(
            tickets,
            pending_path=temporary,
            header_order=header_order,
        )
        snapshot["filename"] = "Pendings.xlsx"
        journal = {
            "schema_version": 1,
            "path": str(target),
            "recreated_at": recreated_at,
            "rows": len(tickets),
            "active_rows": len(active),
            "closure_pending_rows": len(closure_pending),
            "new_hash": new_hash,
            "pendings_snapshot": snapshot,
            "notice": notice,
        }
        with store.transaction(
            "pendings-recreation-prepare",
            {
                "path": str(target),
                "rows": len(tickets),
                "source": "markdown_database",
            },
        ) as staging:
            atomic_write_json(
                staging / "pendings_recreation_journal.json",
                journal,
            )
        if target.exists():
            with store.transaction(
                "pendings-recreation-abort",
                {"reason": "workbook_appeared_during_recreation"},
                backup=False,
            ) as staging:
                (staging / "pendings_recreation_journal.json").unlink(
                    missing_ok=True
                )
            raise WorkbookPublicationError(
                "Pendings.xlsx appeared while Zeus was recreating it. Zeus preserved "
                "the existing file; query again to import it safely."
            )
        os.replace(temporary, target)
        _finalize_pendings_recreation(store, journal)
        return {
            "created": True,
            "path": str(target),
            "rows": len(tickets),
            "activeRows": len(active),
            "closurePendingRows": len(closure_pending),
            "source": "markdown_database",
            "sha256": new_hash,
            "restoredBackup": False,
            "notice": notice,
        }
    except Exception:
        # If replacement already succeeded, the journal deliberately remains
        # available for startup recovery. Otherwise no managed file changed.
        if not target.is_file():
            try:
                with store.transaction(
                    "pendings-recreation-abort",
                    {"reason": "recreation_failed_before_replace"},
                    backup=False,
                ) as staging:
                    (staging / "pendings_recreation_journal.json").unlink(
                        missing_ok=True
                    )
            except Exception:
                pass
        raise
    finally:
        temporary.unlink(missing_ok=True)


def _filter_staging_for_closed(store: ZeusStore, staging: Path, closing_ids: set[str]) -> None:
    path = store.staging_file(staging)
    if not path.exists():
        return
    messages: list[dict[str, Any]] = []
    for line in path.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        import json

        message = json.loads(line)
        remaining = [value for value in message.get("ticket_ids", []) if value not in closing_ids]
        if remaining:
            message["ticket_ids"] = remaining
            messages.append(message)
    text = "".join(__import__("json").dumps(message, ensure_ascii=False) + "\n" for message in messages)
    atomic_write_text(path, text)


def _finalize_publication(store: ZeusStore, journal: dict[str, Any]) -> None:
    closing_ids = {str(value) for value in journal.get("closing_ids", [])}
    summary = {"closing_ids": sorted(closing_ids, reverse=True)}
    with store.transaction("workbook-publication-finalize", summary) as staging:
        for ticket_id in closing_ids:
            store.delete_ticket_bundle(ticket_id, staging)
        _filter_staging_for_closed(store, staging, closing_ids)
        atomic_write_json(staging / "closed_index.json", journal["closed_index"])
        state = store.state(staging)
        state.setdefault("publication_state", {})["last_published_at"] = journal[
            "published_at"
        ]
        state["publication_state"]["pendings_snapshot"] = journal[
            "pendings_snapshot"
        ]
        (staging / "publication_journal.json").unlink(missing_ok=True)
        atomic_write_json(staging / "state.json", state)
    if closing_ids:
        # Internal state snapshots may contain email bodies.  Workbook backups
        # remain available because email data never enters either workbook.
        summary["state_backups_purged_for_email_policy"] = (
            store.purge_state_backups_for_closed_tickets()
        )


def recover_publication(store: ZeusStore, workbook_directory: Path) -> dict[str, Any] | None:
    journal_path = store.current / "publication_journal.json"
    if not journal_path.exists():
        return None
    import json

    journal = json.loads(journal_path.read_text(encoding="utf-8"))
    pending = workbook_directory / "Pendings.xlsx"
    closed = workbook_directory / "Closed.xlsx"
    actual_pending = sha256_file(pending) if pending.exists() else None
    actual_closed = sha256_file(closed) if closed.exists() else None
    if (
        actual_pending == journal.get("new_hashes", {}).get("pendings")
        and actual_closed == journal.get("new_hashes", {}).get("closed")
    ):
        _finalize_publication(store, journal)
        return {"recovered": "finalized", "closing_ids": journal.get("closing_ids", [])}
    if (
        actual_pending == journal.get("old_hashes", {}).get("pendings")
        and actual_closed == journal.get("old_hashes", {}).get("closed")
    ):
        with store.transaction("workbook-publication-abort", {}, backup=False) as staging:
            (staging / "publication_journal.json").unlink(missing_ok=True)
        return {"recovered": "aborted", "closing_ids": []}
    raise WorkbookPublicationError(
        "An interrupted workbook publication left mixed files. Restore the paired "
        "Zeus Backups copies before continuing."
    )


def preview_pendings_restore(
    store: ZeusStore, backup_path: Path
) -> dict[str, Any]:
    backup = read_pendings(backup_path)
    existing_ids = set(store.iter_ticket_ids())
    applicable = sorted(set(backup.records) & existing_ids, reverse=True)
    ignored = sorted(set(backup.records) - existing_ids, reverse=True)
    changed: list[str] = []
    for ticket_id in applicable:
        current = store.read_ticket(ticket_id).get("local", {})
        if current != backup.records[ticket_id]["local"]:
            changed.append(ticket_id)
    return {
        "backup": str(backup.path),
        "applicable_ids": applicable,
        "changed_ids": changed,
        "ignored_closed_or_unknown_ids": ignored,
        "protected_fields_restored": False,
        "email_restored": False,
    }


def _finalize_pendings_restore(store: ZeusStore, journal: dict[str, Any]) -> None:
    summary = {
        "backup": journal["backup"],
        "changed_ids": sorted(journal["local_by_id"], reverse=True),
    }
    with store.transaction("pendings-backup-restore-finalize", summary) as staging:
        for ticket_id, local in journal["local_by_id"].items():
            ticket = store.read_ticket(ticket_id, staging)
            ticket["local"] = deepcopy(local)
            store.write_ticket_bundle(staging, ticket)
        state = store.state(staging)
        state.setdefault("publication_state", {})["pendings_snapshot"] = journal[
            "pendings_snapshot"
        ]
        (staging / "pendings_restore_journal.json").unlink(missing_ok=True)
        atomic_write_json(staging / "state.json", state)


def recover_pendings_restore(
    store: ZeusStore, workbook_directory: Path
) -> dict[str, Any] | None:
    journal_path = store.current / "pendings_restore_journal.json"
    if not journal_path.exists():
        return None
    import json

    journal = json.loads(journal_path.read_text(encoding="utf-8"))
    target = workbook_directory / "Pendings.xlsx"
    actual = sha256_file(target) if target.exists() else None
    if actual == journal.get("new_hash"):
        _finalize_pendings_restore(store, journal)
        return {"recovered": "finalized", "backup": journal.get("backup")}
    if actual == journal.get("old_hash"):
        with store.transaction("pendings-backup-restore-abort", {}, backup=False) as staging:
            (staging / "pendings_restore_journal.json").unlink(missing_ok=True)
        return {"recovered": "aborted", "backup": journal.get("backup")}
    raise WorkbookPublicationError(
        "An interrupted Pendings restore left an unrecognized workbook. Use the "
        "latest Pendings backup before continuing."
    )


def publish_operational_workbooks(
    store: ZeusStore,
    workbook_directory: Path,
    *,
    create_missing: bool = False,
) -> dict[str, Any]:
    store.ensure_layout()
    workbook_directory = workbook_directory.expanduser().resolve()
    workbook_directory.mkdir(parents=True, exist_ok=True)
    recover_pendings_restore(store, workbook_directory)
    recover_publication(store, workbook_directory)
    pending_path = workbook_directory / "Pendings.xlsx"
    closed_path = workbook_directory / "Closed.xlsx"

    pending_book: ManagedWorkbook | None = None
    closed_book: ManagedWorkbook | None = None
    if pending_path.exists():
        pending_book = read_pendings(pending_path)
        import_pendings(store, pending_path)
    elif not create_missing:
        raise WorkbookPublicationError(
            "Pendings.xlsx is missing. Explicitly choose 'Create a fresh Pendings.xlsx' to continue."
        )
    if closed_path.exists():
        closed_book = read_closed(closed_path)
        validate_closed_against_index(closed_book, store.closed_index())
    elif not create_missing:
        raise WorkbookPublicationError(
            "Closed.xlsx is missing. Explicitly confirm creation to continue."
        )

    active = list(store.iter_tickets(status="active"))
    closure_pending = list(store.iter_tickets(status="closure_pending"))
    pending_headers = pending_book.header_order if pending_book else list(PENDING_COLUMNS)
    closed_headers = closed_book.header_order if closed_book else list(PENDING_COLUMNS)
    token = uuid.uuid4().hex
    pending_temp = workbook_directory / f".zeus-pendings-{token}.xlsx"
    closed_temp = workbook_directory / f".zeus-closed-{token}.xlsx"
    try:
        _save_pending_temp(
            store,
            pending_temp,
            active,
            closure_pending,
            header_order=pending_headers,
        )
        closed_ids = _append_closed_rows(
            store,
            closed_temp,
            closure_pending,
            existing_path=closed_path if closed_path.exists() else None,
            header_order=closed_headers,
        )
        generated_pending = read_pendings(pending_temp)
        generated_closed = read_closed(closed_temp)
        expected_active = {ticket["ticket_id"] for ticket in active}
        if set(generated_pending.records) != expected_active:
            raise WorkbookPublicationError("Generated Pendings.xlsx failed ID verification")
        if set(generated_closed.records) != closed_ids:
            raise WorkbookPublicationError("Generated Closed.xlsx failed ID verification")

        if pending_book and sha256_file(pending_path) != pending_book.sha256:
            raise WorkbookPublicationError(
                "Pendings.xlsx changed while publication was being prepared. Retry so the latest saved edits are imported."
            )
        if closed_book and sha256_file(closed_path) != closed_book.sha256:
            raise WorkbookPublicationError(
                "Closed.xlsx changed while publication was being prepared. Retry without modifying it concurrently."
            )

        retention = int(store.config.get("excel", {}).get("backup_retention_count", 10))
        backups = _create_paired_backups(
            workbook_directory,
            pending_path,
            closed_path,
            retention=retention,
        )
        published_at = iso_now()
        journal = {
            "schema_version": 1,
            "published_at": published_at,
            "closing_ids": sorted(
                (ticket["ticket_id"] for ticket in closure_pending), reverse=True
            ),
            "old_hashes": {
                "pendings": sha256_file(pending_path) if pending_path.exists() else None,
                "closed": sha256_file(closed_path) if closed_path.exists() else None,
            },
            "new_hashes": {
                "pendings": sha256_file(pending_temp),
                "closed": sha256_file(closed_temp),
            },
            "pendings_snapshot": {
                "filename": "Pendings.xlsx",
                "sha256": sha256_file(pending_temp),
                "captured_at": published_at,
                "header_order": pending_headers,
                "ticket_ids": sorted(expected_active, reverse=True),
                "protected_values": {
                    ticket["ticket_id"]: deepcopy(ticket.get("upstream", {}).get("fields", {}))
                    for ticket in active
                },
            },
            "closed_index": {
                "schema_version": 1,
                "ticket_ids": sorted(closed_ids, reverse=True),
                "row_count": len(closed_ids),
                "validated_at": published_at,
                "workbook_sha256": sha256_file(closed_temp),
            },
            "backups": backups,
        }
        with store.transaction("workbook-publication-prepare", journal) as staging:
            atomic_write_json(staging / "publication_journal.json", journal)
        try:
            _replace_pair(pending_temp, closed_temp, pending_path, closed_path)
        except Exception:
            with store.transaction("workbook-publication-abort", {}, backup=False) as staging:
                (staging / "publication_journal.json").unlink(missing_ok=True)
            raise
        # The hashes in the snapshot were calculated from the staging file;
        # atomic replacement does not change bytes.
        _finalize_publication(store, journal)
        return {
            "published_at": published_at,
            "active_rows": len(active),
            "closed_appended": len(closure_pending),
            "final_closed_rows": len(closed_ids),
            "pendings_path": str(pending_path),
            "closed_path": str(closed_path),
            "backups": backups,
        }
    finally:
        pending_temp.unlink(missing_ok=True)
        closed_temp.unlink(missing_ok=True)


def restore_pendings_backup(
    store: ZeusStore,
    backup_path: Path,
    workbook_directory: Path,
) -> dict[str, Any]:
    """Restore local fields/styles only, then rewrite current Pendings."""

    backup = read_pendings(backup_path)
    summary = preview_pendings_restore(store, backup_path)
    applicable = summary["applicable_ids"]
    local_by_id = {
        ticket_id: deepcopy(backup.records[ticket_id]["local"])
        for ticket_id in applicable
    }
    # Build and verify the exact post-restore workbook before either authority
    # is changed.  Closure-pending Markdown records remain present here.
    workbook_directory = workbook_directory.expanduser().resolve()
    workbook_directory.mkdir(parents=True, exist_ok=True)
    recover_pendings_restore(store, workbook_directory)
    current_path = workbook_directory / "Pendings.xlsx"
    token = uuid.uuid4().hex
    temporary = workbook_directory / f".zeus-restore-{token}.xlsx"
    tickets = list(store.iter_tickets())
    for ticket in tickets:
        if ticket["ticket_id"] in local_by_id:
            ticket["local"] = deepcopy(local_by_id[ticket["ticket_id"]])
    workbook = Workbook()
    _write_ticket_sheet(workbook, tickets, header_order=backup.header_order)
    _write_report_sheet(
        workbook,
        [ticket for ticket in tickets if ticket["lifecycle"]["status"] == "active"],
        [ticket for ticket in tickets if ticket["lifecycle"]["status"] == "closure_pending"],
        store.config,
    )
    workbook.save(temporary)
    workbook.close()
    read_pendings(temporary)
    new_hash = sha256_file(temporary)
    published_at = iso_now()
    snapshot = {
        "filename": "Pendings.xlsx",
        "sha256": new_hash,
        "captured_at": published_at,
        "header_order": backup.header_order,
        "ticket_ids": sorted((ticket["ticket_id"] for ticket in tickets), reverse=True),
        "protected_values": {
            ticket["ticket_id"]: deepcopy(ticket.get("upstream", {}).get("fields", {}))
            for ticket in tickets
        },
    }
    journal = {
        "schema_version": 1,
        "backup": str(backup_path),
        "old_hash": sha256_file(current_path) if current_path.exists() else None,
        "new_hash": new_hash,
        "local_by_id": local_by_id,
        "pendings_snapshot": snapshot,
    }
    _create_paired_backups(
        workbook_directory,
        current_path,
        workbook_directory / "Closed.xlsx",
        retention=int(store.config.get("excel", {}).get("backup_retention_count", 10)),
    )
    with store.transaction("pendings-backup-restore-prepare", summary) as staging:
        atomic_write_json(staging / "pendings_restore_journal.json", journal)
    try:
        os.replace(temporary, current_path)
    except Exception:
        with store.transaction("pendings-backup-restore-abort", {}, backup=False) as staging:
            (staging / "pendings_restore_journal.json").unlink(missing_ok=True)
        temporary.unlink(missing_ok=True)
        raise
    _finalize_pendings_restore(store, journal)
    summary["rewritten_path"] = str(current_path)
    return summary


# Compatibility helper.  It now creates the two managed files and returns the
# Pendings path rather than a formula-driven third workbook.
def export_progress(store: ZeusStore, output: Path) -> Path:
    directory = output.expanduser().resolve().parent
    result = publish_operational_workbooks(store, directory, create_missing=True)
    source = Path(result["pendings_path"])
    if source != output.expanduser().resolve():
        shutil.copy2(source, output)
    return output


# Preview scripts imported this symbol; generated fields now live on Sheet 2.
EXPORT_COLUMNS = PENDING_COLUMNS
