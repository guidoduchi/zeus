from __future__ import annotations

import os
import shutil
import tempfile
from copy import deepcopy
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from ..excel_export import (
    WorkbookPublicationError,
    recreate_pendings_from_database,
    write_spare_parts_sheet,
)
from ..excel_import import (
    WorkbookValidationError,
    read_pendings,
    validate_pendings_against_state,
)
from ..reconcile import import_pendings
from ..store import ZeusStore
from ..tickets import (
    EDITABLE_LOCAL_COLUMNS,
    LEGACY_SPARE_COLUMNS,
    SPARE_PARTS_CHANGE_KEY,
    normalize_done,
    normalize_local,
    normalize_spare_parts,
)
from ..utils import parse_date, sha256_file
from .errors import ConflictError, FeatureUnavailableError, ValidationError
from .serialization import ticket_revision


class PendingsConflictError(ConflictError):
    code = "pendings_conflict"


class TicketRevisionConflictError(ConflictError):
    code = "ticket_revision_conflict"


def _normalize_local_changes(changes: dict[str, Any]) -> dict[str, Any]:
    if not isinstance(changes, dict) or not changes:
        raise ValidationError("Choose at least one Pendings field to update")
    if "Spare" in changes:
        raise ValidationError("Spare is an export-only value and cannot be edited directly")
    allowed = set(EDITABLE_LOCAL_COLUMNS) | {SPARE_PARTS_CHANGE_KEY}
    unknown = sorted(set(changes) - allowed)
    if unknown:
        raise ValidationError(
            "Only Pendings-owned fields can be edited in Zeus",
            details={"unknownFields": unknown},
        )
    prepared: dict[str, Any] = {}
    for key, value in changes.items():
        if key == SPARE_PARTS_CHANGE_KEY:
            prepared[key] = _normalize_spare_parts_change(value)
            continue
        if isinstance(value, (dict, list, tuple, set)):
            raise ValidationError(f"{key} must be a text, date, number, or blank value")
        if key == "Done?":
            code, corrected = normalize_done(value)
            if corrected:
                raise ValidationError("Done? must be Y, N, P, or ?")
            prepared[key] = code
            continue
        if value is None:
            prepared[key] = None
            continue
        text = str(value)
        if len(text) > 100_000:
            raise ValidationError(f"{key} is too long")
        text = text if key == "Notes" else text.strip()
        if key == "Planned Date" and text.strip():
            planned = parse_date(text)
            if planned is None:
                raise ValidationError("Planned Date must be a recognizable calendar date")
            prepared[key] = planned
        else:
            prepared[key] = text if text else None
    return prepared


def _normalize_spare_parts_change(value: Any) -> list[dict[str, Any]]:
    if not isinstance(value, list):
        raise ValidationError("Spare Parts must be a list of devices")
    if len(value) > 200:
        raise ValidationError("A ticket cannot contain more than 200 damaged devices")
    device_keys = {"device", "model", "parts"}
    part_keys = {"slot", "part", "bom", "faulty_sn", "new_sn"}
    total_parts = 0
    for device_index, device in enumerate(value, start=1):
        if not isinstance(device, dict):
            raise ValidationError(f"Spare Parts device {device_index} must be an object")
        unknown_device = sorted(set(device) - device_keys)
        if unknown_device:
            raise ValidationError(
                f"Spare Parts device {device_index} has unknown fields: "
                + ", ".join(unknown_device)
            )
        for key in ("device", "model"):
            field = device.get(key)
            if field is not None and isinstance(field, (dict, list, tuple, set)):
                raise ValidationError(f"Spare Parts device {device_index} {key} must be text or blank")
            if field is not None and len(str(field)) > 10_000:
                raise ValidationError(f"Spare Parts device {device_index} {key} is too long")
        parts = device.get("parts", [])
        if not isinstance(parts, list):
            raise ValidationError(f"Spare Parts device {device_index} parts must be a list")
        total_parts += len(parts)
        if total_parts > 2_000:
            raise ValidationError("A ticket cannot contain more than 2,000 damaged parts")
        for part_index, part in enumerate(parts, start=1):
            if not isinstance(part, dict):
                raise ValidationError(
                    f"Spare Parts device {device_index}, part {part_index} must be an object"
                )
            unknown_part = sorted(set(part) - part_keys)
            if unknown_part:
                raise ValidationError(
                    f"Spare Parts device {device_index}, part {part_index} has unknown fields: "
                    + ", ".join(unknown_part)
                )
            for key, field in part.items():
                if field is not None and isinstance(field, (dict, list, tuple, set)):
                    raise ValidationError(
                        f"Spare Parts device {device_index}, part {part_index} {key} must be text or blank"
                    )
                if field is not None and len(str(field)) > 10_000:
                    raise ValidationError(
                        f"Spare Parts device {device_index}, part {part_index} {key} is too long"
                    )
    return normalize_spare_parts(value)


def _backup_directory(workbook_directory: Path) -> Path:
    target = workbook_directory / "Zeus Backups" / "Web edits"
    target.mkdir(parents=True, exist_ok=True)
    return target


def _create_edit_backup(path: Path, *, retention: int) -> Path:
    target = _backup_directory(path.parent)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S-%f")
    backup = target / f"Pendings-before-web-edit-{stamp}.xlsx"
    shutil.copy2(path, backup)
    backups = sorted(
        target.glob("Pendings-before-web-edit-*.xlsx"),
        key=lambda item: item.stat().st_mtime,
        reverse=True,
    )
    for stale in backups[max(1, retention) :]:
        stale.unlink(missing_ok=True)
    return backup


def _restore_workbook(backup: Path, destination: Path) -> None:
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=".zeus-web-edit-rollback-",
        suffix=".xlsx",
        dir=destination.parent,
    )
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        shutil.copy2(backup, temporary)
        os.replace(temporary, destination)
    finally:
        temporary.unlink(missing_ok=True)


def edit_ticket_through_pendings(
    store: ZeusStore,
    ticket_id: str,
    changes: dict[str, Any],
    *,
    expected_revision: str,
) -> dict[str, Any]:
    """Edit Pendings first, then import its validated result into Markdown.

    Pendings.xlsx remains the first source of truth.  A browser save is
    therefore implemented as a safe workbook edit rather than a competing
    direct write to the Markdown record.
    """

    if not expected_revision:
        raise ValidationError("A ticket revision is required for safe editing")
    ticket = store.read_ticket(ticket_id)
    actual_revision = ticket_revision(ticket)
    if actual_revision != expected_revision:
        raise TicketRevisionConflictError(
            "This ticket changed after it was opened. Reload it before saving.",
            details={"expectedRevision": expected_revision, "actualRevision": actual_revision},
        )

    workbook_directory = store.configured_directory("workbook_directory")
    if workbook_directory is None:
        raise FeatureUnavailableError(
            "Configure the workbook folder before editing ticket fields in Zeus"
        )
    path = workbook_directory / "Pendings.xlsx"
    recreated: dict[str, Any] | None = None
    if not path.is_file():
        try:
            recreated = recreate_pendings_from_database(
                store,
                workbook_directory,
            )
            if recreated.get("created"):
                recreated["import"] = import_pendings(store, path)
        except (WorkbookPublicationError, WorkbookValidationError, OSError) as exc:
            raise FeatureUnavailableError(
                f"Pendings.xlsx is missing and Zeus could not recreate it safely: {exc}"
            ) from exc

    state = store.state()
    last_import_sha = state.get("pendings_state", {}).get("last_import_sha256")
    current_sha = sha256_file(path)
    if not last_import_sha or current_sha != last_import_sha:
        raise PendingsConflictError(
            "Pendings.xlsx changed outside Zeus. Query data first so those Excel edits are imported.",
            details={"lastImportedSha256": last_import_sha, "currentSha256": current_sha},
        )

    managed = read_pendings(path)
    validate_pendings_against_state(
        managed,
        database_ids=set(store.iter_ticket_ids()),
        snapshot=state.get("publication_state", {}).get("pendings_snapshot"),
    )
    if ticket_id not in managed.records:
        raise PendingsConflictError(
            f"Pendings.xlsx no longer contains SR {ticket_id}. Query data before editing."
        )

    prepared = _normalize_local_changes(changes)
    existing_local = normalize_local(ticket.get("local"))
    proposed_local = deepcopy(existing_local)
    proposed_fields = proposed_local["fields"]
    for key, value in prepared.items():
        if key == SPARE_PARTS_CHANGE_KEY:
            proposed_local["spare_parts"] = value
        else:
            proposed_fields[key] = value
    proposed_local = normalize_local(proposed_local)

    changed_fields = [
        key
        for key, value in prepared.items()
        if (
            existing_local.get("spare_parts") != value
            if key == SPARE_PARTS_CHANGE_KEY
            else _comparable_cell(existing_local["fields"].get(key))
            != _comparable_cell(value)
        )
    ]
    if (
        SPARE_PARTS_CHANGE_KEY in changed_fields
        and existing_local["fields"].get("Spare")
        != proposed_local["fields"].get("Spare")
    ):
        changed_fields.append("Spare")
    if not changed_fields:
        return {
            "changed": False,
            "ticket": ticket,
            "revision": actual_revision,
            "changedFields": [],
            "backup": None,
            "pendingsRecreated": recreated,
        }

    workbook = load_workbook(path, data_only=False, read_only=False, keep_links=True)
    descriptor = -1
    temporary: Path | None = None
    backup: Path | None = None
    try:
        worksheet = workbook.worksheets[0]
        mapping = {
            str(cell.value).strip(): index
            for index, cell in enumerate(worksheet[1], start=1)
            if cell.value not in (None, "")
        }
        row_number = int(managed.records[ticket_id]["row_number"])
        fields_to_write = {
            field: proposed_local["fields"].get(field)
            for field in prepared
            if field != SPARE_PARTS_CHANGE_KEY
        }
        # Compatibility cells are generated from the normalized hierarchy and
        # remain useful in exports, but they are never an independent editor.
        fields_to_write.update(
            {
                field: proposed_local["fields"].get(field)
                for field in [*LEGACY_SPARE_COLUMNS, "Spare"]
            }
        )
        for field, value in fields_to_write.items():
            column_number = mapping.get(field)
            if column_number is None:
                raise WorkbookValidationError(f"Pendings.xlsx is missing the {field} column")
            cell = worksheet.cell(row=row_number, column=column_number)
            cell.value = value
            if field == "Planned Date" and isinstance(value, date):
                cell.number_format = "yyyy-mm-dd"

        sheet_tickets = [
            {
                "ticket_id": managed_ticket_id,
                "local": proposed_local
                if managed_ticket_id == ticket_id
                else record["local"],
            }
            for managed_ticket_id, record in managed.records.items()
        ]
        write_spare_parts_sheet(workbook, sheet_tickets)

        descriptor, temporary_name = tempfile.mkstemp(
            prefix=".zeus-web-edit-",
            suffix=".xlsx",
            dir=path.parent,
        )
        os.close(descriptor)
        descriptor = -1
        temporary = Path(temporary_name)
        workbook.save(temporary)
        candidate = read_pendings(temporary)
        validate_pendings_against_state(
            candidate,
            database_ids=set(store.iter_ticket_ids()),
            snapshot=state.get("publication_state", {}).get("pendings_snapshot"),
        )
        retention = int(store.config.get("excel", {}).get("backup_retention_count", 10))
        backup = _create_edit_backup(path, retention=retention)
        try:
            os.replace(temporary, path)
            temporary = None
        except PermissionError as exc:
            raise FeatureUnavailableError(
                "Pendings.xlsx is locked. Save and close it in Excel, then try again."
            ) from exc

        try:
            import_summary = import_pendings(store, path)
        except Exception:
            _restore_workbook(backup, path)
            raise
    finally:
        workbook.close()
        if descriptor >= 0:
            os.close(descriptor)
        if temporary is not None:
            temporary.unlink(missing_ok=True)

    updated = store.read_ticket(ticket_id)
    store.append_audit(
        "web-local-edit",
        {
            "ticket_id": ticket_id,
            "changed_fields": changed_fields,
            "pendings_backup": backup.name if backup else None,
        },
    )
    return {
        "changed": True,
        "ticket": updated,
        "revision": ticket_revision(updated),
        "changedFields": changed_fields,
        "backup": str(backup) if backup else None,
        "import": import_summary,
        "pendingsRecreated": recreated,
    }


def _comparable_cell(value: Any) -> Any:
    if hasattr(value, "isoformat"):
        try:
            return value.isoformat()
        except (TypeError, ValueError):
            pass
    if value is None:
        return None
    return str(value).strip()
