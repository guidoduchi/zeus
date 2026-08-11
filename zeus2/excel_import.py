from __future__ import annotations

import re
import warnings
from copy import deepcopy
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Color

from .tickets import (
    CLOSED_SCHEMA_DRIFT_COLUMNS,
    LOCAL_COLUMNS,
    PENDING_COLUMNS,
    SPARE_PART_COLUMNS,
    SPARE_PARTS_EXPORT_MARKER,
    SPARE_PARTS_SHEET,
    REQUIRED_UPSTREAM_COLUMNS,
    UPSTREAM_COLUMNS,
    empty_local,
    normalize_local,
    normalize_done,
)
from .utils import normalize_ticket_id, parse_date, sha256_file, to_iso


class WorkbookValidationError(ValueError):
    """Raised when a workbook is not safe to import or publish."""


@dataclass
class ManagedWorkbook:
    path: Path
    sheet_name: str
    header_order: list[str]
    records: dict[str, dict[str, Any]]
    sha256: str
    done_corrections: list[dict[str, Any]]


def _load_workbook(path: Path, *, read_only: bool = False) -> Any:
    try:
        with warnings.catch_warnings():
            warnings.filterwarnings(
                "ignore",
                message="Workbook contains no default style, apply openpyxl's default",
                category=UserWarning,
            )
            return load_workbook(
                path, data_only=False, read_only=read_only, keep_links=True
            )
    except PermissionError as exc:
        raise WorkbookValidationError(
            f"{path.name} cannot be read. Close it in Excel and try again."
        ) from exc
    except Exception as exc:
        raise WorkbookValidationError(
            f"{path.name} is unreadable or is not a valid Excel workbook: {exc}"
        ) from exc


def _clean_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return to_iso(value)
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        stripped = value.strip()
        return stripped if stripped else None
    return value


def _comparable(value: Any, *, column: str | None = None) -> Any:
    cleaned = _clean_value(value)
    if column in {"Report Date", "ResolveBy", "Resolve By Suspend"}:
        from .utils import parse_datetime

        parsed = parse_datetime(cleaned)
        return to_iso(parsed) if parsed is not None else cleaned
    if isinstance(cleaned, float) and cleaned.is_integer():
        return int(cleaned)
    return cleaned


def _color_to_dict(color: Color | None) -> dict[str, Any] | None:
    if color is None or color.type is None:
        return None
    value = getattr(color, color.type, None)
    if value is None:
        return None
    return {"type": color.type, "value": value, "tint": color.tint}


def _style_snapshot(cell: Cell) -> dict[str, Any] | None:
    if not cell.has_style:
        return None
    return {
        "fill": {
            "fill_type": cell.fill.fill_type,
            "fg_color": _color_to_dict(cell.fill.fgColor),
            "bg_color": _color_to_dict(cell.fill.bgColor),
        },
        "font": {
            "name": cell.font.name,
            "size": cell.font.sz,
            "bold": bool(cell.font.bold),
            "italic": bool(cell.font.italic),
            "underline": cell.font.underline,
            "strike": bool(cell.font.strike),
            "color": _color_to_dict(cell.font.color),
            "vertAlign": cell.font.vertAlign,
        },
    }


def _trimmed_headers(worksheet: Any) -> tuple[list[str], dict[str, int]]:
    raw = [cell.value for cell in worksheet[1]]
    original_width = len(raw)
    while raw and (raw[-1] is None or not str(raw[-1]).strip()):
        raw.pop()
    for column in range(len(raw) + 1, original_width + 1):
        for row in range(2, worksheet.max_row + 1):
            value = worksheet.cell(row=row, column=column).value
            if value is not None and str(value).strip():
                raise WorkbookValidationError(
                    f"{worksheet.title} column {column} contains data but has no header"
                )
    headers = [str(value).strip() if value is not None else "" for value in raw]
    if any(not header for header in headers):
        raise WorkbookValidationError(
            f"{worksheet.title} has an empty header between managed columns"
        )
    mapping: dict[str, int] = {}
    duplicates: list[str] = []
    for index, header in enumerate(headers, start=1):
        if header in mapping:
            duplicates.append(header)
        mapping[header] = index
    if duplicates:
        raise WorkbookValidationError(
            f"Duplicate headers: {', '.join(sorted(set(duplicates)))}"
        )
    return headers, mapping


def _validate_headers(
    worksheet: Any, *, kind: str
) -> tuple[list[str], dict[str, int]]:
    headers, mapping = _trimmed_headers(worksheet)
    if kind == "advanced_search":
        required = set(REQUIRED_UPSTREAM_COLUMNS)
        allowed = set(UPSTREAM_COLUMNS)
    elif kind == "pendings":
        required = set(PENDING_COLUMNS)
        allowed = required
    elif kind == "closed":
        required = set(PENDING_COLUMNS)
        allowed = required | set(CLOSED_SCHEMA_DRIFT_COLUMNS)
    else:
        raise ValueError(f"Unknown workbook kind: {kind}")
    missing = sorted(required - set(headers))
    unknown = sorted(set(headers) - allowed)
    if missing:
        raise WorkbookValidationError(
            f"{worksheet.title} is missing required columns: {', '.join(missing)}"
        )
    if unknown:
        raise WorkbookValidationError(
            "Unknown columns are not permitted in the managed ticket table: "
            + ", ".join(unknown)
        )
    return headers, mapping


def _row_is_blank(worksheet: Any, row_number: int, max_column: int) -> bool:
    for column in range(1, max_column + 1):
        value = worksheet.cell(row=row_number, column=column).value
        if value is not None and str(value).strip():
            return False
    return True


def _reject_formulas(workbook: Any, path: Path) -> None:
    found: list[str] = []
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    found.append(f"{worksheet.title}!{cell.coordinate}")
                    if len(found) >= 12:
                        break
            if len(found) >= 12:
                break
        if len(found) >= 12:
            break
    if found:
        raise WorkbookValidationError(
            f"{path.name} contains formulas, which are forbidden: {', '.join(found)}"
        )


def _cell_value(
    worksheet: Any, row_number: int, header_columns: dict[str, int], header: str
) -> Any:
    column = header_columns.get(header)
    return worksheet.cell(row=row_number, column=column).value if column else None


def _source_metadata(path: Path, *, kind: str) -> dict[str, Any]:
    metadata: dict[str, Any] = {
        "kind": kind,
        "filename": path.name,
        "sha256": sha256_file(path),
        "source_timestamp": None,
        "exported_at": None,
    }
    match = re.search(r"(20\d{12})(?!\d)", path.stem)
    if match:
        metadata["source_timestamp"] = match.group(1)
        try:
            exported = datetime.strptime(match.group(1), "%Y%m%d%H%M%S")
            metadata["exported_at"] = to_iso(exported)
        except ValueError:
            pass
    return metadata


def _local_from_row(
    worksheet: Any,
    row_number: int,
    header_columns: dict[str, int],
) -> tuple[dict[str, Any], bool]:
    local = empty_local()
    fields = local["fields"]
    for column in LOCAL_COLUMNS:
        value = _clean_value(
            _cell_value(worksheet, row_number, header_columns, column)
        )
        if column == "Planned Date":
            parsed = parse_date(value)
            value = parsed.isoformat() if parsed else value
        fields[column] = value
    # ``empty_local`` is complete for new records, but at this point we still
    # need to distinguish a legacy flat workbook from one that explicitly
    # contains the normalized Spare Parts worksheet.
    local.pop("spare_parts", None)
    fields["Done?"], corrected = normalize_done(fields.get("Done?"))
    styles: dict[str, dict[str, Any]] = {}
    for header, column_number in header_columns.items():
        snapshot = _style_snapshot(
            worksheet.cell(row=row_number, column=column_number)
        )
        if snapshot is not None:
            styles[header] = snapshot
    local["presentation"] = {"cell_styles": styles}
    return local, corrected


def _positive_index(value: Any, *, label: str, row_number: int) -> int:
    cleaned = _clean_value(value)
    if isinstance(cleaned, float) and cleaned.is_integer():
        cleaned = int(cleaned)
    text = str(cleaned or "").strip()
    if not text.isdigit() or int(text) < 1:
        raise WorkbookValidationError(
            f"{SPARE_PARTS_SHEET} row {row_number}: {label} must be a positive integer"
        )
    return int(text)


def _read_spare_parts(
    workbook: Any,
    *,
    ticket_ids: set[str],
) -> dict[str, list[dict[str, Any]]] | None:
    """Read the optional normalized device/part table.

    A managed normalized sheet contains at least one row per SR.  Tickets with
    no spare-parts data use a blank sentinel row, making deletion explicit and
    preventing an accidentally omitted row from erasing legacy information.
    """

    if SPARE_PARTS_SHEET not in workbook.sheetnames:
        primary = workbook.worksheets[0]
        generated_headers = [
            str(cell.value or "").strip()
            for cell in primary[1]
            if cell.comment
            and SPARE_PARTS_EXPORT_MARKER in str(cell.comment.text or "")
        ]
        if generated_headers:
            raise WorkbookValidationError(
                f"{SPARE_PARTS_SHEET} is missing from a normalized Zeus workbook. "
                "Restore the worksheet from a backup before importing it."
            )
        return None
    worksheet = workbook[SPARE_PARTS_SHEET]
    headers, mapping = _trimmed_headers(worksheet)
    missing = sorted(set(SPARE_PART_COLUMNS) - set(headers))
    unknown = sorted(set(headers) - set(SPARE_PART_COLUMNS))
    if missing:
        raise WorkbookValidationError(
            f"{SPARE_PARTS_SHEET} is missing required columns: {', '.join(missing)}"
        )
    if unknown:
        raise WorkbookValidationError(
            f"{SPARE_PARTS_SHEET} contains unknown columns: {', '.join(unknown)}"
        )

    grouped: dict[str, dict[int, dict[str, Any]]] = {}
    seen_tickets: set[str] = set()
    seen_parts: set[tuple[str, int, int]] = set()
    for row_number in range(2, worksheet.max_row + 1):
        if _row_is_blank(worksheet, row_number, worksheet.max_column):
            continue
        try:
            ticket_id = normalize_ticket_id(
                _cell_value(worksheet, row_number, mapping, "SRNo")
            )
        except ValueError as exc:
            raise WorkbookValidationError(
                f"{SPARE_PARTS_SHEET} row {row_number}: {exc}"
            ) from exc
        if ticket_id not in ticket_ids:
            raise WorkbookValidationError(
                f"{SPARE_PARTS_SHEET} row {row_number} references SR {ticket_id}, "
                "which is absent from the primary sheet"
            )
        seen_tickets.add(ticket_id)

        device_value = _clean_value(
            _cell_value(worksheet, row_number, mapping, "Device #")
        )
        data_columns = [
            "Device",
            "Model",
            "Part #",
            "Slot",
            "Part",
            "BOM",
            "Faulty SN",
            "New SN",
        ]
        values = {
            column: _clean_value(_cell_value(worksheet, row_number, mapping, column))
            for column in data_columns
        }
        if device_value in (None, ""):
            if any(value not in (None, "") for value in values.values()):
                raise WorkbookValidationError(
                    f"{SPARE_PARTS_SHEET} row {row_number}: Device # is required "
                    "when device or part data is present"
                )
            continue

        device_number = _positive_index(
            device_value, label="Device #", row_number=row_number
        )
        devices = grouped.setdefault(ticket_id, {})
        device = devices.setdefault(
            device_number,
            {"device": values["Device"], "model": values["Model"], "parts": []},
        )
        for source, target in (("Device", "device"), ("Model", "model")):
            candidate = values[source]
            if candidate is None:
                continue
            if device[target] not in (None, candidate):
                raise WorkbookValidationError(
                    f"{SPARE_PARTS_SHEET} row {row_number}: Device # {device_number} "
                    f"has conflicting {source} values for SR {ticket_id}"
                )
            device[target] = candidate

        part_data = {
            "slot": values["Slot"],
            "part": values["Part"],
            "bom": values["BOM"],
            "faulty_sn": values["Faulty SN"],
            "new_sn": values["New SN"],
        }
        part_value = values["Part #"]
        if part_value in (None, ""):
            if any(value is not None for value in part_data.values()):
                raise WorkbookValidationError(
                    f"{SPARE_PARTS_SHEET} row {row_number}: Part # is required "
                    "when part data is present"
                )
            continue
        part_number = _positive_index(
            part_value, label="Part #", row_number=row_number
        )
        identity = (ticket_id, device_number, part_number)
        if identity in seen_parts:
            raise WorkbookValidationError(
                f"{SPARE_PARTS_SHEET} contains duplicate SR/Device #/Part #: "
                f"{ticket_id}/{device_number}/{part_number}"
            )
        seen_parts.add(identity)
        device["parts"].append((part_number, part_data))

    missing_tickets = sorted(ticket_ids - seen_tickets, reverse=True)
    if missing_tickets:
        raise WorkbookValidationError(
            f"{SPARE_PARTS_SHEET} is missing an explicit row for SR: "
            + ", ".join(missing_tickets[:20])
        )

    result: dict[str, list[dict[str, Any]]] = {}
    for ticket_id in ticket_ids:
        result[ticket_id] = []
        for _, device in sorted(grouped.get(ticket_id, {}).items()):
            result[ticket_id].append(
                {
                    "device": device["device"],
                    "model": device["model"],
                    "parts": [part for _, part in sorted(device["parts"])],
                }
            )
    return result


def read_pendings(path: Path) -> ManagedWorkbook:
    path = path.expanduser().resolve()
    if not path.is_file():
        raise WorkbookValidationError(f"Pendings.xlsx was not found: {path}")
    workbook = _load_workbook(path)
    try:
        _reject_formulas(workbook, path)
        worksheet = workbook.worksheets[0]
        headers, mapping = _validate_headers(worksheet, kind="pendings")
        records: dict[str, dict[str, Any]] = {}
        duplicates: list[str] = []
        corrections: list[dict[str, Any]] = []
        for row_number in range(2, worksheet.max_row + 1):
            if _row_is_blank(worksheet, row_number, max(len(headers), worksheet.max_column)):
                continue
            raw_id = _cell_value(worksheet, row_number, mapping, "SRNo")
            try:
                ticket_id = normalize_ticket_id(raw_id)
            except ValueError as exc:
                raise WorkbookValidationError(
                    f"Pendings.xlsx row {row_number}: {exc}. Check or delete the "
                    "workbook yourself; Zeus will not replace it automatically."
                ) from exc
            if ticket_id in records:
                duplicates.append(ticket_id)
                continue
            local, corrected = _local_from_row(worksheet, row_number, mapping)
            if corrected:
                corrections.append(
                    {
                        "ticket_id": ticket_id,
                        "cell": worksheet.cell(row_number, mapping["Done?"]).coordinate,
                        "original": _cell_value(worksheet, row_number, mapping, "Done?"),
                        "normalized": "N",
                    }
                )
            protected = {
                column: _clean_value(_cell_value(worksheet, row_number, mapping, column))
                for column in UPSTREAM_COLUMNS
            }
            protected["SRNo"] = ticket_id
            id_cell = worksheet.cell(row=row_number, column=mapping["SRNo"])
            records[ticket_id] = {
                "ticket_id": ticket_id,
                "upstream_fields": protected,
                "local": local,
                "sr_url": id_cell.hyperlink.target if id_cell.hyperlink else None,
                "row_number": row_number,
            }
        if duplicates:
            raise WorkbookValidationError(
                "Pendings.xlsx contains duplicate SRNo values: "
                + ", ".join(sorted(set(duplicates)))
            )
        spare_parts = _read_spare_parts(workbook, ticket_ids=set(records))
        for ticket_id, record in records.items():
            if spare_parts is not None:
                record["local"]["spare_parts"] = spare_parts[ticket_id]
            record["local"] = normalize_local(record["local"])
        return ManagedWorkbook(
            path=path,
            sheet_name=worksheet.title,
            header_order=headers,
            records=records,
            sha256=sha256_file(path),
            done_corrections=corrections,
        )
    finally:
        workbook.close()


def validate_pendings_against_state(
    workbook: ManagedWorkbook,
    *,
    database_ids: set[str],
    snapshot: dict[str, Any] | None,
) -> None:
    workbook_ids = set(workbook.records)
    if database_ids:
        unknown = sorted(workbook_ids - database_ids, reverse=True)
        if unknown:
            raise WorkbookValidationError(
                "Pendings.xlsx contains rows Zeus does not recognize: "
                + ", ".join(unknown)
            )
    if not snapshot:
        return
    exported_ids = {str(value) for value in snapshot.get("ticket_ids", [])}
    missing = sorted(exported_ids - workbook_ids, reverse=True)
    if missing:
        raise WorkbookValidationError(
            "Pendings.xlsx is missing previously exported ticket rows: "
            + ", ".join(missing)
        )
    baseline = snapshot.get("protected_values", {})
    changed: list[str] = []
    for ticket_id in sorted(exported_ids & workbook_ids, reverse=True):
        expected = baseline.get(ticket_id, {})
        actual = workbook.records[ticket_id]["upstream_fields"]
        for column in UPSTREAM_COLUMNS:
            if _comparable(actual.get(column), column=column) != _comparable(
                expected.get(column), column=column
            ):
                changed.append(f"{ticket_id}:{column}")
                if len(changed) >= 20:
                    break
        if len(changed) >= 20:
            break
    if changed:
        raise WorkbookValidationError(
            "Protected Advanced Search fields were changed in Pendings.xlsx: "
            + ", ".join(changed)
        )


def read_closed(path: Path) -> ManagedWorkbook:
    path = path.expanduser().resolve()
    if not path.is_file():
        raise WorkbookValidationError(f"Closed.xlsx was not found: {path}")
    workbook = _load_workbook(path)
    try:
        _reject_formulas(workbook, path)
        worksheet = workbook.worksheets[0]
        headers, mapping = _validate_headers(worksheet, kind="closed")
        records: dict[str, dict[str, Any]] = {}
        duplicates: list[str] = []
        corrections: list[dict[str, Any]] = []
        for row_number in range(2, worksheet.max_row + 1):
            if _row_is_blank(worksheet, row_number, max(len(headers), worksheet.max_column)):
                continue
            raw_id = _cell_value(worksheet, row_number, mapping, "SRNo")
            try:
                ticket_id = normalize_ticket_id(raw_id)
            except ValueError as exc:
                raise WorkbookValidationError(
                    f"Closed.xlsx row {row_number}: {exc}"
                ) from exc
            if ticket_id in records:
                duplicates.append(ticket_id)
                continue
            local, _ = _local_from_row(worksheet, row_number, mapping)
            protected = {
                column: _clean_value(
                    _cell_value(worksheet, row_number, mapping, column)
                )
                for column in UPSTREAM_COLUMNS
            }
            protected["SRNo"] = ticket_id
            id_cell = worksheet.cell(row=row_number, column=mapping["SRNo"])
            records[ticket_id] = {
                "ticket_id": ticket_id,
                "upstream_fields": protected,
                "row_number": row_number,
                "local": local,
                "sr_url": id_cell.hyperlink.target if id_cell.hyperlink else None,
            }
        if duplicates:
            raise WorkbookValidationError(
                "Closed.xlsx contains duplicate SRNo values: "
                + ", ".join(sorted(set(duplicates)))
            )
        spare_parts = _read_spare_parts(workbook, ticket_ids=set(records))
        for ticket_id, record in records.items():
            if spare_parts is not None:
                record["local"]["spare_parts"] = spare_parts[ticket_id]
            record["local"] = normalize_local(record["local"])
        return ManagedWorkbook(
            path=path,
            sheet_name=worksheet.title,
            header_order=headers,
            records=records,
            sha256=sha256_file(path),
            done_corrections=corrections,
        )
    finally:
        workbook.close()


def validate_closed_against_index(
    workbook: ManagedWorkbook, closed_index: dict[str, Any]
) -> None:
    indexed = {str(value) for value in closed_index.get("ticket_ids", [])}
    actual = set(workbook.records)
    missing = sorted(indexed - actual, reverse=True)
    if missing:
        raise WorkbookValidationError(
            "Closed.xlsx is missing or has altered historical SR numbers: "
            + ", ".join(missing)
        )
    initialized = closed_index.get("validated_at") is not None
    unexpected = sorted(actual - indexed, reverse=True) if initialized else []
    if unexpected:
        raise WorkbookValidationError(
            "Closed.xlsx contains historical SR rows not created by Zeus: "
            + ", ".join(unexpected)
        )


def read_advanced_search(
    path: Path, *, config: dict[str, Any] | None = None
) -> tuple[dict[str, dict[str, Any]], dict[str, Any]]:
    path = path.expanduser().resolve()
    if not path.is_file():
        raise WorkbookValidationError(f"Advanced Search workbook not found: {path}")
    workbook = _load_workbook(path)
    try:
        _reject_formulas(workbook, path)
        worksheet = (
            workbook["Service Request"]
            if "Service Request" in workbook.sheetnames
            else workbook.worksheets[0]
        )
        _, mapping = _validate_headers(worksheet, kind="advanced_search")
        source = _source_metadata(path, kind="advanced_search")
        records: dict[str, dict[str, Any]] = {}
        duplicates: list[str] = []
        for row_number in range(2, worksheet.max_row + 1):
            if _row_is_blank(worksheet, row_number, worksheet.max_column):
                continue
            raw_id = _cell_value(worksheet, row_number, mapping, "SRNo")
            try:
                ticket_id = normalize_ticket_id(raw_id)
            except ValueError as exc:
                raise WorkbookValidationError(
                    f"{path.name} row {row_number}: {exc}"
                ) from exc
            if ticket_id in records:
                duplicates.append(ticket_id)
                continue
            fields = {
                column: _clean_value(
                    _cell_value(worksheet, row_number, mapping, column)
                )
                for column in UPSTREAM_COLUMNS
            }
            fields["SRNo"] = ticket_id
            id_cell = worksheet.cell(row=row_number, column=mapping["SRNo"])
            records[ticket_id] = {
                "ticket_id": ticket_id,
                "upstream_fields": fields,
                "source": deepcopy(source),
                "sr_url": id_cell.hyperlink.target if id_cell.hyperlink else None,
                "row_number": row_number,
            }
        if duplicates:
            raise WorkbookValidationError(
                f"{path.name} contains duplicate SRNo values: "
                + ", ".join(sorted(set(duplicates)))
            )
        if not records:
            raise WorkbookValidationError(f"{path.name} contains no ticket rows")
        return records, source
    finally:
        workbook.close()


def find_latest_advanced_search(
    path_or_directory: Path, config: dict[str, Any]
) -> Path:
    candidate = path_or_directory.expanduser().resolve()
    if candidate.is_file():
        if candidate.name.startswith("~$"):
            raise WorkbookValidationError(
                "Temporary Excel lock files cannot be synchronized"
            )
        return candidate
    if not candidate.is_dir():
        raise WorkbookValidationError(f"Source path not found: {candidate}")
    pattern = config.get("advanced_search", {}).get("glob", "Advanced Search*.xlsx")
    files = [
        path
        for path in candidate.glob(pattern)
        if path.is_file() and not path.name.startswith("~$")
    ]
    if not files:
        raise WorkbookValidationError(
            f"No files matching {pattern!r} were found in {candidate}"
        )

    def sort_key(path: Path) -> tuple[str, float, str]:
        match = re.search(r"(20\d{12})(?!\d)", path.stem)
        timestamp = match.group(1) if match else ""
        return timestamp, path.stat().st_mtime, path.name.lower()

    return max(files, key=sort_key)


# Compatibility name retained for scripts from the first 2.x preview.
def read_legacy_workbook(
    path: Path, *, kind: str, config: dict[str, Any]
) -> tuple[dict[str, dict[str, Any]], dict[str, Any]]:
    managed = read_pendings(path) if "pending" in kind else read_closed(path)
    return managed.records, _source_metadata(path, kind=kind)
