from __future__ import annotations

import re
import shutil
import tempfile
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from docx.document import Document as DocumentObject
from docx.table import _Cell, Table
from docx.text.paragraph import Paragraph

from .aging import aging_for_ticket
from .store import ZeusStore
from .tickets import workflow_code
from .utils import iso_now, normalize_ticket_id, sha256_file


PLACEHOLDER_PATTERN = re.compile(r"{{\s*([A-Za-z][A-Za-z0-9_.-]*)\s*}}")


class MopGenerationError(RuntimeError):
    """Raised when a MOP template cannot be safely generated."""


def _snake(value: str) -> str:
    value = re.sub(r"([a-z0-9])([A-Z])", r"\1_\2", value)
    return re.sub(r"[^a-z0-9]+", "_", value.lower()).strip("_")


def _format_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Y" if value else "N"
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (list, tuple, set)):
        return "\n".join(_format_value(item) for item in value if item is not None)
    return str(value)


def build_placeholder_values(store: ZeusStore, ticket_id: str) -> dict[str, str]:
    normalized_id = normalize_ticket_id(ticket_id)
    ticket = store.read_ticket(normalized_id)
    notes = store.read_notes(normalized_id)
    fields = ticket.get("upstream", {}).get("fields", {})
    local = ticket.get("local", {})
    fields_local = local.get("fields", {})
    aging = aging_for_ticket(ticket, store.config)

    values: dict[str, Any] = {
        "ticket_id": normalized_id,
        "srno": normalized_id,
        "done": workflow_code(ticket),
        "planned_mw_date": fields_local.get("Planned Date"),
        "site": fields_local.get("Site"),
        "cloud": fields_local.get("Cloud"),
        "model": fields_local.get("Model"),
        "models": fields_local.get("Model"),
        "device": fields_local.get("Device"),
        "devices": fields_local.get("Device"),
        "slot": fields_local.get("Slot"),
        "slots": fields_local.get("Slot"),
        "part": fields_local.get("Part"),
        "parts": fields_local.get("Part"),
        "bom": fields_local.get("BOM"),
        "boms": fields_local.get("BOM"),
        "old_sn": fields_local.get("Old SN"),
        "new_sn": fields_local.get("New SN"),
        "spare": fields_local.get("Spare"),
        "related_sr": fields_local.get("RelatedSR"),
        "notes": notes.strip(),
        "age_days": aging.ticket_age_days,
        "planned_days": aging.planned_days,
        "resolve_days": aging.resolve_days,
        "resolve_state": aging.resolve_state,
    }
    for key, value in fields.items():
        values[_snake(key)] = value
    for key, value in local.get("mop_fields", {}).items():
        values[key] = value
        values[_snake(key)] = value
    return {key: _format_value(value) for key, value in values.items()}


def _iter_table_paragraphs(table: Table) -> Iterable[Paragraph]:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                yield paragraph
            for nested in cell.tables:
                yield from _iter_table_paragraphs(nested)


def _iter_document_paragraphs(document: DocumentObject) -> Iterable[Paragraph]:
    yield from document.paragraphs
    for table in document.tables:
        yield from _iter_table_paragraphs(table)
    for section in document.sections:
        yield from section.header.paragraphs
        for table in section.header.tables:
            yield from _iter_table_paragraphs(table)
        yield from section.footer.paragraphs
        for table in section.footer.tables:
            yield from _iter_table_paragraphs(table)


def _replace_in_paragraph(
    paragraph: Paragraph,
    values: dict[str, str],
    missing: set[str],
) -> int:
    runs = paragraph.runs
    if not runs:
        return 0
    full_text = "".join(run.text for run in runs)
    matches = list(PLACEHOLDER_PATTERN.finditer(full_text))
    if not matches:
        return 0
    replaced = 0
    run_starts: list[int] = []
    cursor = 0
    for run in runs:
        run_starts.append(cursor)
        cursor += len(run.text)

    def locate(position: int, *, end: bool = False) -> tuple[int, int]:
        if position == len(full_text):
            return len(runs) - 1, len(runs[-1].text)
        for index, start in enumerate(run_starts):
            length = len(runs[index].text)
            if start <= position < start + length:
                return index, position - start
            if end and position == start + length:
                return index, length
        return len(runs) - 1, len(runs[-1].text)

    for match in reversed(matches):
        key = match.group(1)
        if key not in values:
            missing.add(key)
            continue
        start_run, start_offset = locate(match.start())
        end_run, end_offset = locate(match.end(), end=True)
        prefix = runs[start_run].text[:start_offset]
        suffix = runs[end_run].text[end_offset:]
        runs[start_run].text = prefix + values[key] + suffix
        for index in range(start_run + 1, end_run + 1):
            runs[index].text = ""
        replaced += 1
    return replaced


def render_template(
    template_path: Path,
    output_path: Path,
    values: dict[str, str],
    *,
    allow_missing: bool = False,
) -> dict[str, Any]:
    template_path = template_path.expanduser().resolve()
    if not template_path.is_file() or template_path.suffix.lower() != ".docx":
        raise MopGenerationError(f"MOP template must be an existing .docx file: {template_path}")
    document = Document(template_path)
    missing: set[str] = set()
    replacements = 0
    for paragraph in _iter_document_paragraphs(document):
        replacements += _replace_in_paragraph(paragraph, values, missing)
    if missing and not allow_missing:
        raise MopGenerationError(
            "Template placeholders have no ticket value: " + ", ".join(sorted(missing))
        )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return {
        "replacements": replacements,
        "missing_placeholders": sorted(missing),
    }


def generate_mop(
    store: ZeusStore,
    ticket_id: str,
    template_path: Path,
    *,
    allow_missing: bool = False,
    copy_to: Path | None = None,
) -> dict[str, Any]:
    normalized_id = normalize_ticket_id(ticket_id)
    ticket = store.read_ticket(normalized_id)
    values = build_placeholder_values(store, normalized_id)
    current_version = int(ticket.get("mop", {}).get("versions") or 0)
    next_version = current_version + 1
    filename = f"MOP-{normalized_id}-v{next_version:03d}.docx"
    temporary_root = Path(tempfile.mkdtemp(prefix=".mop-", dir=store.root)).resolve()
    temporary_output = temporary_root / filename
    try:
        render_summary = render_template(
            template_path,
            temporary_output,
            values,
            allow_missing=allow_missing,
        )
        summary: dict[str, Any] = {
            "ticket_id": normalized_id,
            "template": template_path.name,
            "template_sha256": sha256_file(template_path),
            "version": next_version,
            "filename": filename,
            **render_summary,
        }
        with store.transaction("mop-generate", summary) as staging:
            staged_ticket = store.read_ticket(normalized_id, staging)
            destination = store.ticket_dir(normalized_id, staging) / "mops" / filename
            destination.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(temporary_output, destination)
            staged_ticket["mop"] = {
                "latest": f"mops/{filename}",
                "versions": next_version,
                "generated_at": iso_now(),
                "template": template_path.name,
                "template_sha256": summary["template_sha256"],
            }
            store.write_ticket_bundle(staging, staged_ticket)
        final_path = store.ticket_dir(normalized_id) / "mops" / filename
        if copy_to is not None:
            copy_destination = copy_to.expanduser().resolve()
            if copy_destination.is_dir() or copy_destination.suffix.lower() != ".docx":
                copy_destination.mkdir(parents=True, exist_ok=True)
                copy_destination = copy_destination / filename
            else:
                copy_destination.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(final_path, copy_destination)
            summary["copied_to"] = str(copy_destination)
        summary["path"] = str(final_path)
        return summary
    finally:
        shutil.rmtree(temporary_root, ignore_errors=True)
