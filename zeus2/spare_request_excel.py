from __future__ import annotations

import os
import re
import shutil
import uuid
from copy import copy, deepcopy
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter, range_boundaries

from .spare_requests import (
    ECUADOR_TIMEZONE,
    aging_color,
    dispatch_age_days,
    item_status,
    normalize_rma,
    normalize_spare_sr,
    request_description,
    request_history,
)
from .tickets import PENDING_COLUMNS
from .utils import iso_now, parse_date, parse_datetime, sha256_file


REQUEST_SHEET = "Spare Parts Request"
FAULTY_TAG_SHEET = "Faulty Tag Template"
RETURN_SHEET = "Fault Tag Return"
SPARE_ARCHIVE_SHEET = "Spare Requests"
SPARE_ARCHIVE_EMAIL_SHEET = "Spare Request Emails"

LIGHT_RED = "F4CCCC"
LIGHT_GREEN = "D9EAD3"
ARCHIVE_HEADERS = [
    "TT",
    "RMA",
    "RMA Aliases",
    "Email inactivity days",
    "Emails received",
    "Emails sent",
    "Spare SR",
    "Request ID",
    "Status",
    "Archived At",
    "Archive Reason",
    "Client Initials",
    "Customer",
    "Site",
    "Address",
    "Cloud",
    "Requested BOM",
    "Delivered BOM",
    "Description",
    "Part",
    "Model",
    "Device",
    "Slot",
    "Faulty SN",
    "New SN",
    "Return Condition",
    "Dispatch Date",
    "Return Export",
    "RT",
    "Notes",
    "Source",
    "Item ID",
]
LEGACY_ARCHIVE_HEADERS = [
    header for header in ARCHIVE_HEADERS if header != "RMA Aliases"
]
ARCHIVE_EMAIL_HEADERS = [
    "TT",
    "Spare SR",
    "RMA",
    "Item ID",
    "Message ID",
    "Timestamp",
    "Direction",
    "Sender",
    "Sender Address",
    "Subject",
    "Body",
]


class SpareRequestWorkbookError(RuntimeError):
    pass


def archived_rma_values(row: dict[str, Any]) -> set[str]:
    """Return every current or historical RMA reserved by an archive row."""

    candidates = [row.get("RMA")]
    candidates.extend(re.split(r"[\r\n,;]+", str(row.get("RMA Aliases") or "")))
    return {
        normalized
        for candidate in candidates
        if str(candidate or "").strip()
        for normalized in [normalize_rma(candidate, required=True)]
        if normalized
    }


def _require_template(path: Path, expected_sheet_count: int) -> tuple[Path, tuple[str, ...]]:
    resolved = path.expanduser().resolve()
    if not resolved.is_file() or resolved.suffix.lower() != ".xlsx":
        raise SpareRequestWorkbookError(f"Spare Request template was not found: {resolved}")
    workbook = load_workbook(resolved, read_only=True, data_only=False)
    try:
        sheet_names = tuple(workbook.sheetnames)
        if len(sheet_names) != expected_sheet_count:
            raise SpareRequestWorkbookError(
                f"Template must contain exactly {expected_sheet_count} worksheet(s)"
            )
    finally:
        workbook.close()
    return resolved, sheet_names


def _shifted_range(value: str, row_offset: int) -> str:
    min_col, min_row, max_col, max_row = range_boundaries(value)
    return (
        f"{get_column_letter(min_col)}{min_row + row_offset}:"
        f"{get_column_letter(max_col)}{max_row + row_offset}"
    )


def _insert_rows_preserving_merges(worksheet: Any, index: int, amount: int) -> None:
    if amount <= 0:
        return
    shifted: list[tuple[str, int]] = []
    spanning: list[str] = []
    for merged in list(worksheet.merged_cells.ranges):
        text = str(merged)
        if merged.min_row >= index:
            shifted.append((text, amount))
            worksheet.unmerge_cells(text)
        elif merged.max_row >= index:
            spanning.append(text)
            worksheet.unmerge_cells(text)
    worksheet.insert_rows(index, amount)
    for text, offset in shifted:
        worksheet.merge_cells(_shifted_range(text, offset))
    for text in spanning:
        min_col, min_row, max_col, max_row = range_boundaries(text)
        worksheet.merge_cells(
            start_row=min_row,
            start_column=min_col,
            end_row=max_row + amount,
            end_column=max_col,
        )


def _copy_two_row_block(worksheet: Any, source_row: int, target_row: int) -> None:
    source_merges = [
        str(merged)
        for merged in worksheet.merged_cells.ranges
        if source_row <= merged.min_row and merged.max_row <= source_row + 1
    ]
    target_merges = [
        str(merged)
        for merged in worksheet.merged_cells.ranges
        if target_row <= merged.min_row and merged.max_row <= target_row + 1
    ]
    for merged in target_merges:
        worksheet.unmerge_cells(merged)
    for row_offset in range(2):
        source_height = worksheet.row_dimensions[source_row + row_offset].height
        worksheet.row_dimensions[target_row + row_offset].height = source_height
        for column in range(1, worksheet.max_column + 1):
            source = worksheet.cell(source_row + row_offset, column)
            target = worksheet.cell(target_row + row_offset, column)
            target._style = copy(source._style)
            target.number_format = source.number_format
            target.protection = copy(source.protection)
            target.alignment = copy(source.alignment)
            target.value = None
    offset = target_row - source_row
    for merged in source_merges:
        shifted = _shifted_range(merged, offset)
        if shifted not in {str(value) for value in worksheet.merged_cells.ranges}:
            worksheet.merge_cells(shifted)


def _ensure_two_row_blocks(
    worksheet: Any,
    *,
    first_row: int,
    existing_blocks: int,
    required_blocks: int,
    insert_at: int,
    style_source_row: int,
) -> None:
    if required_blocks <= existing_blocks:
        return
    additional = required_blocks - existing_blocks
    _insert_rows_preserving_merges(worksheet, insert_at, additional * 2)
    for index in range(additional):
        target = first_row + (existing_blocks + index) * 2
        _copy_two_row_block(worksheet, style_source_row, target)


def _set_merged_value(worksheet: Any, coordinate: str, value: Any) -> None:
    worksheet[coordinate] = value


def _excel_date(value: Any) -> date | datetime | None:
    parsed_date = parse_date(value)
    if parsed_date is not None:
        return parsed_date
    parsed_datetime = parse_datetime(value)
    if parsed_datetime is not None and parsed_datetime.tzinfo is not None:
        parsed_datetime = parsed_datetime.astimezone(ECUADOR_TIMEZONE).replace(tzinfo=None)
    return parsed_datetime


def _request_applied_date(request: dict[str, Any]) -> date:
    value = request.get("export", {}).get("created_at") or request.get("created_at")
    parsed = parse_datetime(value)
    if parsed is None:
        return datetime.now(ECUADOR_TIMEZONE).date()
    if parsed.tzinfo is not None:
        parsed = parsed.astimezone(ECUADOR_TIMEZONE)
    return parsed.date()


def _populate_request_sheet(worksheet: Any, request: dict[str, Any]) -> None:
    profile = request.get("profile", {})
    requester = profile.get("requester", {})
    contact = profile.get("contact", {})
    _set_merged_value(
        worksheet,
        "G2",
        profile.get("customer_organization") or profile.get("customer_name"),
    )
    _set_merged_value(worksheet, "J2", requester.get("name"))
    _set_merged_value(worksheet, "G4", _request_applied_date(request))
    worksheet["G4"].number_format = "yyyy-mm-dd"
    _set_merged_value(worksheet, "J4", requester.get("email"))
    _set_merged_value(worksheet, "G6", requester.get("phone"))
    _set_merged_value(worksheet, "J6", request.get("tt"))
    _set_merged_value(worksheet, "F11", request.get("spare_sr"))
    _set_merged_value(worksheet, "B28", profile.get("site_address"))
    _set_merged_value(worksheet, "F31", contact.get("name"))
    _set_merged_value(worksheet, "F33", contact.get("phone"))
    _set_merged_value(worksheet, "F35", contact.get("email"))

    lines = list(request.get("request_lines") or [])
    _ensure_two_row_blocks(
        worksheet,
        first_row=15,
        existing_blocks=4,
        required_blocks=len(lines),
        insert_at=23,
        style_source_row=21,
    )
    for index in range(max(4, len(lines))):
        row = 15 + index * 2
        for column in (2, 3, 4, 6, 8, 10, 11):
            worksheet.cell(row, column).value = None
    for index, line in enumerate(lines):
        row = 15 + index * 2
        worksheet.cell(row, 2).value = line.get("bom")
        worksheet.cell(row, 3).value = int(line.get("amount") or 1)
        worksheet.cell(row, 4).value = request_description(line)
        serials = list(line.get("faulty_sns") or [])
        if not serials and line.get("faulty_sn"):
            serials = [
                value.strip()
                for value in re.split(r"\r?\n", str(line.get("faulty_sn")))
                if value.strip()
            ]
        serial_cell = worksheet.cell(row, 6)
        serial_cell.value = "\n".join(serials) or None
        alignment = copy(serial_cell.alignment)
        alignment.wrap_text = True
        serial_cell.alignment = alignment


def _faulty_tag_rows(request: dict[str, Any]) -> list[dict[str, Any]]:
    """Create exactly one Fault Tag row for each requested physical unit.

    Diagnostic component serials describe why that unit is being replaced;
    they do not create additional replacement units or future RMAs.
    """

    items = list(request.get("items") or [])
    if items:
        return [
            {
                "bom": item.get("requested_bom"),
                "description": item.get("requested_description"),
                "part": item.get("part"),
                "model": item.get("model"),
                "device": item.get("device"),
                "slot": item.get("slot"),
                "faulty_sn": item.get("faulty_sn"),
                "report_date": item.get("report_date"),
                "notes": item.get("notes"),
            }
            for item in items
        ]

    rows: list[dict[str, Any]] = []
    for line in request.get("request_lines") or []:
        serials = list(line.get("faulty_sns") or [])
        if not serials and line.get("faulty_sn"):
            serials = [
                value.strip()
                for value in re.split(r"\r?\n", str(line.get("faulty_sn")))
                if value.strip()
            ]
        combined_serials = "\n".join(serials) or None
        for _ in range(max(1, int(line.get("amount") or 1))):
            rows.append({**line, "faulty_sn": combined_serials})
    return rows


def _populate_faulty_tag_sheet(worksheet: Any, request: dict[str, Any]) -> None:
    profile = request.get("profile", {})
    requester = profile.get("requester", {})
    _set_merged_value(
        worksheet,
        "G2",
        profile.get("customer_organization") or profile.get("customer_name"),
    )
    _set_merged_value(worksheet, "K2", request.get("spare_sr"))
    _set_merged_value(worksheet, "G4", requester.get("name"))
    _set_merged_value(worksheet, "K4", requester.get("phone"))
    _set_merged_value(worksheet, "G6", requester.get("email"))
    _set_merged_value(worksheet, "K6", request.get("tt"))
    rows = _faulty_tag_rows(request)
    # The supplied first item row contains broken #REF! formulas. Match the
    # valid row-14 formatting, then write every item as a direct value.
    _copy_two_row_block(worksheet, 14, 12)
    _ensure_two_row_blocks(
        worksheet,
        first_row=12,
        existing_blocks=5,
        required_blocks=len(rows),
        insert_at=22,
        style_source_row=20,
    )
    for index in range(max(5, len(rows))):
        row = 12 + index * 2
        for column in (2, 4, 6, 7, 8, 9, 10, 12, 13):
            worksheet.cell(row, column).value = None
    for index, item in enumerate(rows):
        row = 12 + index * 2
        worksheet.cell(row, 2).value = item.get("bom")
        description_cell = worksheet.cell(row, 4)
        description_cell.value = request_description(item)
        description_alignment = copy(description_cell.alignment)
        description_alignment.wrap_text = True
        description_cell.alignment = description_alignment
        serial_cell = worksheet.cell(row, 6)
        serial_cell.value = item.get("faulty_sn")
        serial_alignment = copy(serial_cell.alignment)
        serial_alignment.wrap_text = True
        serial_cell.alignment = serial_alignment
        worksheet.cell(row, 7).value = profile.get("site_code")
        fault_date = _excel_date(item.get("report_date"))
        worksheet.cell(row, 8).value = fault_date
        if fault_date is not None:
            worksheet.cell(row, 8).number_format = "yyyy-mm-dd"


def _next_revision_path(path: Path) -> Path:
    if not path.exists():
        return path
    stem = re.sub(r"-r\d+$", "", path.stem, flags=re.IGNORECASE)
    revision = 2
    while True:
        candidate = path.with_name(f"{stem}-r{revision}{path.suffix}")
        if not candidate.exists():
            return candidate
        revision += 1


def _atomic_save(workbook: Any, destination: Path) -> Path:
    destination.parent.mkdir(parents=True, exist_ok=True)
    final = _next_revision_path(destination)
    temporary = final.with_name(f".{final.stem}-{uuid.uuid4().hex}.xlsx")
    try:
        workbook.save(temporary)
        os.replace(temporary, final)
    finally:
        temporary.unlink(missing_ok=True)
    return final


def _validate_initial_export(
    path: Path,
    request: dict[str, Any],
    expected_sheet_names: tuple[str, ...],
) -> None:
    workbook = load_workbook(path, read_only=False, data_only=False, keep_links=True)
    try:
        if workbook.sheetnames != list(expected_sheet_names):
            raise SpareRequestWorkbookError("Generated request workbook changed template sheets")
        application = workbook.worksheets[0]
        exported_boms = [
            str(application.cell(15 + index * 2, 2).value or "")
            for index in range(len(request.get("request_lines") or []))
        ]
        expected = [str(line.get("bom") or "") for line in request.get("request_lines") or []]
        if exported_boms != expected:
            raise SpareRequestWorkbookError("Generated request workbook failed BOM verification")
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and "#REF!" in cell.value:
                        raise SpareRequestWorkbookError(
                            f"Generated request workbook contains a broken reference at {worksheet.title}!{cell.coordinate}"
                        )
    finally:
        workbook.close()


def export_initial_request(
    template_path: Path,
    export_root: Path,
    request: dict[str, Any],
    filename: str,
) -> Path:
    template, sheet_names = _require_template(template_path, 2)
    workbook = load_workbook(template, read_only=False, data_only=False, keep_links=True)
    try:
        _populate_request_sheet(workbook.worksheets[0], request)
        _populate_faulty_tag_sheet(workbook.worksheets[1], request)
        destination = export_root.expanduser().resolve() / "Requests" / filename
        final = _atomic_save(workbook, destination)
    finally:
        workbook.close()
    try:
        _validate_initial_export(final, request, sheet_names)
    except Exception:
        final.unlink(missing_ok=True)
        raise
    return final


def _return_description(item: dict[str, Any]) -> str:
    return request_description(
        {
            "description": item.get("requested_description"),
            "part": item.get("part"),
            "model": item.get("model"),
            "device": item.get("device"),
            "slot": item.get("slot"),
        }
    )


def _return_sn(item: dict[str, Any], condition: str) -> str | None:
    return item.get("new_sn") if condition == "New" else None


def export_return_workbook(
    template_path: Path,
    export_root: Path,
    selections: list[tuple[dict[str, Any], dict[str, Any], str]],
    *,
    fault_tag_id: str | None = None,
    return_site: dict[str, Any] | None = None,
    today: date | None = None,
) -> dict[str, Any]:
    if not selections:
        raise SpareRequestWorkbookError("Choose at least one RMA to return")
    first_request = selections[0][0]
    first_profile = first_request.get("profile", {})
    supplied_site = return_site if isinstance(return_site, dict) else {}
    site = str(supplied_site.get("code") or first_profile.get("site_code") or "").strip()
    cloud = str(supplied_site.get("cloud") or first_profile.get("cloud") or "").strip()
    address = str(
        supplied_site.get("address") or first_profile.get("site_address") or ""
    ).strip()
    if not site or not cloud:
        raise SpareRequestWorkbookError("Every return item requires a site and cloud")
    if not address:
        raise SpareRequestWorkbookError("Every return requires an actual return-site address")
    source_sites = {
        (
            str(request.get("profile", {}).get("site_code") or "").strip(),
            str(request.get("profile", {}).get("cloud") or "").strip(),
        )
        for request, _, _ in selections
    }
    if len(source_sites) > 1 and not return_site:
        raise SpareRequestWorkbookError(
            "Items come from different sites. Choose the actual return site before export."
        )
    for request, item, condition in selections:
        normalize_spare_sr(request.get("spare_sr"), required=True)
        normalize_rma(item.get("rma"), required=True)
        if condition not in {"Faulty", "New"}:
            raise SpareRequestWorkbookError("Return condition must be Faulty or New")

    template, sheet_names = _require_template(template_path, 1)
    workbook = load_workbook(template, read_only=False, data_only=False, keep_links=True)
    warnings: list[str] = []
    try:
        worksheet = workbook.worksheets[0]
        contact = first_profile.get("contact", {})
        worksheet["A2"] = f"*Customer’s Name：{first_profile.get('customer_name') or ''}"
        worksheet["B3"] = address
        worksheet["B4"] = contact.get("name")
        worksheet["B5"] = contact.get("phone")
        item_count = len(selections)
        available_rows = 4
        if item_count > available_rows:
            _insert_rows_preserving_merges(worksheet, 12, item_count - available_rows)
        for index in range(max(available_rows, item_count)):
            row = 8 + index
            if index:
                for column in range(1, 10):
                    source = worksheet.cell(8, column)
                    target = worksheet.cell(row, column)
                    target._style = copy(source._style)
                    target.alignment = copy(source.alignment)
                    target.protection = copy(source.protection)
            for column in range(1, 10):
                worksheet.cell(row, column).value = None
                worksheet.cell(row, column).fill = copy(worksheet.cell(8, column).fill)
        for index, (request, item, condition) in enumerate(selections):
            row = 8 + index
            return_sn = _return_sn(item, condition)
            if not return_sn:
                warnings.append(
                    f"{item['rma']} has no return serial; confirm or enter it in the exported workbook."
                )
            worksheet.cell(row, 1).value = item.get("delivered_bom") or item.get("requested_bom")
            worksheet.cell(row, 2).value = _return_description(item)
            worksheet.cell(row, 3).value = return_sn
            fault_date = _excel_date(item.get("report_date"))
            worksheet.cell(row, 4).value = fault_date
            if fault_date is not None:
                worksheet.cell(row, 4).number_format = "yyyy-mm-dd"
            worksheet.cell(row, 5).value = condition
            worksheet.cell(row, 8).value = request.get("spare_sr")
            worksheet.cell(row, 9).value = item.get("rma")
            fill = PatternFill("solid", fgColor=LIGHT_RED if condition == "Faulty" else LIGHT_GREEN)
            for column in range(1, 10):
                worksheet.cell(row, column).fill = copy(fill)

        day = today or datetime.now(ECUADOR_TIMEZONE).date()
        rmas = [str(item.get("rma")) for _, item, _ in selections]
        tracking = str(fault_tag_id or f"FT-{day.strftime('%y%m%d')}000000")
        filename = "–".join([tracking, cloud, day.strftime("%Y%m%d"), *rmas]) + ".xlsx"
        destination = export_root.expanduser().resolve() / "Returns" / filename
        final = _atomic_save(workbook, destination)
    finally:
        workbook.close()

    check = load_workbook(final, read_only=True, data_only=False)
    try:
        if check.sheetnames != list(sheet_names):
            raise SpareRequestWorkbookError("Generated return workbook changed template sheets")
        worksheet = check.worksheets[0]
        values = [worksheet.cell(8 + index, 9).value for index in range(len(selections))]
        if values != rmas:
            raise SpareRequestWorkbookError("Generated return workbook failed RMA verification")
    finally:
        check.close()
    return {
        "path": str(final),
        "filename": final.name,
        "subject": f"[FAULT TAG {tracking}] [{site}] RMA " + " ".join(rmas),
        "warnings": warnings,
    }


def _archive_email_inactivity(request: dict[str, Any], archived_at: str) -> int | None:
    last = parse_datetime(request.get("email", {}).get("last_activity_at"))
    archived = parse_datetime(archived_at)
    if last is None or archived is None:
        return None
    if last.tzinfo is not None:
        last = last.astimezone(ECUADOR_TIMEZONE)
    if archived.tzinfo is not None:
        archived = archived.astimezone(ECUADOR_TIMEZONE)
    return max(0, (archived.date() - last.date()).days)


def _ensure_archive_sheet(workbook: Any, name: str, headers: list[str]) -> Any:
    if name in workbook.sheetnames:
        worksheet = workbook[name]
        current = [str(cell.value or "").strip() for cell in worksheet[1]]
        if (
            name == SPARE_ARCHIVE_SHEET
            and headers == ARCHIVE_HEADERS
            and current == LEGACY_ARCHIVE_HEADERS
        ):
            alias_column = ARCHIVE_HEADERS.index("RMA Aliases") + 1
            worksheet.insert_cols(alias_column)
            source = worksheet.cell(1, alias_column - 1)
            target = worksheet.cell(1, alias_column)
            target.value = "RMA Aliases"
            target._style = copy(source._style)
            target.alignment = copy(source.alignment)
            target.protection = copy(source.protection)
            worksheet.auto_filter.ref = (
                f"A1:{get_column_letter(len(ARCHIVE_HEADERS))}1"
            )
            current = [str(cell.value or "").strip() for cell in worksheet[1]]
        if current != headers:
            raise SpareRequestWorkbookError(f"{name} in Closed.xlsx has an unsupported schema")
        return worksheet
    worksheet = workbook.create_sheet(name)
    worksheet.append(headers)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    worksheet.sheet_view.showGridLines = False
    for cell in worksheet[1]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        font = copy(cell.font)
        font.color = "FFFFFF"
        font.bold = True
        cell.font = font
    return worksheet


def _new_closed_workbook() -> Any:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Closed"
    worksheet.append(PENDING_COLUMNS)
    return workbook


def append_archived_item(
    closed_path: Path,
    request: dict[str, Any],
    item: dict[str, Any],
    *,
    reason: str,
    note: str | None,
    archived_at: str | None = None,
) -> dict[str, Any]:
    closed_path = closed_path.expanduser().resolve()
    closed_path.parent.mkdir(parents=True, exist_ok=True)
    timestamp = archived_at or iso_now()
    if closed_path.exists():
        workbook = load_workbook(closed_path, read_only=False, data_only=False, keep_links=True)
    else:
        workbook = _new_closed_workbook()
    temporary = closed_path.with_name(f".Closed-spare-{uuid.uuid4().hex}.xlsx")
    backup: Path | None = None
    try:
        archive = _ensure_archive_sheet(workbook, SPARE_ARCHIVE_SHEET, ARCHIVE_HEADERS)
        emails = _ensure_archive_sheet(
            workbook, SPARE_ARCHIVE_EMAIL_SHEET, ARCHIVE_EMAIL_HEADERS
        )
        item_id = str(item.get("item_id") or "")
        existing_ids = {
            str(archive.cell(row, ARCHIVE_HEADERS.index("Item ID") + 1).value or "")
            for row in range(2, archive.max_row + 1)
        }
        if item_id not in existing_ids:
            profile = request.get("profile", {})
            email = request.get("email", {})
            row = {
                "TT": request.get("tt"),
                "RMA": item.get("rma"),
                "RMA Aliases": "\n".join(item.get("rma_aliases") or []) or None,
                "Email inactivity days": _archive_email_inactivity(request, timestamp),
                "Emails received": int(email.get("total_received") or 0),
                "Emails sent": int(email.get("total_sent") or 0),
                "Spare SR": request.get("spare_sr"),
                "Request ID": request.get("request_id"),
                "Status": reason,
                "Archived At": timestamp,
                "Archive Reason": reason,
                "Client Initials": profile.get("client_initials"),
                "Customer": profile.get("customer_name"),
                "Site": profile.get("site_code"),
                "Address": profile.get("site_address"),
                "Cloud": profile.get("cloud"),
                "Requested BOM": item.get("requested_bom"),
                "Delivered BOM": item.get("delivered_bom"),
                "Description": item.get("requested_description"),
                "Part": item.get("part"),
                "Model": item.get("model"),
                "Device": item.get("device"),
                "Slot": item.get("slot"),
                "Faulty SN": item.get("faulty_sn"),
                "New SN": item.get("new_sn"),
                "Return Condition": item.get("return_condition"),
                "Dispatch Date": item.get("dispatch_at"),
                "Return Export": item.get("return_export_filename"),
                "RT": item.get("rt"),
                "Notes": note or item.get("notes"),
                "Source": request.get("source"),
                "Item ID": item_id,
            }
            archive.append([row.get(header) for header in ARCHIVE_HEADERS])
            for message in email.get("messages", []):
                associated = set(message.get("item_ids") or [])
                if associated and item_id not in associated:
                    continue
                email_row = {
                    "TT": request.get("tt"),
                    "Spare SR": request.get("spare_sr"),
                    "RMA": item.get("rma"),
                    "Item ID": item_id,
                    "Message ID": message.get("message_key"),
                    "Timestamp": message.get("timestamp"),
                    "Direction": message.get("direction"),
                    "Sender": message.get("sender"),
                    "Sender Address": message.get("sender_address"),
                    "Subject": message.get("subject"),
                    "Body": message.get("latest_reply_body") or message.get("body"),
                }
                emails.append([email_row.get(header) for header in ARCHIVE_EMAIL_HEADERS])
        workbook.save(temporary)
        verified = load_workbook(temporary, read_only=True, data_only=False)
        try:
            archived_ids = {
                str(row[ARCHIVE_HEADERS.index("Item ID")].value or "")
                for row in verified[SPARE_ARCHIVE_SHEET].iter_rows(min_row=2)
            }
            if item_id not in archived_ids:
                raise SpareRequestWorkbookError("Closed.xlsx failed Spare Request archive verification")
        finally:
            verified.close()
        if closed_path.exists():
            backup_dir = closed_path.parent / "Zeus Backups" / "Spare Requests"
            backup_dir.mkdir(parents=True, exist_ok=True)
            backup = backup_dir / f"Closed_{datetime.now().strftime('%Y%m%d_%H%M%S_%f')}.xlsx"
            shutil.copy2(closed_path, backup)
        os.replace(temporary, closed_path)
    finally:
        workbook.close()
        temporary.unlink(missing_ok=True)
    return {
        "closedPath": str(closed_path),
        "backup": str(backup) if backup else None,
        "itemId": item.get("item_id"),
        "archivedAt": timestamp,
    }


def read_archived_items(closed_path: Path) -> list[dict[str, Any]]:
    if not closed_path.is_file():
        return []
    workbook = load_workbook(closed_path, read_only=True, data_only=False)
    try:
        if SPARE_ARCHIVE_SHEET not in workbook.sheetnames:
            return []
        worksheet = workbook[SPARE_ARCHIVE_SHEET]
        headers = [str(cell.value or "").strip() for cell in worksheet[1]]
        if tuple(headers) not in {
            tuple(ARCHIVE_HEADERS),
            tuple(LEGACY_ARCHIVE_HEADERS),
        }:
            raise SpareRequestWorkbookError(
                f"{SPARE_ARCHIVE_SHEET} in Closed.xlsx has an unsupported schema"
            )
        rows: list[dict[str, Any]] = []
        for values in worksheet.iter_rows(min_row=2, values_only=True):
            if not any(value not in (None, "") for value in values):
                continue
            row = dict(zip(headers, values))
            row.setdefault("RMA Aliases", None)
            rows.append(row)
        return rows
    finally:
        workbook.close()


def purge_archived_items(
    closed_path: Path,
    *,
    item_ids: set[str] | None = None,
    archived_before: datetime | None = None,
) -> dict[str, Any]:
    if not closed_path.is_file():
        return {"removed": 0, "remaining": 0}
    workbook = load_workbook(closed_path, read_only=False, data_only=False, keep_links=True)
    temporary = closed_path.with_name(f".Closed-purge-{uuid.uuid4().hex}.xlsx")
    removed_ids: set[str] = set()
    try:
        if SPARE_ARCHIVE_SHEET not in workbook.sheetnames:
            return {"removed": 0, "remaining": 0}
        archive = workbook[SPARE_ARCHIVE_SHEET]
        headers = [str(cell.value or "").strip() for cell in archive[1]]
        if tuple(headers) not in {
            tuple(ARCHIVE_HEADERS),
            tuple(LEGACY_ARCHIVE_HEADERS),
        }:
            raise SpareRequestWorkbookError(
                f"{SPARE_ARCHIVE_SHEET} in Closed.xlsx has an unsupported schema"
            )
        item_column = headers.index("Item ID") + 1
        archived_column = headers.index("Archived At") + 1
        for row in range(archive.max_row, 1, -1):
            item_id = str(archive.cell(row, item_column).value or "")
            timestamp = parse_datetime(archive.cell(row, archived_column).value)
            selected = item_ids is not None and item_id in item_ids
            expired = (
                archived_before is not None
                and timestamp is not None
                and (
                    timestamp.replace(tzinfo=None)
                    if timestamp.tzinfo is not None
                    else timestamp
                )
                < (
                    archived_before.replace(tzinfo=None)
                    if archived_before.tzinfo is not None
                    else archived_before
                )
            )
            if selected or expired:
                removed_ids.add(item_id)
                archive.delete_rows(row)
        if SPARE_ARCHIVE_EMAIL_SHEET in workbook.sheetnames and removed_ids:
            emails = workbook[SPARE_ARCHIVE_EMAIL_SHEET]
            email_headers = [str(cell.value or "").strip() for cell in emails[1]]
            if email_headers == ARCHIVE_EMAIL_HEADERS:
                email_item_column = email_headers.index("Item ID") + 1
                for row in range(emails.max_row, 1, -1):
                    if str(emails.cell(row, email_item_column).value or "") in removed_ids:
                        emails.delete_rows(row)
        if removed_ids:
            workbook.save(temporary)
            os.replace(temporary, closed_path)
        remaining = max(0, archive.max_row - 1)
        return {"removed": len(removed_ids), "remaining": remaining}
    finally:
        workbook.close()
        temporary.unlink(missing_ok=True)


def purge_spare_archive_backups(
    closed_path: Path,
    *,
    before: datetime | None = None,
    remove_all: bool = False,
) -> int:
    """Discard Closed backups that could retain purged Spare Request data."""

    root = closed_path.expanduser().resolve().parent / "Zeus Backups"
    candidates = [
        *root.glob("Closed_*.xlsx"),
        *(root / "Spare Requests").glob("Closed_*.xlsx"),
    ]
    cutoff = (
        before.replace(tzinfo=None) if before is not None and before.tzinfo is not None else before
    )
    removed = 0
    for path in candidates:
        try:
            modified = datetime.fromtimestamp(path.stat().st_mtime)
        except OSError:
            continue
        if not remove_all and (cutoff is None or modified >= cutoff):
            continue
        path.unlink(missing_ok=True)
        removed += 1
    return removed
