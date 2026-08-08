from __future__ import annotations

import argparse
import os
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook

# Playwright starts this file directly from ``frontend/``.  Put the checkout at
# the front of sys.path so the fixture always exercises this branch, even when
# another Zeus build is installed in the runner environment.
PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from zeus2.config import save_config
from zeus2.excel_export import publish_operational_workbooks
from zeus2.main import main
from zeus2.reference_data import save_user_profile
from zeus2.spare_request_excel import (
    FAULTY_TAG_SHEET,
    REQUEST_SHEET,
    RETURN_SHEET,
    append_archived_item,
)
from zeus2.spare_requests import (
    create_request_record,
    normalize_profile,
    normalize_request_lines,
    request_subject,
)
from zeus2.startup import run_startup
from zeus2.store import ZeusStore
from zeus2.tickets import LOCAL_COLUMNS, UPSTREAM_COLUMNS, normalize_local


def pending_row(ticket_id: str, **values: object) -> dict[str, object]:
    """Build an independent Pendings fixture for the installed-app browser test."""

    summary = values.pop("summary", None)
    row: dict[str, object] = {
        "SRNo": ticket_id,
        "Problem Summary": summary or f"Ticket {ticket_id}",
        "Report Date": "2026-07-01 10:00:00",
        "Customer Contact": "Customer",
        "Customer Severity": values.pop("severity", "Minor"),
        "Product": "Product",
        "Current Handler": "Handler A",
        "Status": "L1-Work in Progress",
        "ResolveBy": "2026-08-20 10:00:00",
        "Resolve By Suspend": "2026-08-25 10:00:00",
        **{column: None for column in LOCAL_COLUMNS},
        "Done?": "N",
    }
    row.update(values)
    return row


def write_advanced(path: Path, rows: list[dict[str, object]]) -> None:
    """Write the discovery source used by the database-first browser fixture."""

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Service Request"
    worksheet.append(list(UPSTREAM_COLUMNS))
    for row in rows:
        worksheet.append([row.get(column) for column in UPSTREAM_COLUMNS])
        worksheet.cell(worksheet.max_row, 1).hyperlink = (
            f"https://example.invalid/sr/{row['SRNo']}"
        )
    workbook.save(path)
    workbook.close()


def prepare_fixture(root: Path) -> Path:
    home = root / "home"
    workbooks = root / "workbooks"
    downloads = root / "downloads"
    spare_exports = root / "spare-exports"
    workbooks.mkdir()
    downloads.mkdir()
    spare_exports.mkdir()
    store = ZeusStore(home / "data", config_home=home)
    store.ensure_layout()
    save_user_profile(
        store.root,
        {
            "name": "Zeus Browser User",
            "email": "browser.user@example.com",
            "phone": "+593 99 000 0000",
            "username": "browser-user",
        },
    )
    config = store.config
    config["paths"]["workbook_directory"] = str(workbooks)
    config["paths"]["advanced_search_directory"] = str(downloads)
    config["paths"]["outlook_store_path"] = None
    request_template = root / "spare-request-template.xlsx"
    request_book = Workbook()
    request_book.active.title = REQUEST_SHEET
    request_book.create_sheet(FAULTY_TAG_SHEET)
    request_book.save(request_template)
    request_book.close()
    return_template = root / "spare-return-template.xlsx"
    return_book = Workbook()
    return_book.active.title = RETURN_SHEET
    return_book.save(return_template)
    return_book.close()
    config["paths"]["spare_parts_export_directory"] = str(spare_exports)
    config["paths"]["spare_request_template_path"] = str(request_template)
    config["paths"]["spare_return_template_path"] = str(return_template)
    config["advanced_search"]["poll_interval_minutes"] = 0
    save_config(home, config)

    severities = ("Critical", "Major", "Minor", "Minor")
    rows = [
        pending_row(
            str(39400000 + index),
            summary=f"Ecuador | Claro | Filtered network incident {index:02d}",
            report_date=f"2026-07-{max(1, 31 - index):02d} 10:00:00",
            severity=severities[index % len(severities)],
            **{
                "Planned Date": "Unplanned",
                "Site": "GYE" if index % 2 else "UIO",
                "Cloud": "FusionSphere",
                "Model": "2288H V5",
                "Device": f"server-{index:02d}",
                "Slot": "Slot 1",
                "Part": "Disk",
                "BOM": f"BOM-{index:02d}",
                "Old SN": f"FAULTY-{index:02d}",
                "New SN": "pending",
                "Notes": "Awaiting field confirmation",
                "Done?": "N",
            },
        )
        for index in range(1, 35)
    ]
    write_advanced(
        downloads / "Advanced Search(Service Request)20260808010101.xlsx",
        rows,
    )
    run_startup(store)
    with store.transaction("e2e-database-work-fields", {}) as staging:
        for row in rows:
            ticket_id = str(row["SRNo"])
            ticket = store.read_ticket(ticket_id, staging)
            # Omit spare_parts so normalization performs the one-time legacy
            # scalar-to-hierarchy migration inside the authoritative database.
            ticket["local"] = normalize_local(
                {"fields": {column: row.get(column) for column in LOCAL_COLUMNS}}
            )
            store.write_ticket_bundle(staging, ticket)
    legacy = store.read_ticket("39400001")
    quoted_lines = "\n".join(
        f"Quoted historical line {index:03d}: prior diagnostic context"
        for index in range(1, 241)
    )
    raw_body = (
        "Newest field response: the DIMM alarm is clear.\n\n"
        "Regards/Saludos cordiales,\n\n"
        "De: Previous Engineer <previous@example.com>\n"
        "Enviado el: lunes, 27 de julio de 2026 18:30\n"
        "Para: Cloud Support <cloud@example.com>\n"
        "CC: Operations <operations@example.com>\n"
        "Asunto: RE: [SR 39400001] DIMM MCE error\n\n"
        f"{quoted_lines}"
    )
    legacy["email"].update(
        {
            "total_received": 1,
            "last_activity_at": "2026-07-27T18:30:14Z",
            "last_direction": "received",
            # Intentionally omit the compact-reply metadata.  This reproduces
            # retained Zeus records created before the web workstation.
            "messages": [
                {
                    "message_key": "legacy-browser-reply-chain",
                    "timestamp": "2026-07-27T18:30:14Z",
                    "direction": "received",
                    "subject": "RE: [SR 39400001] DIMM MCE error",
                    "sender": "field.engineer@example.com",
                    "body": raw_body,
                }
            ],
        }
    )
    store.write_ticket_bundle(store.current, legacy)
    # Keep one finalized SR in the real-browser fixture.  Its normalized spare
    # record must remain in the Spare Parts workspace through Closed.xlsx even
    # after the live Markdown bundle is removed.
    finalized = store.read_ticket("39400002")
    finalized["lifecycle"]["status"] = "closure_pending"
    with store.transaction("e2e-finalized-spare-parts", {}) as staging:
        store.write_ticket_bundle(staging, finalized)
    publish_operational_workbooks(store, workbooks, create_missing=True)
    profile = normalize_profile(
        {
            "clientInitials": "CNT",
            "customerName": "Customer Network Team",
            "siteCode": "UIO",
            "siteAddress": "Quito operations center",
            "cloud": "FusionSphere",
            "requester": {"name": "Zeus User", "email": "user@example.com"},
            "contact": {"name": "Customer", "email": "customer@example.com"},
        }
    )
    lines = normalize_request_lines(
        [
            {
                "bom": "BOM-01",
                "amount": 2,
                "description": "Disk",
                "part": "Disk",
                "model": "2288H V5",
                "device": "server-01",
                "slot": "Slot 1",
                "faultySn": "FAULTY-01",
                "reportDate": "2026-07-01",
            }
        ]
    )
    active = create_request_record(
        request_id="260808123456",
        tt="39400001",
        source="ticket",
        profile=profile,
        lines=lines,
        export_path=None,
        subject=request_subject("260808123456", "39400001", lines),
    )
    active["spare_sr"] = "SR4956964"
    active["items"][0].update(
        {
            "rma": "C3209937826",
            "attendance_confirmed_at": "2026-08-02T10:00:00-05:00",
            "attendance_source": "email",
            "delivered_bom": "02540255",
            "new_sn": "NEW-01",
            "dispatch_at": "2026-08-03T10:00:00-05:00",
            "dispatch_source": "email",
        }
    )
    active["items"][1].update(
        {
            "attendance_confirmed_at": "2026-08-02T10:00:00-05:00",
            "attendance_source": "email",
        }
    )
    with store.transaction("e2e-spare-request", {}) as staging:
        store.write_spare_request(staging, active)
    completed = create_request_record(
        request_id="260807123456",
        tt="39400003",
        source="ticket",
        profile=profile,
        lines=normalize_request_lines([{**lines[0], "amount": 1}]),
        export_path=None,
        subject="completed",
    )
    completed["spare_sr"] = "SR4956963"
    completed["items"][0].update(
        {
            "rma": "C3209937825",
            "delivered_bom": "BOM-03",
            "new_sn": "NEW-03",
            "dispatch_at": "2026-07-15T10:00:00-05:00",
        }
    )
    append_archived_item(
        workbooks / "Closed.xlsx",
        completed,
        completed["items"][0],
        reason="returned",
        note="E2E completed fixture",
    )
    return home


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--port", type=int, default=8765)
    args = parser.parse_args()
    with tempfile.TemporaryDirectory(prefix="zeus3-e2e-") as temporary:
        os.environ["ZEUS_HOME"] = str(prepare_fixture(Path(temporary)))
        raise SystemExit(
            main(
                [
                    "serve",
                    "--port",
                    str(args.port),
                    "--no-browser",
                    "--no-tray",
                    "--console",
                ]
            )
        )
