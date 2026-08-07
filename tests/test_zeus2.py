from __future__ import annotations

import json
import os
import re
import tempfile
import unittest
from concurrent.futures import Future
from copy import deepcopy
from contextlib import redirect_stderr, redirect_stdout
from datetime import date, datetime
from io import StringIO
from pathlib import Path
from unittest.mock import patch

from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from zeus2.aging import aging_for_ticket, calculate_ticket_facts, report_sort_key
from zeus2.cli import (
    ENABLE_ECHO_INPUT,
    ENABLE_EXTENDED_FLAGS,
    ENABLE_LINE_INPUT,
    ENABLE_MOUSE_INPUT,
    ENABLE_QUICK_EDIT_MODE,
    ENABLE_VIRTUAL_TERMINAL_INPUT,
    ENABLE_VIRTUAL_TERMINAL_PROCESSING,
    ENABLE_WINDOW_INPUT,
    MouseEvent,
    ZeusTUI,
    _decode_escape_sequence,
    _display_width,
    _edit_outlook_store_setting,
    _first_run_setup,
    _interactive_config,
    _parse_sgr_mouse,
    _screen_mode,
    _truncate_ansi,
    _windows_console_mode_plan,
    main,
)
from zeus2.config import (
    SETTING_SPECS,
    ensure_config,
    load_config,
    prepare_outlook_store_path,
    save_config,
    scan_outlook_store_files,
    set_dotted,
)
from zeus2.excel_export import (
    REPORT_COLUMNS,
    WorkbookPublicationError,
    publish_operational_workbooks,
    recover_pendings_restore,
    recover_publication,
    restore_pendings_backup,
)
from zeus2.excel_import import (
    WorkbookValidationError,
    read_pendings,
    validate_pendings_against_state,
)
from zeus2.mail import (
    MailSyncError,
    _select_outlook_store,
    commit_fetched_messages,
    extract_ticket_ids,
    fetch_outlook_messages,
    interval_due,
    strip_quoted_history,
    synchronize_staged_email,
)
from zeus2.mop import generate_mop
from zeus2.operations import ReadOnlyTicketError, append_note, update_ticket
from zeus2.reconcile import (
    ReconciliationError,
    import_pendings,
    initialize_closed_index,
    sync_advanced_search,
)
from zeus2.startup import StartupResult, run_startup
from zeus2.store import ZeusStore
from zeus2.tickets import (
    CLOSED_SCHEMA_DRIFT_COLUMNS,
    LOCAL_COLUMNS,
    PENDING_COLUMNS,
    UPSTREAM_COLUMNS,
    empty_local,
    new_ticket,
)
from zeus2.utils import normalize_ticket_id, sha256_file


def upstream_row(
    ticket_id: str,
    *,
    summary: str | None = None,
    report_date: str = "2026-07-01 10:00:00",
    severity: str = "Minor",
    status: str = "L1-Work in Progress",
    handler: str = "Handler A",
) -> dict[str, object]:
    return {
        "SRNo": ticket_id,
        "Problem Summary": summary or f"Ticket {ticket_id}",
        "Report Date": report_date,
        "Customer Contact": "Customer",
        "Customer Severity": severity,
        "Product": "Product",
        "Current Handler": handler,
        "Status": status,
        "ResolveBy": "2026-08-20 10:00:00",
        "Resolve By Suspend": "2026-08-25 10:00:00",
    }


def pending_row(ticket_id: str, **values: object) -> dict[str, object]:
    summary = values.pop("summary", None)
    row: dict[str, object] = {
        **upstream_row(ticket_id, summary=str(summary) if summary is not None else None),
        **{column: None for column in LOCAL_COLUMNS},
        "Done?": "N",
    }
    row.update(values)
    return row


def write_managed(
    path: Path,
    rows: list[dict[str, object]],
    *,
    headers: list[str] | None = None,
    sheet_name: str = "Pendings",
    style_cells: dict[tuple[str, str], tuple[str, str]] | None = None,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name
    headers = headers or list(PENDING_COLUMNS)
    worksheet.append(headers)
    for row in rows:
        worksheet.append([row.get(column) for column in headers])
        ticket_id = str(row["SRNo"])
        for column, colors in {
            key[1]: value
            for key, value in (style_cells or {}).items()
            if key[0] == ticket_id
        }.items():
            cell = worksheet.cell(worksheet.max_row, headers.index(column) + 1)
            cell.fill = PatternFill("solid", fgColor=colors[0])
            cell.font = Font(color=colors[1], bold=True)
    workbook.save(path)
    workbook.close()


def write_advanced(path: Path, rows: list[dict[str, object]]) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Service Request"
    worksheet.append(UPSTREAM_COLUMNS)
    for row in rows:
        worksheet.append([row.get(column) for column in UPSTREAM_COLUMNS])
        worksheet.cell(worksheet.max_row, 1).hyperlink = (
            f"https://example.invalid/sr/{row['SRNo']}"
        )
    workbook.save(path)
    workbook.close()


class ZeusCase(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.home = self.root / "home"
        self.books = self.root / "books"
        self.downloads = self.root / "downloads"
        self.books.mkdir()
        self.downloads.mkdir()
        self.store = ZeusStore(self.home / "data", config_home=self.home)
        self.store.ensure_layout()
        config = self.store.config
        config["paths"]["workbook_directory"] = str(self.books)
        config["paths"]["advanced_search_directory"] = str(self.downloads)
        self.store.save_config(config)

    def tearDown(self) -> None:
        self.temp.cleanup()

    def seed(
        self,
        pending_rows: list[dict[str, object]],
        advanced_rows: list[dict[str, object]] | None = None,
        closed_rows: list[dict[str, object]] | None = None,
    ) -> Path:
        write_managed(self.books / "Pendings.xlsx", pending_rows)
        write_managed(
            self.books / "Closed.xlsx",
            closed_rows or [],
            sheet_name="Closed",
        )
        source = self.downloads / "Advanced Search(Service Request)20260801010101.xlsx"
        write_advanced(source, advanced_rows if advanced_rows is not None else pending_rows)
        import_pendings(self.store, self.books / "Pendings.xlsx")
        initialize_closed_index(self.store, self.books / "Closed.xlsx")
        sync_advanced_search(self.store, source)
        return source


class UtilityAndConfigTests(ZeusCase):
    def test_ticket_ids_are_exactly_eight_digits(self) -> None:
        self.assertEqual(normalize_ticket_id(12345678), "12345678")
        self.assertEqual(normalize_ticket_id("12345678.0"), "12345678")
        for invalid in ("1234567", "123456789", 1234, 12345678.5):
            with self.assertRaises(ValueError):
                normalize_ticket_id(invalid)

    def test_locked_defaults_and_config_validation(self) -> None:
        config = load_config(self.home)
        self.assertEqual(config["advanced_search"]["poll_interval_minutes"], 15)
        self.assertEqual(config["email"]["fetch_interval_days"], 7)
        self.assertEqual(config["email"]["sync_mode"], "scheduled")
        self.assertEqual(config["email"]["sync_interval_days"], 7)
        self.assertEqual(config["email"]["retained_message_count"], 7)
        self.assertTrue(config["email"]["fetch_new_ticket_history_automatically"])
        bad = deepcopy(config)
        bad["email"]["fetch_interval_days"] = -2
        with self.assertRaises(ValueError):
            save_config(self.home, bad)

    def test_calendar_interval_convention(self) -> None:
        today = date(2026, 8, 6)
        self.assertTrue(interval_due(-1, "2026-08-06", today=today))
        self.assertFalse(interval_due(0, None, today=today))
        self.assertTrue(interval_due(7, "2026-07-30 23:59:59", today=today))
        self.assertFalse(interval_due(7, "2026-07-31 00:00:00", today=today))

    def test_outlook_configuration_requires_an_exact_store_file(self) -> None:
        folder = self.root / "email"
        folder.mkdir()
        with self.assertRaisesRegex(ValueError, "points to a folder"):
            prepare_outlook_store_path(folder)
        with self.assertRaisesRegex(ValueError, "must name an .ost or .pst"):
            prepare_outlook_store_path(folder / "mailbox")
        expected = (folder / "primary.ost").resolve()
        self.assertEqual(prepare_outlook_store_path(expected), expected)

    def test_screenshot_duplicate_outlook_key_repairs_the_canonical_setting(self) -> None:
        """Regression: the old raw editor accepted a non-canonical root key."""

        folder = r"D:\OutlookData"
        expected = r"D:\OutlookData\primary-mailbox.ost"
        saved = self.store.config
        saved["paths"]["outlook_store_path"] = folder
        saved["outlook_store_path"] = expected
        self.store.config_file.write_text(json.dumps(saved), encoding="utf-8")

        repaired = load_config(self.home)

        self.assertEqual(repaired["paths"]["outlook_store_path"], expected)
        self.assertNotIn("outlook_store_path", repaired)

        ensure_config(self.home)
        saved_again = json.loads(self.store.config_file.read_text(encoding="utf-8"))
        self.assertEqual(saved_again["paths"]["outlook_store_path"], expected)
        self.assertNotIn("outlook_store_path", saved_again)

    def test_existing_folder_setting_is_repaired_before_startup(self) -> None:
        folder = self.root / "email"
        folder.mkdir()
        config = self.store.config
        config["paths"]["outlook_store_path"] = str(folder)
        expected = folder / "primary.ost"
        expected.touch()
        self.store.config_file.write_text(json.dumps(config), encoding="utf-8")
        with patch("builtins.input", side_effect=["", "1"]):
            _first_run_setup(self.store)
        self.assertEqual(
            self.store.config["paths"]["outlook_store_path"],
            str(expected.resolve()),
        )

    def test_outlook_store_picker_scans_selects_and_deselects(self) -> None:
        folder = self.root / "email"
        folder.mkdir()
        ignored = folder / "readme.txt"
        ignored.touch()
        archive = folder / "archive.pst"
        primary = folder / "primary.ost"
        archive.touch()
        primary.touch()
        # Windows may expose the temporary directory through its 8.3 alias
        # while ``Path.resolve()`` returns the long spelling. Compare the
        # canonical paths that Zeus intentionally persists.
        self.assertEqual(
            scan_outlook_store_files(folder),
            [archive.resolve(), primary.resolve()],
        )

        config = self.store.config
        with patch("builtins.input", side_effect=[str(folder), "2"]):
            changed = _edit_outlook_store_setting(config, base=self.home)
        self.assertTrue(changed)
        self.assertEqual(
            config["paths"]["outlook_store_path"],
            str(primary.resolve()),
        )

        with patch("builtins.input", side_effect=["", "2"]):
            changed = _edit_outlook_store_setting(config, base=self.home)
        self.assertTrue(changed)
        self.assertIsNone(config["paths"]["outlook_store_path"])

    def test_outlook_store_picker_failure_preserves_the_setting(self) -> None:
        folder = self.root / "empty-email-folder"
        folder.mkdir()
        config = self.store.config
        before = deepcopy(config)
        with patch("builtins.input", return_value=str(folder)):
            with self.assertRaisesRegex(ValueError, "No .ost or .pst files"):
                _edit_outlook_store_setting(config, base=self.home)
        self.assertEqual(config, before)

    def test_unknown_raw_configuration_key_is_rejected(self) -> None:
        config = self.store.config
        with self.assertRaisesRegex(ValueError, "Unknown configuration setting"):
            set_dotted(config, "outlook_store_path", "D:/Email/primary.ost")
        config["outlook_store_path"] = "D:/Email/primary.ost"
        with self.assertRaisesRegex(ValueError, "Unknown configuration setting"):
            save_config(self.home, config)

    def test_configuration_menu_lists_items_and_changes_only_the_selection(self) -> None:
        before = self.store.config
        output = StringIO()
        poll_index = next(
            index
            for index, spec in enumerate(SETTING_SPECS, start=1)
            if spec.key == "advanced_search.poll_interval_minutes"
        )
        with patch("builtins.input", side_effect=[str(poll_index), "30", ""]):
            with redirect_stdout(output):
                _interactive_config(self.store)

        after = self.store.config
        self.assertEqual(after["advanced_search"]["poll_interval_minutes"], 30)
        self.assertEqual(after["email"], before["email"])
        rendered = output.getvalue()
        self.assertIn("Data query interval", rendered)
        self.assertIn("Outlook mailbox store", rendered)
        self.assertNotIn('"schema_version"', rendered)

    def test_outlook_folder_warning_suggests_the_unique_ost(self) -> None:
        class FakeStore:
            def __init__(self, path: str):
                self.FilePath = path

        class FakeStores:
            def __init__(self, paths: list[str]):
                self._stores = [FakeStore(path) for path in paths]
                self.Count = len(self._stores)

            def Item(self, index: int) -> FakeStore:
                return self._stores[index - 1]

        class FakeNamespace:
            def __init__(self, paths: list[str]):
                self.Stores = FakeStores(paths)

        namespace = FakeNamespace(
            [
                r"D:\OutlookData\primary-mailbox.ost",
                r"D:\OutlookData\archive-2026.pst",
            ]
        )
        with self.assertRaises(MailSyncError) as caught:
            _select_outlook_store(namespace, Path(r"D:\OutlookData"))
        message = str(caught.exception)
        self.assertIn("folder, not a store file", message)
        self.assertIn(r"D:\OutlookData\primary-mailbox.ost", message)

        selected = _select_outlook_store(
            namespace, Path(r"d:\outlookdata\PRIMARY-MAILBOX.OST")
        )
        self.assertEqual(selected.FilePath, r"D:\OutlookData\primary-mailbox.ost")


class StoreTests(ZeusCase):
    def test_markdown_is_the_self_contained_ticket_record(self) -> None:
        ticket = new_ticket(
            "12345678",
            upstream_row("12345678"),
            {"filename": "source.xlsx"},
        )
        ticket["local"]["fields"]["Notes"] = "A visible note"
        ticket["email"]["total_received"] = 1
        ticket["email"]["messages"] = [
            {
                "message_key": "a" * 64,
                "timestamp": "2026-08-01T10:00:00",
                "direction": "received",
                "subject": "SR 12345678",
                "body": "Complete body",
            }
        ]
        with self.store.transaction("seed", {}) as staging:
            self.store.write_ticket_bundle(staging, ticket)
        path = self.store.ticket_file("12345678")
        text = path.read_text(encoding="utf-8")
        self.assertTrue(path.name == "12345678.md")
        self.assertIn("# SR 12345678", text)
        self.assertIn("Complete body", text)
        self.assertFalse((path.parent / "ticket.json").exists())
        self.assertEqual(
            self.store.read_ticket("12345678")["local"]["fields"]["Notes"],
            "A visible note",
        )

    def test_failed_transaction_keeps_previous_record(self) -> None:
        ticket = new_ticket("12345678", upstream_row("12345678"), {})
        with self.store.transaction("seed", {}) as staging:
            self.store.write_ticket_bundle(staging, ticket)
        with self.assertRaises(RuntimeError):
            with self.store.transaction("fail", {}) as staging:
                changed = self.store.read_ticket("12345678", staging)
                changed["local"]["fields"]["Notes"] = "should roll back"
                self.store.write_ticket_bundle(staging, changed)
                raise RuntimeError("boom")
        self.assertIsNone(
            self.store.read_ticket("12345678")["local"]["fields"]["Notes"]
        )


class WorkbookImportTests(ZeusCase):
    def test_reordered_headers_styles_and_done_normalization(self) -> None:
        headers = list(reversed(PENDING_COLUMNS))
        path = self.books / "Pendings.xlsx"
        write_managed(
            path,
            [pending_row("12345678", **{"Done?": "bad", "Notes": "hello"})],
            headers=headers,
            style_cells={("12345678", "Notes"): ("FF0000", "00FF00")},
        )
        managed = read_pendings(path)
        record = managed.records["12345678"]
        self.assertEqual(managed.header_order, headers)
        self.assertEqual(record["local"]["fields"]["Done?"], "N")
        self.assertEqual(managed.done_corrections[0]["ticket_id"], "12345678")
        style = record["local"]["presentation"]["cell_styles"]["Notes"]
        self.assertEqual(style["fill"]["fg_color"]["value"][-6:], "FF0000")
        self.assertEqual(style["font"]["color"]["value"][-6:], "00FF00")

    def test_formulas_unknown_columns_invalid_and_duplicate_ids_are_rejected(self) -> None:
        formula = self.books / "formula.xlsx"
        write_managed(formula, [pending_row("12345678")])
        workbook = load_workbook(formula)
        workbook.active["N2"] = "=1+1"
        workbook.save(formula)
        workbook.close()
        with self.assertRaises(WorkbookValidationError):
            read_pendings(formula)

        unknown = self.books / "unknown.xlsx"
        write_managed(
            unknown,
            [pending_row("12345678", Custom="x")],
            headers=PENDING_COLUMNS + ["Custom"],
        )
        with self.assertRaises(WorkbookValidationError):
            read_pendings(unknown)

        duplicate = self.books / "duplicate.xlsx"
        write_managed(
            duplicate,
            [pending_row("12345678"), pending_row("12345678")],
        )
        with self.assertRaises(WorkbookValidationError):
            read_pendings(duplicate)

        invalid = self.books / "invalid.xlsx"
        write_managed(invalid, [pending_row("1234567")])
        with self.assertRaises(WorkbookValidationError):
            read_pendings(invalid)

        headerless = self.books / "headerless.xlsx"
        write_managed(headerless, [pending_row("12345678")])
        workbook = load_workbook(headerless)
        workbook.active.cell(2, len(PENDING_COLUMNS) + 2, "hidden data")
        workbook.save(headerless)
        workbook.close()
        with self.assertRaises(WorkbookValidationError):
            read_pendings(headerless)

    def test_protected_snapshot_distinguishes_upstream_refresh_from_excel_edit(self) -> None:
        pending = self.books / "Pendings.xlsx"
        write_managed(pending, [pending_row("12345678", summary="Old")])
        import_pendings(self.store, pending)
        snapshot = self.store.state()["publication_state"]["pendings_snapshot"]

        # The database may now hold a newer Advanced Search value while the
        # still-unpublished workbook correctly retains its previous snapshot.
        ticket = self.store.read_ticket("12345678")
        ticket["upstream"]["fields"]["Problem Summary"] = "New upstream"
        with self.store.transaction("upstream-test", {}) as staging:
            self.store.write_ticket_bundle(staging, ticket)
        managed = read_pendings(pending)
        validate_pendings_against_state(
            managed,
            database_ids={"12345678"},
            snapshot=snapshot,
        )

        workbook = load_workbook(pending)
        workbook.active["B2"] = "User changed protected field"
        workbook.save(pending)
        workbook.close()
        with self.assertRaises(WorkbookValidationError):
            validate_pendings_against_state(
                read_pendings(pending),
                database_ids={"12345678"},
                snapshot=snapshot,
            )

    def test_missing_exported_row_blocks_but_unexported_new_ticket_may_be_absent(self) -> None:
        pending = self.books / "Pendings.xlsx"
        write_managed(pending, [pending_row("12345678")])
        import_pendings(self.store, pending)
        snapshot = self.store.state()["publication_state"]["pendings_snapshot"]
        new = new_ticket("12345679", upstream_row("12345679"), {})
        with self.store.transaction("new-ticket", {}) as staging:
            self.store.write_ticket_bundle(staging, new)
        validate_pendings_against_state(
            read_pendings(pending),
            database_ids={"12345678", "12345679"},
            snapshot=snapshot,
        )
        write_managed(pending, [])
        with self.assertRaises(WorkbookValidationError):
            validate_pendings_against_state(
                read_pendings(pending),
                database_ids={"12345678", "12345679"},
                snapshot=snapshot,
            )


class ReconciliationTests(ZeusCase):
    def test_sources_own_disjoint_fields_and_closure_is_deferred(self) -> None:
        self.seed(
            [pending_row("12345678", **{"Planned Date": "2026-08-12", "Notes": "local"})],
            [upstream_row("12345678", summary="Online", handler="New handler")],
        )
        ticket = self.store.read_ticket("12345678")
        self.assertEqual(ticket["upstream"]["fields"]["Problem Summary"], "Online")
        self.assertEqual(ticket["local"]["fields"]["Notes"], "local")

        second = self.downloads / "Advanced Search(Service Request)20260802010101.xlsx"
        write_advanced(second, [upstream_row("12345679", summary="New")])
        result = sync_advanced_search(self.store, second)
        self.assertEqual(result["added_ids"], ["12345679"])
        self.assertEqual(result["closure_pending_ids"], ["12345678"])
        self.assertEqual(
            self.store.read_ticket("12345678")["lifecycle"]["status"],
            "closure_pending",
        )
        self.assertTrue(self.store.ticket_file("12345678").exists())
        self.assertEqual(
            self.store.read_ticket("12345679")["local"]["fields"]["Done?"],
            "N",
        )

        third = self.downloads / "Advanced Search(Service Request)20260803010101.xlsx"
        write_advanced(third, [upstream_row("12345678"), upstream_row("12345679")])
        sync_advanced_search(self.store, third)
        self.assertEqual(
            self.store.read_ticket("12345678")["lifecycle"]["status"], "active"
        )

    def test_same_filename_changed_contents_and_closed_reappearance_are_blocked(self) -> None:
        source = self.seed([pending_row("12345678")])
        write_advanced(source, [upstream_row("12345678", summary="mutated")])
        with self.assertRaises(ReconciliationError):
            sync_advanced_search(self.store, source)

        index = self.store.closed_index()
        index["ticket_ids"] = ["12345679"]
        with self.store.transaction("closed-index-test", {}) as staging:
            from zeus2.utils import atomic_write_json

            atomic_write_json(staging / "closed_index.json", index)
        newer = self.downloads / "Advanced Search(Service Request)20260802010101.xlsx"
        write_advanced(newer, [upstream_row("12345679")])
        with self.assertRaises(ReconciliationError):
            sync_advanced_search(self.store, newer)


class EmailTests(ZeusCase):
    def setUp(self) -> None:
        super().setUp()
        self.seed([pending_row("12345678"), pending_row("12345679")])

    def test_subject_classifier(self) -> None:
        known = {"12345678", "12345679"}
        self.assertEqual(
            extract_ticket_ids("RE: SR 12345678 / issue", known), ["12345678"]
        )
        self.assertEqual(
            extract_ticket_ids("References 12345678 and 12345679", known),
            ["12345678", "12345679"],
        )
        self.assertEqual(
            extract_ticket_ids(
                "Spare Request: SR12345678 // TT: 12345679", known
            ),
            ["12345679"],
        )
        self.assertEqual(
            extract_ticket_ids("Spare Request: SR12345678 only", known), []
        )

    def _messages(self) -> list[dict[str, Any]]:
        return [
            {
                "message_id": "one",
                "ticket_ids": ["12345678"],
                "timestamp": "2026-08-01 10:00:00",
                "direction": "received",
                "subject": "SR 12345678 first",
                "body": "body one",
            },
            {
                "message_id": "two",
                "ticket_ids": ["12345678", "12345679"],
                "timestamp": "2026-08-02 10:00:00",
                "direction": "sent",
                "subject": "Both tickets",
                "body": "body two",
            },
        ]

    def test_fetch_stages_without_exposing_values_then_syncs_and_deduplicates(self) -> None:
        fetched = commit_fetched_messages(
            self.store,
            self._messages(),
            fetched_ticket_ids=["12345678", "12345679"],
            full_scan=True,
            synchronize=False,
        )
        self.assertEqual(fetched["new_staged_messages"], 2)
        self.assertEqual(
            self.store.read_ticket("12345678")["email"]["total_received"], 0
        )
        self.assertEqual(self.store.state()["email_state"]["staged_message_count"], 2)

        synchronized = synchronize_staged_email(self.store)
        self.assertEqual(synchronized["new_message_associations"], 3)
        first = self.store.read_ticket("12345678")["email"]
        second = self.store.read_ticket("12345679")["email"]
        self.assertEqual((first["total_received"], first["total_sent"]), (1, 1))
        self.assertEqual((second["total_received"], second["total_sent"]), (0, 1))
        self.assertEqual(first["last_direction"], "sent")
        self.assertEqual(first["messages"][0]["body"], "body two")

        commit_fetched_messages(
            self.store,
            self._messages(),
            fetched_ticket_ids=["12345678", "12345679"],
            full_scan=False,
            synchronize=True,
        )
        self.assertEqual(
            self.store.read_ticket("12345678")["email"]["total_received"], 1
        )

    def test_retained_count_zero_keeps_totals_without_bodies(self) -> None:
        config = self.store.config
        config["email"]["retained_message_count"] = 0
        self.store.save_config(config)
        commit_fetched_messages(
            self.store,
            self._messages(),
            fetched_ticket_ids=["12345678", "12345679"],
            full_scan=True,
            synchronize=True,
        )
        email = self.store.read_ticket("12345678")["email"]
        self.assertEqual(email["total_received"] + email["total_sent"], 2)
        self.assertEqual(email["messages"], [])

    def test_worker_failure_is_normalized_and_logged(self) -> None:
        failed: Future[object] = Future()
        failed.set_exception(RuntimeError("transient MAPI initialization failure"))

        with patch("zeus2.mail._OUTLOOK_EXECUTOR.submit", return_value=failed):
            with self.assertRaises(MailSyncError) as caught:
                fetch_outlook_messages(self.store)

        self.assertIsInstance(caught.exception.__cause__, RuntimeError)
        self.assertIn("Classic Outlook", str(caught.exception))
        diagnostic_log = self.home / "logs" / "zeus.log"
        self.assertIn(
            "transient MAPI initialization failure",
            diagnostic_log.read_text(encoding="utf-8"),
        )

    def test_reply_history_gets_a_compact_view_without_losing_the_raw_body(self) -> None:
        message = {
            "message_id": "reply-chain",
            "ticket_ids": ["12345678"],
            "timestamp": "2026-08-03 10:00:00",
            "direction": "received",
            "subject": "RE: SR 12345678 first",
            "body": (
                "The replacement was installed and the alarm is clear.\r\n\r\n"
                "From: First Engineer <first@example.com>\r\n"
                "Sent: Saturday, August 1, 2026 10:00 AM\r\n"
                "To: Support <support@example.com>\r\n"
                "Subject: SR 12345678 first\r\n\r\n"
                "This is the entire older conversation."
            ),
        }
        commit_fetched_messages(
            self.store,
            [message],
            fetched_ticket_ids=["12345678"],
            full_scan=True,
            synchronize=True,
        )
        retained = self.store.read_ticket("12345678")["email"]["messages"][0]
        self.assertEqual(
            retained["latest_reply_body"],
            "The replacement was installed and the alarm is clear.",
        )
        self.assertIn("This is the entire older conversation.", retained["body"])
        markdown = self.store.ticket_file("12345678").read_text(encoding="utf-8")
        self.assertIn(
            "The replacement was installed and the alarm is clear.", markdown
        )
        self.assertNotIn("This is the entire older conversation.", markdown)

    def test_localized_and_quoted_reply_markers_are_compacted_safely(self) -> None:
        spanish = (
            "Nueva actualización del cliente.\n\n"
            "De: Ingeniero Uno <one@example.com>\n"
            "Enviado el: lunes, 3 de agosto de 2026\n"
            "Para: Soporte <support@example.com>\n"
            "Asunto: RE: SR 12345678\n\n"
            "Conversación anterior"
        )
        cleaned, hidden, hidden_lines = strip_quoted_history(spanish)
        self.assertEqual(cleaned, "Nueva actualización del cliente.")
        self.assertTrue(hidden)
        self.assertGreaterEqual(hidden_lines, 5)

        quoted = "Newest answer\n\n> previous answer\n> older answer"
        self.assertEqual(strip_quoted_history(quoted)[0], "Newest answer")

        isolated = "Diagnostic output:\nFrom: switch-a\nPackets are healthy."
        self.assertEqual(strip_quoted_history(isolated), (isolated, False, 0))

        french = (
            "Nouvelle réponse.\n\nDe: Ingénieur\nEnvoyé: lundi\n"
            "À: Assistance\nObjet: RE: SR 12345678\n\nAncien échange"
        )
        self.assertEqual(strip_quoted_history(french)[0], "Nouvelle réponse.")


class InterfaceRegressionTests(ZeusCase):
    def test_long_unselected_summary_cannot_wrap_and_scroll_header_away(self) -> None:
        self.seed(
            [
                pending_row("12345679", summary="Short summary"),
                pending_row(
                    "12345678",
                    summary=(
                        "Ecuador | Claro | FusionSphere | Follow up of the DHCP "
                        "Service Is Unavailable HOST=nfvuio1fsc01 "
                        "HOST=nfvuio1fsc02 HOST=nfvuio1fsc03 activities | Site: NFV UIO"
                    ),
                ),
            ]
        )
        tui = ZeusTUI(self.store, StartupResult())
        tui.sort_index = tui.SORTS.index("sr")
        tui.selected = 0  # Keep the long row unselected, matching image 1.
        width = 96
        lines = tui._dashboard_lines(width, 20)
        ansi = re.compile(r"\x1b\[[0-?]*[ -/]*[@-~]")
        visible_lengths = [len(ansi.sub("", line)) for line in lines]
        self.assertLessEqual(max(visible_lengths), width)

    def test_redraw_erases_terminal_scrollback(self) -> None:
        self.seed([pending_row("12345678")])
        tui = ZeusTUI(self.store, StartupResult())
        with patch.object(tui, "_screen_size", return_value=(80, 16)):
            output = StringIO()
            with redirect_stdout(output):
                tui._draw()
        self.assertIn("\x1b[3J", output.getvalue())

    def test_styled_text_is_clipped_by_visible_cells_not_escape_bytes(self) -> None:
        rendered = "\x1b[91m" + "x" * 80 + "\x1b[0m"
        clipped = _truncate_ansi(rendered, 20, ellipsis="…")
        self.assertEqual(_display_width(clipped), 20)
        self.assertTrue(clipped.endswith("\x1b[0m"))

    def test_keyboard_pages_dashboard_and_scrolls_ticket_detail(self) -> None:
        self.seed(
            [pending_row(str(12345670 + index)) for index in range(1, 10)]
        )
        tui = ZeusTUI(self.store, StartupResult())
        tui.sort_index = tui.SORTS.index("sr")
        tui._dashboard_lines(100, 9)
        page_size = tui._dashboard_page_size
        tui._handle_key("pagedown")
        self.assertEqual(tui.selected, page_size)
        tui._handle_key("home")
        self.assertEqual(tui.selected, 0)
        tui._handle_key("end")
        self.assertEqual(tui.selected, 8)

        tui.view = "detail"
        tui._body_lines(80, 8)
        tui._handle_key("down")
        self.assertEqual(tui.detail_offset, 1)
        tui._handle_key("pagedown")
        self.assertGreater(tui.detail_offset, 1)
        tui._handle_key("home")
        self.assertEqual(tui.detail_offset, 0)

    def test_mouse_click_opens_ticket_and_wheel_navigates(self) -> None:
        self.seed(
            [pending_row(str(12345670 + index)) for index in range(1, 7)]
        )
        tui = ZeusTUI(self.store, StartupResult())
        tui.sort_index = tui.SORTS.index("sr")
        with patch.object(tui, "_screen_size", return_value=(100, 16)):
            with redirect_stdout(StringIO()):
                tui._draw()
        ticket_rows = {
            target[1]: row
            for row, target in tui._screen_targets.items()
            if target[0] == "ticket"
        }
        self.assertIn(1, ticket_rows)
        tui._handle_mouse(MouseEvent(10, ticket_rows[1], "click", button="left"))
        self.assertEqual((tui.view, tui.selected), ("detail", 1))

        tui.view = "dashboard"
        tui.selected = 5
        tui._handle_mouse(MouseEvent(10, 8, "wheel", delta=1))
        self.assertEqual(tui.selected, 2)

    def test_email_picker_shows_one_reply_and_supports_click_or_arrows(self) -> None:
        self.seed([pending_row("12345678")])
        commit_fetched_messages(
            self.store,
            [
                {
                    "message_id": "older",
                    "ticket_ids": ["12345678"],
                    "timestamp": "2026-08-01 10:00:00",
                    "direction": "received",
                    "subject": "Older",
                    "body": "OLDER UNIQUE BODY",
                },
                {
                    "message_id": "newer",
                    "ticket_ids": ["12345678"],
                    "timestamp": "2026-08-02 10:00:00",
                    "direction": "sent",
                    "subject": "Newer",
                    "body": (
                        "NEWER UNIQUE BODY\n\n"
                        "From: Earlier Engineer <earlier@example.com>\n"
                        "Sent: Friday, July 31, 2026 10:00 AM\n"
                        "To: Support <support@example.com>\n"
                        "Subject: Older chain\n\n"
                        "QUOTED HISTORY ONLY"
                    ),
                },
            ],
            fetched_ticket_ids=["12345678"],
            full_scan=True,
            synchronize=True,
        )
        tui = ZeusTUI(self.store, StartupResult())
        tui.view = "detail"
        detail = "\n".join(tui._detail_lines(self.store.read_ticket("12345678"), 100))
        self.assertIn("NEWER UNIQUE BODY", detail)
        self.assertNotIn("QUOTED HISTORY ONLY", detail)
        self.assertNotIn("OLDER UNIQUE BODY", detail)

        tui._handle_key("T")
        detail = "\n".join(tui._detail_lines(self.store.read_ticket("12345678"), 100))
        self.assertIn("QUOTED HISTORY ONLY", detail)
        tui._handle_key("T")

        tui._handle_key("right")
        detail = "\n".join(tui._detail_lines(self.store.read_ticket("12345678"), 100))
        self.assertIn("OLDER UNIQUE BODY", detail)
        self.assertNotIn("NEWER UNIQUE BODY", detail)

        with patch.object(tui, "_screen_size", return_value=(100, 22)):
            with redirect_stdout(StringIO()):
                tui._draw()
        email_rows = {
            target[1]: row
            for row, target in tui._screen_targets.items()
            if target[0] == "email"
        }
        self.assertIn(0, email_rows)
        tui._handle_mouse(MouseEvent(10, email_rows[0], "click", button="left"))
        self.assertEqual(tui.email_selected, 0)

    def test_sgr_mouse_input_contract(self) -> None:
        self.assertEqual(
            _parse_sgr_mouse("[<0;12;7M"),
            MouseEvent(12, 7, "click", button="left"),
        )
        self.assertEqual(
            _parse_sgr_mouse("[<64;4;9M"),
            MouseEvent(4, 9, "wheel", delta=1),
        )

    def test_windows_terminal_enables_click_and_wheel_transport(self) -> None:
        """Windows Terminal must be asked to send mouse events, not just parse them."""

        output = StringIO()
        with patch("zeus2.cli.os.name", "nt"), redirect_stdout(output):
            with _screen_mode():
                pass

        rendered = output.getvalue()
        self.assertIn("\033[?1000h", rendered)
        self.assertIn("\033[?1006h", rendered)
        self.assertIn("\033[?1000l", rendered)
        self.assertIn("\033[?1006l", rendered)

    def test_windows_mode_forces_transition_and_accepts_both_mouse_formats(self) -> None:
        original_input = (
            ENABLE_MOUSE_INPUT
            | ENABLE_QUICK_EDIT_MODE
            | ENABLE_LINE_INPUT
            | ENABLE_ECHO_INPUT
        )
        transition, interactive, output = _windows_console_mode_plan(
            original_input, 0
        )

        self.assertFalse(transition & ENABLE_MOUSE_INPUT)
        self.assertFalse(transition & ENABLE_VIRTUAL_TERMINAL_INPUT)
        self.assertTrue(transition & ENABLE_EXTENDED_FLAGS)
        self.assertTrue(interactive & ENABLE_MOUSE_INPUT)
        self.assertTrue(interactive & ENABLE_WINDOW_INPUT)
        self.assertTrue(interactive & ENABLE_VIRTUAL_TERMINAL_INPUT)
        self.assertFalse(interactive & ENABLE_QUICK_EDIT_MODE)
        self.assertFalse(interactive & ENABLE_LINE_INPUT)
        self.assertFalse(interactive & ENABLE_ECHO_INPUT)
        self.assertTrue(output & ENABLE_VIRTUAL_TERMINAL_PROCESSING)
        self.assertEqual(
            _decode_escape_sequence("[<65;4;9M"),
            MouseEvent(4, 9, "wheel", delta=-1),
        )


class AgingTests(ZeusCase):
    def test_thresholds_unplanned_no_mail_and_suspend(self) -> None:
        config = self.store.config
        facts = calculate_ticket_facts(
            report_date="2026-07-07",
            planned_date=None,
            resolve_by="2026-08-01",
            resolve_by_suspend="2026-09-01",
            status="Customer Agreed Suspend",
            last_activity_at=None,
            config=config,
            today=date(2026, 8, 6),
        )
        self.assertEqual(facts.ticket_age_days, 30)
        self.assertEqual(facts.ticket_age_color, "red")
        self.assertEqual((facts.planned_label, facts.planned_color), ("Unplanned", "amber"))
        self.assertEqual((facts.communication_label, facts.communication_color), ("No email found", "grey"))
        self.assertEqual((facts.resolve_label, facts.resolve_color), ("Suspended", None))

        due = calculate_ticket_facts(
            report_date="2026-07-23",
            planned_date="2026-08-08",
            resolve_by="2026-08-09",
            resolve_by_suspend=None,
            status="Open",
            last_activity_at="2026-08-03",
            config=config,
            today=date(2026, 8, 6),
        )
        self.assertEqual(due.ticket_age_color, "yellow")
        self.assertEqual(due.planned_color, "yellow")
        self.assertEqual(due.resolve_color, "yellow")
        self.assertEqual(due.communication_color, "yellow")

    def test_report_sort_overdue_then_unplanned_then_future(self) -> None:
        tickets = []
        for ticket_id, planned in (
            ("12345678", "2026-08-10"),
            ("12345679", None),
            ("12345680", "2026-08-01"),
        ):
            ticket = new_ticket(ticket_id, upstream_row(ticket_id), {})
            ticket["local"]["fields"]["Planned Date"] = planned
            tickets.append(ticket)
        with patch("zeus2.aging.datetime") as mocked:
            mocked.now.return_value = datetime(2026, 8, 6).astimezone()
            ordered = sorted(tickets, key=lambda ticket: report_sort_key(ticket, self.store.config))
        self.assertEqual([ticket["ticket_id"] for ticket in ordered], ["12345680", "12345679", "12345678"])


class PublicationTests(ZeusCase):
    def test_publication_appends_closure_generates_report_and_deletes_after_verify(self) -> None:
        self.seed(
            [
                pending_row("12345678", **{"Notes": "keep", "Done?": "P"}),
                pending_row("12345679"),
            ],
            [upstream_row("12345679")],
            [pending_row("11111111", **{"Notes": "historical"})],
        )
        # Style and value of the existing Closed row must survive append.
        workbook = load_workbook(self.books / "Closed.xlsx")
        workbook.active["A2"].fill = PatternFill("solid", fgColor="ABCDEF")
        workbook.save(self.books / "Closed.xlsx")
        workbook.close()
        result = publish_operational_workbooks(self.store, self.books)
        self.assertEqual(result["closed_appended"], 1)
        self.assertFalse(self.store.ticket_file("12345678").exists())
        self.assertIn("12345678", self.store.closed_index()["ticket_ids"])

        pending = load_workbook(self.books / "Pendings.xlsx", data_only=False)
        self.assertEqual(pending.sheetnames[:3], ["Pendings", "Spare Parts", "Report"])
        self.assertEqual(pending["Pendings"]["A2"].value, "12345679")
        formulas = [
            cell.coordinate
            for sheet in pending.worksheets
            for row in sheet.iter_rows()
            for cell in row
            if cell.data_type == "f"
        ]
        self.assertEqual(formulas, [])
        report = pending["Report"]
        self.assertEqual(report["A1"].value, "Zeus ticket report")
        self.assertIn("Active", [report.cell(3, col).value for col in range(1, 10)])
        pending.close()

        closed = load_workbook(self.books / "Closed.xlsx")
        self.assertEqual(closed.active["A2"].value, "11111111")
        self.assertEqual(closed.active["A2"].fill.fgColor.rgb[-6:], "ABCDEF")
        ids = {str(closed.active.cell(row, 1).value) for row in range(2, closed.active.max_row + 1)}
        self.assertEqual(ids, {"11111111", "12345678"})
        spare_ids = {
            str(closed["Spare Parts"].cell(row, 1).value)
            for row in range(2, closed["Spare Parts"].max_row + 1)
        }
        self.assertEqual(spare_ids, {"11111111", "12345678"})
        closed.close()

        # Internal snapshots can contain Outlook bodies, so final closure must
        # purge them while leaving the Excel-only paired backups available.
        self.assertEqual(self.store.list_backups(), [])
        self.assertTrue(list((self.books / "Zeus Backups").glob("Pendings_*.xlsx")))

    def test_publication_recovery_finalizes_replaced_workbooks(self) -> None:
        self.seed(
            [pending_row("12345678"), pending_row("12345679")],
            [upstream_row("12345679")],
        )
        real_finalize = __import__("zeus2.excel_export", fromlist=["_finalize_publication"])._finalize_publication
        with patch(
            "zeus2.excel_export._finalize_publication",
            side_effect=RuntimeError("simulated crash after pair replacement"),
        ):
            with self.assertRaises(RuntimeError):
                publish_operational_workbooks(self.store, self.books)
        self.assertTrue((self.store.current / "publication_journal.json").exists())
        self.assertTrue(self.store.ticket_file("12345678").exists())
        with patch("zeus2.excel_export._finalize_publication", real_finalize):
            result = recover_publication(self.store, self.books)
        self.assertEqual(result["recovered"], "finalized")
        self.assertFalse(self.store.ticket_file("12345678").exists())
        self.assertFalse((self.store.current / "publication_journal.json").exists())

    def test_pair_replacement_rolls_back_when_closed_is_locked(self) -> None:
        self.seed([pending_row("12345678")])
        publish_operational_workbooks(self.store, self.books)
        pending_before = (self.books / "Pendings.xlsx").read_bytes()
        closed_before = (self.books / "Closed.xlsx").read_bytes()
        real_replace = os.replace

        def fail_closed(source: object, destination: object) -> None:
            source_path = Path(source)
            destination_path = Path(destination)
            if (
                destination_path.resolve() == (self.books / "Closed.xlsx").resolve()
                and source_path.name.startswith(".zeus-closed-")
            ):
                raise PermissionError("simulated Excel lock")
            real_replace(source, destination)

        with patch("zeus2.excel_export.os.replace", side_effect=fail_closed):
            with self.assertRaises(PermissionError):
                publish_operational_workbooks(self.store, self.books)
        self.assertEqual((self.books / "Pendings.xlsx").read_bytes(), pending_before)
        self.assertEqual((self.books / "Closed.xlsx").read_bytes(), closed_before)
        self.assertFalse((self.store.current / "publication_journal.json").exists())

    def test_restore_imports_only_local_fields_and_rewrites_current_pendings(self) -> None:
        self.seed([pending_row("12345678", **{"Notes": "current"})])
        publish_operational_workbooks(self.store, self.books)
        backup = self.root / "old.xlsx"
        write_managed(
            backup,
            [pending_row("12345678", summary="old protected", **{"Notes": "restored", "Done?": "P"})],
        )
        result = restore_pendings_backup(self.store, backup, self.books)
        ticket = self.store.read_ticket("12345678")
        self.assertEqual(ticket["local"]["fields"]["Notes"], "restored")
        self.assertEqual(ticket["local"]["fields"]["Done?"], "P")
        self.assertNotEqual(
            ticket["upstream"]["fields"]["Problem Summary"], "old protected"
        )
        workbook = load_workbook(self.books / "Pendings.xlsx", read_only=True)
        headers = [cell.value for cell in workbook.worksheets[0][1]]
        notes_column = headers.index("Notes") + 1
        self.assertEqual(workbook.worksheets[0].cell(2, notes_column).value, "restored")
        workbook.close()
        self.assertIn("rewritten_path", result)

    def test_restore_recovery_finalizes_rewritten_pendings(self) -> None:
        self.seed([pending_row("12345678", **{"Notes": "current"})])
        publish_operational_workbooks(self.store, self.books)
        backup = self.root / "old.xlsx"
        write_managed(backup, [pending_row("12345678", **{"Notes": "restored"})])
        real_finalize = __import__("zeus2.excel_export", fromlist=["_finalize_pendings_restore"])._finalize_pendings_restore
        with patch(
            "zeus2.excel_export._finalize_pendings_restore",
            side_effect=RuntimeError("simulated crash after Pendings replacement"),
        ):
            with self.assertRaises(RuntimeError):
                restore_pendings_backup(self.store, backup, self.books)
        self.assertEqual(
            self.store.read_ticket("12345678")["local"]["fields"]["Notes"],
            "current",
        )
        with patch("zeus2.excel_export._finalize_pendings_restore", real_finalize):
            result = recover_pendings_restore(self.store, self.books)
        self.assertEqual(result["recovered"], "finalized")
        self.assertEqual(
            self.store.read_ticket("12345678")["local"]["fields"]["Notes"],
            "restored",
        )

    def test_initialized_closed_index_rejects_unmanaged_rows(self) -> None:
        self.seed(
            [pending_row("12345678")],
            closed_rows=[pending_row("11111111")],
        )
        workbook = load_workbook(self.books / "Closed.xlsx")
        workbook.active.append(
            [pending_row("22222222").get(column) for column in PENDING_COLUMNS]
        )
        workbook.save(self.books / "Closed.xlsx")
        workbook.close()
        with self.assertRaises(WorkbookValidationError):
            initialize_closed_index(self.store, self.books / "Closed.xlsx")


class StartupAndReadOnlyTests(ZeusCase):
    def test_startup_imports_pendings_then_advanced_without_publishing(self) -> None:
        write_managed(
            self.books / "Pendings.xlsx",
            [pending_row("12345678", **{"Notes": "from Excel"})],
        )
        write_managed(self.books / "Closed.xlsx", [], sheet_name="Closed")
        source = self.downloads / "Advanced Search(Service Request)20260801010101.xlsx"
        write_advanced(source, [upstream_row("12345678", summary="Online")])
        pending_hash = sha256_file(self.books / "Pendings.xlsx")
        closed_hash = sha256_file(self.books / "Closed.xlsx")
        result = run_startup(self.store)
        ticket = self.store.read_ticket("12345678")
        self.assertEqual(ticket["local"]["fields"]["Notes"], "from Excel")
        self.assertEqual(ticket["upstream"]["fields"]["Problem Summary"], "Online")
        self.assertEqual(sha256_file(self.books / "Pendings.xlsx"), pending_hash)
        self.assertEqual(sha256_file(self.books / "Closed.xlsx"), closed_hash)
        self.assertTrue(result.advanced_search_valid)

    def test_invalid_pendings_warns_but_valid_advanced_still_updates(self) -> None:
        self.seed([pending_row("12345678")])
        workbook = load_workbook(self.books / "Pendings.xlsx")
        workbook.active["N2"] = "=1+1"
        workbook.save(self.books / "Pendings.xlsx")
        workbook.close()
        newer = self.downloads / "Advanced Search(Service Request)20260802010101.xlsx"
        write_advanced(newer, [upstream_row("12345678", summary="Updated online")])
        result = run_startup(self.store)
        self.assertTrue(any("PENDINGS WARNING" in warning for warning in result.warnings))
        self.assertTrue(result.advanced_search_valid)
        self.assertEqual(
            self.store.read_ticket("12345678")["upstream"]["fields"]["Problem Summary"],
            "Updated online",
        )

    def test_transient_outlook_startup_failure_is_nonfatal_and_logged(self) -> None:
        """A cold Outlook/COM failure must not close the Zeus dashboard."""

        self.seed([pending_row("12345678")])
        mailbox = self.root / "mailbox.ost"
        mailbox.touch()
        config = self.store.config
        config["paths"]["outlook_store_path"] = str(mailbox)
        config["email"]["fetch_interval_days"] = -1
        self.store.save_config(config)

        with patch(
            "zeus2.startup.fetch_and_commit_outlook",
            side_effect=RuntimeError("RPC server unavailable during Outlook warm-up"),
        ):
            result = run_startup(self.store)

        self.assertTrue(
            any("EMAIL FETCH WARNING" in warning for warning in result.warnings)
        )
        diagnostic_log = self.home / "logs" / "zeus.log"
        self.assertTrue(diagnostic_log.is_file())
        self.assertIn("RPC server unavailable", diagnostic_log.read_text(encoding="utf-8"))

    def test_top_level_failure_reports_the_persistent_log(self) -> None:
        stderr = StringIO()
        with (
            patch("zeus2.cli.create_store", return_value=self.store),
            patch("zeus2.cli.interactive", side_effect=RuntimeError("dashboard crash")),
            redirect_stderr(stderr),
        ):
            exit_code = main([])

        self.assertEqual(exit_code, 1)
        self.assertIn("Diagnostic log:", stderr.getvalue())
        self.assertIn(
            "dashboard crash",
            (self.home / "logs" / "zeus.log").read_text(encoding="utf-8"),
        )

    def test_app_side_ticket_edits_are_forbidden(self) -> None:
        self.seed([pending_row("12345678")])
        with self.assertRaises(ReadOnlyTicketError):
            update_ticket(self.store, "12345678", workflow_code="Y")
        with self.assertRaises(ReadOnlyTicketError):
            append_note(self.store, "12345678", "note")

    def test_mop_generation_remains_an_output_not_a_ticket_editor(self) -> None:
        self.seed([pending_row("12345678", **{"Site": "GYE"})])
        template = self.root / "template.docx"
        document = Document()
        document.add_paragraph("Ticket {{ ticket_id }} at {{ site }}")
        document.save(template)
        result = generate_mop(self.store, "12345678", template)
        generated = Document(result["path"])
        self.assertEqual(generated.paragraphs[0].text, "Ticket 12345678 at GYE")


if __name__ == "__main__":
    unittest.main()
