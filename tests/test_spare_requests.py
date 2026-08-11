from __future__ import annotations

import shutil
import tempfile
import unittest
from copy import deepcopy
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook

from zeus2.application.errors import ValidationError
from zeus2.application.serialization import (
    COLUMN_DEFINITIONS,
    SPARE_REQUEST_COLUMN_DEFINITIONS,
    dashboard_payload,
    serialize_spare_request_detail,
    spare_requests_dashboard_payload,
    ticket_revision,
)
from zeus2.application.service import ApplicationService
from zeus2.config import load_config, save_config
from zeus2.fault_tags import fault_tag_status
from zeus2.spare_request_excel import (
    ARCHIVE_HEADERS,
    FAULTY_TAG_SHEET,
    LEGACY_ARCHIVE_HEADERS,
    REQUEST_SHEET,
    RETURN_SHEET,
    SPARE_ARCHIVE_SHEET,
    export_initial_request,
    export_return_workbook,
    append_archived_item,
    read_archived_items,
)
from zeus2.spare_request_mail import (
    _request_id_from_message,
    apply_spare_request_messages,
    parse_dispatch_notification,
)
from zeus2.spare_requests import (
    ECUADOR_TIMEZONE,
    SpareRequestError,
    create_request_record,
    item_status,
    lifecycle_stage,
    normalize_profile,
    normalize_request_lines,
    request_filename,
    request_subject,
)
from zeus2.store import ZeusStore


def profile_input() -> dict:
    return {
        "clientInitials": "CNT",
        "customerName": "Customer Network Team",
        "siteCode": "UIO1",
        "siteName": "Quito",
        "siteAddress": "Av. Example 123",
        "cloud": "Ecuador Cloud",
        "requester": {
            "name": "Zeus User",
            "email": "user@example.com",
            "phone": "+593000000000",
        },
        "contact": {
            "name": "Customer Contact",
            "email": "contact@example.com",
            "phone": "+593111111111",
        },
    }


def lines_input(amount: int = 2) -> list[dict]:
    return [
        {
            "bom": "02312RCC",
            "amount": amount,
            "description": "Controller board",
            "part": "Controller board",
            "model": "S6730",
            "device": "SW-UIO-01",
            "slot": "",
            "faultySn": "FAULTY-1",
            "reportDate": "2026-07-01",
        }
    ]


def request_record(amount: int = 2) -> dict:
    profile = normalize_profile(profile_input())
    lines = normalize_request_lines(lines_input(amount))
    request_id = "260808123456"
    return create_request_record(
        request_id=request_id,
        tt="39416095",
        source="manual",
        profile=profile,
        lines=lines,
        export_path=None,
        subject=request_subject(request_id, "39416095", lines),
        created_at="2026-08-08T12:34:56-05:00",
    )


def source_ticket() -> dict:
    return {
        "schema_version": 2,
        "ticket_id": "39416095",
        "lifecycle": {"status": "active"},
        "upstream": {
            "fields": {
                "SRNo": "39416095",
                "Report Date": "2026-07-01 10:00:00",
                "Customer Severity": "Minor",
                "Problem Summary": "Controller board failure",
            }
        },
        "local": {
            "fields": {"Done?": "N", "Site": "UIO1", "Cloud": "Ecuador Cloud"},
            "spare_parts": [
                {
                    "device": "SW-UIO-01",
                    "model": "S6730",
                    "faulty_sns": ["FAULTY-1"],
                    "parts": [
                        {
                            "slot": "1/0/1",
                            "part": "Controller board",
                            "bom": "02312RCC",
                            "notes": None,
                            "new_sn": None,
                        }
                    ],
                }
            ],
        },
        "email": {
            "total_received": 0,
            "total_sent": 0,
            "last_activity_at": None,
            "messages": [],
        },
        "mop": {"latest": None, "versions": 0},
        "updated_at": None,
    }


def create_request_template(path: Path) -> None:
    workbook = Workbook()
    request = workbook.active
    request.title = REQUEST_SHEET
    for row in range(1, 42):
        for column in range(1, 14):
            request.cell(row, column).value = None
    faulty = workbook.create_sheet(FAULTY_TAG_SHEET)
    for row in range(1, 32):
        for column in range(1, 14):
            faulty.cell(row, column).value = None
    workbook.save(path)
    workbook.close()


def create_return_template(path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = RETURN_SHEET
    for row in range(1, 16):
        for column in range(1, 10):
            worksheet.cell(row, column).value = None
    workbook.save(path)
    workbook.close()


class SpareRequestDomainTests(unittest.TestCase):
    def test_email_and_spare_summary_columns_match_their_badge_capacity(self) -> None:
        service_columns = {column["key"]: column for column in COLUMN_DEFINITIONS}
        spare_columns = {
            column["key"]: column for column in SPARE_REQUEST_COLUMN_DEFINITIONS
        }

        self.assertEqual(service_columns["emailLabel"]["width"], 140)
        self.assertEqual(service_columns["spareBadges"]["width"], 116)
        self.assertEqual(spare_columns["emailLabel"]["width"], 140)

    def test_detail_serialization_supplies_collections_before_database_upgrade(self) -> None:
        request = request_record(1)
        request["creation_method"] = "zeus_create"
        request["items"][0].pop("rma_aliases")
        request["email"].pop("messages")

        detail = serialize_spare_request_detail(request)

        self.assertEqual(detail["items"][0]["rma_aliases"], [])
        self.assertEqual(detail["email"]["messages"], [])
        self.assertEqual(detail["items"][0]["lifecycle"]["stage"], 0)
        self.assertEqual(detail["creationMethod"], "zeus_create")

    def test_legacy_archive_adds_the_rma_alias_column_without_shifting_data(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            closed = Path(directory) / "Closed.xlsx"
            workbook = Workbook()
            archive = workbook.active
            archive.title = SPARE_ARCHIVE_SHEET
            archive.append(LEGACY_ARCHIVE_HEADERS)
            legacy = {header: None for header in LEGACY_ARCHIVE_HEADERS}
            legacy.update({
                "TT": "39416095",
                "RMA": "C3209937001",
                "Email inactivity days": 17,
                "Item ID": "legacy-0001",
            })
            archive.append([legacy[header] for header in LEGACY_ARCHIVE_HEADERS])
            workbook.save(closed)
            workbook.close()

            request = request_record(1)
            item = request["items"][0]
            item["rma"] = "C3209937002"
            item["rma_aliases"] = ["C3209937000"]
            append_archived_item(
                closed, request, item, reason="returned", note="Returned"
            )

            workbook = load_workbook(closed, read_only=True, data_only=True)
            try:
                archive = workbook[SPARE_ARCHIVE_SHEET]
                self.assertEqual([cell.value for cell in archive[1]], ARCHIVE_HEADERS)
                self.assertEqual(archive.cell(2, ARCHIVE_HEADERS.index("RMA") + 1).value, "C3209937001")
                self.assertIsNone(archive.cell(2, ARCHIVE_HEADERS.index("RMA Aliases") + 1).value)
                self.assertEqual(archive.cell(2, ARCHIVE_HEADERS.index("Email inactivity days") + 1).value, 17)
            finally:
                workbook.close()
            rows = read_archived_items(closed)
            self.assertEqual(rows[1]["RMA Aliases"], "C3209937000")

    def test_customer_email_and_phone_are_required_for_new_requests(self) -> None:
        for missing in ("email", "phone"):
            profile = profile_input()
            profile["contact"][missing] = ""
            with self.subTest(missing=missing), self.assertRaisesRegex(
                SpareRequestError,
                f"Customer contact {missing}",
            ):
                normalize_profile(profile)

    def test_request_quantity_rejects_non_whole_json_values(self) -> None:
        for amount in (True, 1.5, "1.5", "1e2"):
            with self.subTest(amount=amount):
                with self.assertRaisesRegex(SpareRequestError, "whole number"):
                    normalize_request_lines(lines_input(amount))

    def test_one_physical_unit_keeps_all_component_serials_in_one_fault_tag(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            template = root / "request.xlsx"
            create_request_template(template)
            raw_profile = profile_input()
            raw_profile.pop("clientInitials")
            raw_profile["contact"] = {
                "name": "Juan Piguave",
                "email": "juan@example.com",
                "phone": "+593111111111",
            }
            profile = normalize_profile(raw_profile)
            lines = normalize_request_lines(
                [
                    {
                        **lines_input(amount=1)[0],
                        "bom": "SERVER-001",
                        "description": "Complete replacement server",
                        "faultySn": "CPU-SN-001\nMEMORY,SN,002\nMEZZ-SN-003",
                    }
                ]
            )
            request = create_request_record(
                request_id="260808123456",
                tt="39416095",
                source="manual",
                profile=profile,
                lines=lines,
                export_path=None,
                subject=request_subject("260808123456", "39416095", lines),
                created_at="2026-08-08T12:34:56-05:00",
            )

            self.assertEqual(profile["client_initials"], "JP")
            self.assertEqual(len(request["items"]), 1)
            self.assertEqual(
                request["items"][0]["faulty_sns"],
                ["CPU-SN-001", "MEMORY,SN,002", "MEZZ-SN-003"],
            )
            filename = request_filename(
                request["request_id"], request["tt"], profile, lines
            )
            self.assertIn("SP–JP–", filename)
            exported = export_initial_request(template, root / "exports", request, filename)
            workbook = load_workbook(exported, read_only=True, data_only=False)
            try:
                application = workbook[REQUEST_SHEET]
                faulty = workbook[FAULTY_TAG_SHEET]
                self.assertEqual(application["B15"].value, "SERVER-001")
                self.assertEqual(application["C15"].value, 1)
                self.assertEqual(
                    application["F15"].value,
                    "CPU-SN-001\nMEMORY,SN,002\nMEZZ-SN-003",
                )
                self.assertEqual(application["G2"].value, "Customer Network Team")
                self.assertEqual(application["F31"].value, "Juan Piguave")
                self.assertEqual(
                    [faulty[f"F{row}"].value for row in (12, 14, 16)],
                    ["CPU-SN-001\nMEMORY,SN,002\nMEZZ-SN-003", None, None],
                )
                self.assertEqual(
                    [faulty[f"B{row}"].value for row in (12, 14, 16)],
                    ["SERVER-001", None, None],
                )
                self.assertTrue(faulty["F12"].alignment.wrap_text)
            finally:
                workbook.close()

    def test_export_identity_subject_and_quantity_expansion(self) -> None:
        request = request_record(3)
        filename = request_filename(
            request["request_id"], request["tt"], request["profile"], request["request_lines"]
        )
        self.assertEqual(
            filename,
            "SP–CNT–UIO1–Ecuador Cloud–3x02312RCC–39416095–260808123456.xlsx",
        )
        self.assertEqual(len(request["items"]), 3)
        self.assertEqual(
            [item["item_id"] for item in request["items"]],
            ["260808123456-0001", "260808123456-0002", "260808123456-0003"],
        )
        self.assertEqual(
            request["export"]["subject"],
            "[TT 39416095] [SPARE PARTS REQUEST] 260808123456 Controller board for S6730 SW-UIO-01",
        )

    def test_newline_slots_derive_quantity_and_assign_one_slot_per_item(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            template = root / "request.xlsx"
            create_request_template(template)
            lines = normalize_request_lines(
                [{
                    **lines_input(amount=99)[0],
                    "part": "DIMM",
                    "description": "DIMM",
                    "slot": "DIMM101\nDIMM203\nDIMM103\nDIMM101",
                    "faultySns": "SERVER-SN-01\nMEMORY-SN-02",
                    "notes": "Memory diagnostics completed before this request.",
                }]
            )
            self.assertEqual(lines[0]["amount"], 3)
            self.assertEqual(lines[0]["slots"], ["DIMM101", "DIMM203", "DIMM103"])
            request = create_request_record(
                request_id="260808123456",
                tt="39416095",
                source="manual",
                profile=normalize_profile(profile_input()),
                lines=lines,
                export_path=None,
                subject=request_subject("260808123456", "39416095", lines),
                created_at="2026-08-08T12:34:56-05:00",
            )
            self.assertEqual(len(request["items"]), 3)
            self.assertEqual(
                [item["slot"] for item in request["items"]],
                ["DIMM101", "DIMM203", "DIMM103"],
            )
            self.assertEqual(
                {item["requested_bom"] for item in request["items"]},
                {"02312RCC"},
            )
            self.assertTrue(all(item["notes"] == lines[0]["notes"] for item in request["items"]))
            self.assertTrue(
                all(
                    item["faulty_sns"] == ["SERVER-SN-01", "MEMORY-SN-02"]
                    for item in request["items"]
                )
            )
            exported = export_initial_request(
                template,
                root / "exports",
                request,
                request_filename(request["request_id"], request["tt"], request["profile"], lines),
            )
            workbook = load_workbook(exported, read_only=True)
            try:
                self.assertEqual(workbook.worksheets[0]["C15"].value, 3)
                self.assertEqual(
                    [workbook.worksheets[1][f"B{row}"].value for row in (12, 14, 16)],
                    ["02312RCC"] * 3,
                )
                descriptions = [
                    workbook.worksheets[1][f"D{row}"].value for row in (12, 14, 16)
                ]
                for slot, description in zip(lines[0]["slots"], descriptions, strict=True):
                    self.assertIn(slot, description)
            finally:
                workbook.close()

    def test_export_preserves_neutral_template_sheet_names(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            request_template = root / "request.xlsx"
            return_template = root / "return.xlsx"
            create_request_template(request_template)
            create_return_template(return_template)
            request_book = load_workbook(request_template)
            request_book.worksheets[0].title = "Local Request Form"
            request_book.worksheets[1].title = "Local Fault Evidence"
            request_book.save(request_template)
            request_book.close()
            return_book = load_workbook(return_template)
            return_book.worksheets[0].title = "Local Return Form"
            return_book.save(return_template)
            return_book.close()

            request = request_record(1)
            exported = export_initial_request(
                request_template,
                root / "exports",
                request,
                request_filename(request["request_id"], request["tt"], request["profile"], request["request_lines"]),
            )
            check = load_workbook(exported, read_only=True)
            try:
                self.assertEqual(check.sheetnames, ["Local Request Form", "Local Fault Evidence"])
            finally:
                check.close()

            request["spare_sr"] = "SR4956964"
            request["items"][0].update(
                {
                    "rma": "C3209937821",
                    "delivered_bom": "02540255",
                    "new_sn": "NEW-1",
                }
            )
            returned_path = export_return_workbook(
                return_template,
                root / "exports",
                [(request, request["items"][0], "Faulty")],
            )["path"]
            returned = load_workbook(returned_path, read_only=True)
            try:
                self.assertEqual(returned.sheetnames, ["Local Return Form"])
                self.assertEqual(returned.worksheets[0]["I8"].value, "C3209937821")
            finally:
                returned.close()

    def test_request_and_return_exports_preserve_template_sheet_contracts(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            request_template = root / "request.xlsx"
            return_template = root / "return.xlsx"
            create_request_template(request_template)
            create_return_template(return_template)
            request = request_record(7)
            filename = request_filename(
                request["request_id"], request["tt"], request["profile"], request["request_lines"]
            )
            exported = export_initial_request(request_template, root / "exports", request, filename)
            workbook = load_workbook(exported, read_only=True)
            try:
                self.assertEqual(workbook.sheetnames, [REQUEST_SHEET, FAULTY_TAG_SHEET])
                self.assertEqual(workbook[FAULTY_TAG_SHEET]["B24"].value, "02312RCC")
                self.assertEqual(
                    [workbook[FAULTY_TAG_SHEET][f"F{row}"].value for row in range(12, 26, 2)],
                    ["FAULTY-1"] * 7,
                )
            finally:
                workbook.close()

            request["spare_sr"] = "SR4956964"
            for index, item in enumerate(request["items"][:2], start=1):
                item["rma"] = f"C320993782{index}"
                item["delivered_bom"] = "02540255"
                item["new_sn"] = f"NEW-{index}"
            result = export_return_workbook(
                return_template,
                root / "exports",
                [
                    (request, request["items"][0], "Faulty"),
                    (request, request["items"][1], "New"),
                ],
            )
            returned = load_workbook(result["path"], read_only=True)
            try:
                self.assertEqual(returned.sheetnames, [RETURN_SHEET])
                self.assertIsNone(returned[RETURN_SHEET]["C8"].value)
                self.assertEqual(returned[RETURN_SHEET]["C9"].value, "NEW-2")
            finally:
                returned.close()


class SpareRequestMailTests(unittest.TestCase):
    def test_previous_rma_alias_still_matches_dispatch_email(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            config = load_config(root / "home")
            config["email"]["dispatch_notification_sender"] = "dispatch@example.test"
            save_config(root / "home", config)
            store = ZeusStore(root / "data", config_home=root / "home")
            record = request_record(1)
            record["spare_sr"] = "SR4956964"
            record["request_sent_at"] = "2026-08-06T08:00:00-05:00"
            record["request_sent_source"] = "manual"
            item = record["items"][0]
            item["rma"] = "C3209937827"
            item["rma_aliases"] = ["C3209937826"]
            item["attendance_confirmed_at"] = "2026-08-07T08:00:00-05:00"
            item["attendance_source"] = "manual"
            with store.transaction("seed-alias", {}) as staging:
                store.write_spare_request(staging, record)

            with store.transaction("dispatch-by-alias", {}) as staging:
                apply_spare_request_messages(store, staging, [{
                    "message_key": "dispatch-alias",
                    "timestamp": "2026-08-08T08:00:00-05:00",
                    "direction": "received",
                    "sender": "Dispatch Service",
                    "sender_address": "dispatch@example.test",
                    "subject": "Dispatch delivery",
                    "body": "",
                    "html_body": "<table><tr><th>Order No.</th><th>Line No.</th><th>Item</th><th>QTY</th></tr><tr><td>C3209937826</td><td>10</td><td>02540255</td><td>1</td></tr></table>",
                }])
            updated = store.read_spare_request(record["request_id"])["items"][0]
            self.assertEqual(updated["rma"], "C3209937827")
            self.assertEqual(updated["rma_aliases"], ["C3209937826"])
            self.assertEqual(updated["dispatch_message_key"], "dispatch-alias")

    def test_fault_tag_timestamp_never_masquerades_as_a_request_id(self) -> None:
        self.assertIsNone(
            _request_id_from_message(
                {"subject": "[FAULT TAG FT-260810120000] [UIO1] RMA C3209937826"}
            )
        )
        self.assertEqual(
            _request_id_from_message(
                {
                    "subject": "Request 260810115959 with Fault Tag FT-260810120000"
                }
            ),
            "260810115959",
        )

    def test_trusted_emails_advance_request_fault_tag_and_warehouse_stages(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            config = load_config(root / "home")
            config["email"].update(
                {
                    "request_confirmation_sender": "request-confirmation@example.test",
                    "dispatch_notification_sender": "dispatch@example.test",
                    "warehouse_sender_domain": "@warehouse.example.test",
                }
            )
            save_config(root / "home", config)
            store = ZeusStore(root / "data", config_home=root / "home")
            store.ensure_layout()
            seeded = request_record(1)
            seeded["creation_method"] = "zeus_export"
            with store.transaction("seed", {}) as staging:
                store.write_spare_request(staging, seeded)
            messages = [
                {
                    "message_key": "request-sent",
                    "timestamp": "2026-08-06T08:00:00-05:00",
                    "direction": "sent",
                    "sender": "Zeus User",
                    "sender_address": "user@example.com",
                    "subject": "[TT 39416095] [SPARE PARTS REQUEST] 260808123456",
                    "body": "Request attached",
                    "html_body": "",
                },
                {
                    "message_key": "confirmation",
                    "timestamp": "2026-08-07T08:00:00-05:00",
                    "direction": "received",
                    "sender": "Request Service",
                    "sender_address": "request-confirmation@example.test",
                    "subject": "TT 39416095 request 260808123456",
                    "body": "TT: 39416095",
                    "html_body": "<table><tr><th>SR</th><th>RMA</th><th>ITEM</th></tr><tr><td>SR4956964</td><td>C3209937826</td><td>02312RCC</td></tr></table>",
                },
                {
                    "message_key": "dispatch",
                    "timestamp": "2026-08-08T08:00:00-05:00",
                    "direction": "received",
                    "sender": "Dispatch Service",
                    "sender_address": "dispatch@example.test",
                    "subject": "Dispatch delivery",
                    "body": "",
                    "html_body": "<table><tr><th>Order No.</th><th>Line No.</th><th>Item</th><th>QTY</th></tr><tr><td>C3209937826</td><td>10</td><td>02540255</td><td>1</td></tr></table>",
                },
            ]
            with store.transaction("mail-stages-1-3", {}) as staging:
                apply_spare_request_messages(
                    store,
                    staging,
                    messages,
                    now=datetime(2026, 8, 8, 12, 0, tzinfo=ECUADOR_TIMEZONE),
                )
            request = store.read_spare_request("260808123456")
            item = request["items"][0]
            self.assertEqual(request["request_sent_source"], "email")
            self.assertEqual(lifecycle_stage(item, request), 3)

            with store.transaction("fault-tag-generated", {}) as staging:
                request = store.read_spare_request("260808123456", staging)
                request["items"][0]["fault_tag_generated_at"] = "2026-08-09T08:00:00-05:00"
                request["items"][0]["fault_tag_generated_source"] = "zeus_export"
                store.write_spare_request(staging, request)
            later_messages = [
                {
                    "message_key": "fault-tag-sent",
                    "timestamp": "2026-08-09T09:00:00-05:00",
                    "direction": "sent",
                    "sender": "Zeus User",
                    "sender_address": "user@example.com",
                    "subject": "[FAULT TAG] RMA C3209937826",
                    "body": "Fault Tag attached for C3209937826",
                    "html_body": "",
                },
                {
                    "message_key": "warehouse-confirmed",
                    "timestamp": "2026-08-10T09:00:00-05:00",
                    "direction": "received",
                    "sender": "Warehouse",
                    "sender_address": "agent@warehouse.example.test",
                    "subject": "SR4956964 C3209937826 RT12345678",
                    "body": "SR4956964 C3209937826 RT12345678",
                    "html_body": "",
                },
            ]
            with store.transaction("mail-stages-5-6", {}) as staging:
                result = apply_spare_request_messages(
                    store,
                    staging,
                    later_messages,
                    now=datetime(2026, 8, 10, 12, 0, tzinfo=ECUADOR_TIMEZONE),
                )
            self.assertEqual(result["unmatched_messages"], 0)
            request = store.read_spare_request("260808123456")
            item = request["items"][0]
            self.assertIsNone(item["fault_tag_sent_source"])
            self.assertEqual(item["warehouse_confirmation_source"], "email")
            self.assertEqual(item["rt"], "RT12345678")
            self.assertEqual(lifecycle_stage(item, request), 5)
            self.assertEqual(item_status(item, request), "awaiting_user_confirmation")

    def test_out_of_order_dispatch_and_partial_stock_resolve_per_unit(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            config = load_config(root / "home")
            config["email"].update(
                {
                    "request_confirmation_sender": "request-confirmation@example.test",
                    "dispatch_notification_sender": "dispatch@example.test",
                    "warehouse_sender_domain": "@warehouse.example.test",
                }
            )
            save_config(root / "home", config)
            store = ZeusStore(root / "data", config_home=root / "home")
            store.ensure_layout()
            with store.transaction("seed", {}) as staging:
                store.write_spare_request(staging, request_record(3))
            dispatch = {
                "message_key": "dispatch",
                "timestamp": "2026-08-07T09:00:00-05:00",
                "direction": "received",
                "sender": "Dispatch Service",
                "sender_address": "dispatch@example.test",
                "subject": "Spare Request SR4956964 TT 39416095",
                "body": "",
                "html_body": "<table><tr><th>Order No.</th><th>Line No.</th><th>Item</th><th>QTY</th><th>SN</th></tr><tr><td>C3209937826</td><td>10</td><td>02540255</td><td>1</td><td>NEW-1</td></tr></table>",
            }
            confirmation = {
                "message_key": "confirmation",
                "timestamp": "2026-08-08T10:00:00-05:00",
                "direction": "received",
                "sender": "Request Service",
                "sender_address": "request-confirmation@example.test",
                "subject": "TT 39416095 request 260808123456",
                "body": "TT: 39416095",
                "html_body": "<table><tr><th>SR</th><th>RMA</th><th>ITEM</th><th>ITEM Description</th></tr><tr><td>SR4956964</td><td>C3209937826</td><td>02312RCC</td><td>Controller board</td></tr><tr><td>SR4956964</td><td>C3209937827</td><td>02312RCC</td><td>Controller board</td></tr></table>",
            }
            with store.transaction("mail", {}) as staging:
                result = apply_spare_request_messages(
                    store,
                    staging,
                    [dispatch, confirmation],
                    now=datetime(2026, 8, 8, 12, 0, tzinfo=ECUADOR_TIMEZONE),
                )
            self.assertEqual(result["unmatched_messages"], 0)
            request = store.read_spare_request("260808123456")
            self.assertEqual(request["spare_sr"], "SR4956964")
            self.assertEqual(request["items"][0]["rma"], "C3209937826")
            self.assertEqual(request["items"][0]["delivered_bom"], "02540255")
            self.assertEqual(request["items"][0]["new_sn"], "NEW-1")
            self.assertEqual(item_status(request["items"][0], request), "dispatched")
            self.assertEqual(item_status(request["items"][1], request), "awaiting_dispatch")
            self.assertEqual(item_status(request["items"][2], request), "awaiting_stock")

    def test_dispatch_plain_text_fallback_uses_order_number_as_rma(self) -> None:
        fact = parse_dispatch_notification(
            {
                "subject": "Dispatch delivery",
                "sender_address": "dispatch@example.test",
                "body": "Order No.  Line No.  Item  QTY  SN\nC3209937826  10  02540255  1  NEW-1",
                "html_body": "",
            },
            "dispatch@example.test",
        )
        self.assertIsNotNone(fact)
        assert fact is not None
        self.assertEqual(fact["assignments"][0]["rma"], "C3209937826")
        self.assertEqual(fact["assignments"][0]["line_number"], "10")
        self.assertEqual(fact["assignments"][0]["delivered_bom"], "02540255")
        self.assertIsNone(
            parse_dispatch_notification(
                {
                    "sender_address": "attacker@example.com",
                    "body": "C3209937826  10  02540255  1  NEW-1",
                },
                "dispatch@example.test",
            )
        )


class SpareRequestApplicationTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.home = self.root / "home"
        self.data = self.root / "data"
        self.exports = self.root / "exports"
        self.workbooks = self.root / "workbooks"
        self.workbooks.mkdir()
        self.request_template = self.root / "request.xlsx"
        self.return_template = self.root / "return.xlsx"
        create_request_template(self.request_template)
        create_return_template(self.return_template)
        config = load_config(self.home)
        config["paths"].update(
            {
                "spare_parts_export_directory": str(self.exports),
                "spare_request_template_path": str(self.request_template),
                "spare_return_template_path": str(self.return_template),
                "workbook_directory": str(self.workbooks),
            }
        )
        save_config(self.home, config)
        self.store = ZeusStore(self.data, config_home=self.home)
        self.service = ApplicationService(self.store)

    def tearDown(self) -> None:
        self.service.stop()
        self.temporary.cleanup()

    def detect_request_sent(
        self, request: dict, *, message_key: str | None = None
    ) -> None:
        request_id = str(request.get("requestId") or request.get("request_id") or "")
        message = {
            "message_key": message_key or f"request-email-{request_id}",
            "timestamp": "2026-08-08T13:00:00-05:00",
            "direction": "sent",
            "sender": "Zeus User",
            "sender_address": "user@example.com",
            "subject": f"[TT 39416095] [SPARE PARTS REQUEST] {request_id}",
            "body": "Request attached",
            "html_body": "",
        }
        with self.store.transaction("sent-request-email", {}) as staging:
            apply_spare_request_messages(self.store, staging, [message])

    def request_at_spare_replaced(self, amount: int = 1) -> dict:
        request = self.service.export_spare_request(
            {
                "source": "manual",
                "ticketId": "39416095",
                "reportDate": "2026-07-01",
                "profile": profile_input(),
                "lines": lines_input(amount),
            }
        )["request"]
        item_ids = [item["item_id"] for item in request["items"]]
        self.detect_request_sent(request)
        request = self.service.spare_request(request["requestId"])
        request = self.service.edit_spare_request(
            request["requestId"],
            expected_revision=request["revision"],
            changes={"spareSr": "SR4956964", "note": "Confirmed for test"},
            item_updates=[
                {
                    "itemId": item_id,
                    "rma": f"C32099378{26 + index:02d}",
                    "newSn": f"NEW-{index + 1}",
                }
                for index, item_id in enumerate(item_ids)
            ],
        )["request"]
        for _ in range(3):
            self.service.bulk_spare_lifecycle(item_ids=item_ids, action="advance")
        request = self.service.spare_request(request["requestId"])
        self.assertTrue(
            all(item["lifecycle"]["stage"] == 4 for item in request["items"])
        )
        return request

    def test_manual_request_email_confirmation_is_shared_and_idempotent(self) -> None:
        request = self.service.export_spare_request(
            {
                "source": "manual",
                "ticketId": "39416095",
                "profile": profile_input(),
                "lines": lines_input(2),
            }
        )["request"]
        item_ids = [item["item_id"] for item in request["items"]]

        dashboard = spare_requests_dashboard_payload(self.store, view="active")
        self.assertTrue(
            all(
                row["canAdvance"] for row in dashboard["spareRequests"]
            )
        )
        confirmation = datetime.now(ECUADOR_TIMEZONE) - timedelta(minutes=5)
        result = self.service.bulk_spare_lifecycle(
            item_ids=[item_ids[0]],
            action="advance",
            confirmed_at=confirmation.isoformat(),
        )
        self.assertEqual(set(result["items"]), set(item_ids))
        request = self.service.spare_request(request["requestId"])
        self.assertEqual([item["lifecycle"]["stage"] for item in request["items"]], [1, 1])
        self.assertEqual(request["requestSentAt"], confirmation.isoformat(timespec="seconds"))
        self.assertEqual(request["history"][-1]["summary"]["items"], sorted(item_ids))

        self.detect_request_sent(request)
        request = self.service.spare_request(request["requestId"])
        self.assertEqual([item["lifecycle"]["stage"] for item in request["items"]], [1, 1])
        self.assertEqual(request["email"]["total_sent"], 1)
        persisted = self.store.read_spare_request(request["requestId"])
        self.assertEqual(persisted["request_sent_source"], "manual")
        self.assertEqual(
            persisted["request_sent_at"], confirmation.isoformat(timespec="seconds")
        )
        self.assertTrue(persisted["request_sent_message_key"])
        self.assertTrue(
            spare_requests_dashboard_payload(self.store, view="active")[
                "spareRequests"
            ][0]["rollbackRequiresDoubleConfirmation"]
        )
        self.assertTrue(request["items"][0]["rollbackRequiresDoubleConfirmation"])
        request = self.service.edit_spare_request(
            request["requestId"],
            expected_revision=request["revision"],
            changes={"spareSr": "SR4956964"},
            item_updates=[{"itemId": item_ids[0], "rma": "C3209937826"}],
        )["request"]
        self.service.bulk_spare_lifecycle(
            item_ids=[item_ids[0]], action="advance"
        )
        with self.assertRaisesRegex(ValidationError, "later-stage item"):
            self.service.bulk_spare_lifecycle(item_ids=item_ids, action="rollback")
        self.service.bulk_spare_lifecycle(
            item_ids=[item_ids[0]], action="rollback"
        )
        self.service.bulk_spare_lifecycle(
            item_ids=[item_ids[1]],
            action="rollback",
            email_override_confirmed=True,
            note="The shared outbound email was linked to the wrong request.",
        )
        request = self.service.spare_request(request["requestId"])
        self.assertEqual([item["lifecycle"]["stage"] for item in request["items"]], [0, 0])

    def test_service_dashboard_spare_badges_count_units_at_day_twenty(self) -> None:
        self.store.write_ticket_bundle(self.store.current, source_ticket())
        now = datetime.now(ECUADOR_TIMEZONE)

        def active_request(request_id: str, stage: int, age_days: int | None = None) -> dict:
            record = request_record(1)
            record["request_id"] = request_id
            record["spare_sr"] = f"SR{int(request_id[-7:]):07d}"
            item = record["items"][0]
            item["item_id"] = f"{request_id}-0001"
            item["rma"] = f"C{int(request_id[-10:]):010d}"
            if stage >= 1:
                record["request_sent_at"] = now.isoformat(timespec="seconds")
                record["request_sent_source"] = "manual"
            if stage >= 2:
                item["attendance_confirmed_at"] = now.isoformat(timespec="seconds")
                item["attendance_source"] = "manual"
            if stage >= 3:
                dispatched = now - timedelta(days=age_days or 0)
                item["dispatch_at"] = dispatched.isoformat(timespec="seconds")
                item["dispatch_source"] = "manual"
            if stage >= 4:
                item["replacement_confirmed_at"] = now.isoformat(timespec="seconds")
                item["replacement_confirmation_source"] = "manual"
            if stage >= 5:
                item["warehouse_candidate_at"] = now.isoformat(timespec="seconds")
                item["warehouse_confirmation_source"] = "email"
            return record

        records = [
            active_request("260810120001", 0),
            active_request("260810120002", 3, 19),
            active_request("260810120003", 3, 20),
            active_request("260810120004", 5, 22),
        ]
        with self.store.transaction("badge-fixtures", {}) as staging:
            for record in records:
                self.store.write_spare_request(staging, record)

        returned_archive = deepcopy(records[3])
        returned_archive["request_id"] = "260810120005"
        returned_archive["items"][0]["item_id"] = "260810120005-0001"
        returned_archive["items"][0]["rma"] = "C3209937995"
        append_archived_item(
            self.workbooks / "Closed.xlsx",
            returned_archive,
            returned_archive["items"][0],
            reason="returned",
            note="Returned",
        )
        cancelled_archive = deepcopy(returned_archive)
        cancelled_archive["request_id"] = "260810120006"
        cancelled_archive["items"][0]["item_id"] = "260810120006-0001"
        cancelled_archive["items"][0]["rma"] = "C3209937996"
        append_archived_item(
            self.workbooks / "Closed.xlsx",
            cancelled_archive,
            cancelled_archive["items"][0],
            reason="cancelled",
            note="Cancelled",
        )

        badges = dashboard_payload(self.store)["tickets"][0]["spareBadges"]
        self.assertEqual(
            badges,
            {"pendingDispatch": 2, "dispatched": 1, "overdue": 1, "returned": 2},
        )

    def test_email_backed_rollback_suppresses_the_same_message_on_resync(self) -> None:
        request = self.service.export_spare_request(
            {
                "source": "manual",
                "ticketId": "39416095",
                "profile": profile_input(),
                "lines": lines_input(1),
            }
        )["request"]
        message = {
            "message_key": "request-email-sent",
            "timestamp": "2026-08-08T13:00:00-05:00",
            "direction": "sent",
            "sender": "Zeus User",
            "sender_address": "user@example.com",
            "subject": f"[TT 39416095] [SPARE PARTS REQUEST] {request['requestId']}",
            "body": "Request attached",
            "html_body": "",
        }
        with self.store.transaction("sent-request-email", {}) as staging:
            apply_spare_request_messages(self.store, staging, [message])
        request = self.service.spare_request(request["requestId"])
        item_id = request["items"][0]["item_id"]
        self.assertEqual(request["items"][0]["lifecycle"]["stage"], 1)
        self.assertEqual(request["email"]["total_sent"], 1)

        with self.assertRaisesRegex(ValidationError, "second confirmation"):
            self.service.bulk_spare_lifecycle(
                item_ids=[item_id], action="rollback"
            )
        self.service.bulk_spare_lifecycle(
            item_ids=[item_id],
            action="rollback",
            email_override_confirmed=True,
            note="The outbound email used the wrong request attachment.",
        )
        rolled_back = self.store.read_spare_request(request["requestId"])
        self.assertEqual(lifecycle_stage(rolled_back["items"][0], rolled_back), 0)
        self.assertEqual(rolled_back["email"]["total_sent"], 1)
        self.assertEqual(
            rolled_back["lifecycle_suppressions"][0]["message_key"],
            "request-email-sent",
        )

        with self.store.transaction("resync-same-request-email", {}) as staging:
            apply_spare_request_messages(self.store, staging, [message])
        resynced = self.store.read_spare_request(request["requestId"])
        self.assertEqual(lifecycle_stage(resynced["items"][0], resynced), 0)
        self.assertEqual(resynced["email"]["total_sent"], 1)
        self.assertEqual(len(resynced["email"]["messages"]), 1)

    def test_fault_tag_locks_reexports_and_can_be_deleted_without_lifecycle_change(self) -> None:
        request = self.request_at_spare_replaced()
        item_id = request["items"][0]["item_id"]
        exported = self.service.export_spare_return(
            [{"itemId": item_id, "condition": "Faulty"}]
        )
        fault_tag_id = exported["faultTagId"]
        self.assertRegex(fault_tag_id, r"^FT-\d{12}$")
        self.assertEqual(
            self.service.spare_request(request["requestId"])["items"][0]["lifecycle"]["stage"],
            4,
        )

        current = self.service.spare_request(request["requestId"])
        corrected = self.service.edit_spare_request(
            request["requestId"],
            expected_revision=current["revision"],
            changes={"note": "Supplier corrected the RMA before the tag was sent."},
            item_updates=[{"itemId": item_id, "rma": "C3209937999"}],
        )["request"]
        self.assertEqual(corrected["items"][0]["rma_aliases"], ["C3209937826"])
        unlocked_tag = self.store.read_fault_tag(fault_tag_id, completed=False)
        self.assertEqual(unlocked_tag["members"][0]["rma"], "C3209937999")
        self.assertTrue(unlocked_tag["export"]["needs_reexport"])

        revised = self.service.reexport_fault_tag(fault_tag_id)
        self.assertEqual(revised["faultTag"]["faultTagId"], fault_tag_id)
        self.assertEqual(len(revised["faultTag"]["export"]["revisions"]), 2)
        self.assertFalse(self.store.read_fault_tag(fault_tag_id)["export"]["needs_reexport"])
        sent = {
            "message_key": "fault-tag-sent",
            "timestamp": "2026-08-09T09:00:00-05:00",
            "direction": "sent",
            "sender": "Zeus User",
            "sender_address": "user@example.com",
            "subject": f"Fault Tag {fault_tag_id}",
            "body": f"Fault Tag batch {fault_tag_id}",
            "html_body": "",
        }
        with self.store.transaction("fault-tag-email", {}) as staging:
            apply_spare_request_messages(self.store, staging, [sent])
        locked = self.service.fault_tag(fault_tag_id)
        self.assertTrue(locked["locked"])
        self.assertEqual(locked["email"]["message_key"], "fault-tag-sent")
        self.assertEqual(
            self.service.spare_request(request["requestId"])["items"][0]["lifecycle"]["stage"],
            4,
        )
        current = self.service.spare_request(request["requestId"])
        with self.assertRaisesRegex(ValidationError, "Delete Fault Tag"):
            self.service.edit_spare_request(
                request["requestId"],
                expected_revision=current["revision"],
                changes={"note": "A second supplier correction arrived."},
                item_updates=[{"itemId": item_id, "rma": "C3209937888"}],
            )

        with self.assertRaisesRegex(ValidationError, "Delete the active Fault Tag"):
            self.service.archive_spare_items(
                item_ids=[item_id],
                reason="cancelled",
                note="This item was linked to the wrong batch.",
            )

        deleted = self.service.delete_fault_tag(fault_tag_id)
        self.assertFalse(deleted["lifecycleChanged"])
        released = self.service.spare_request(request["requestId"])["items"][0]
        self.assertEqual(released["fault_tag_ids"], [])
        self.assertEqual(released["lifecycle"]["stage"], 4)
        current = self.service.spare_request(request["requestId"])
        corrected_after_delete = self.service.edit_spare_request(
            request["requestId"],
            expected_revision=current["revision"],
            changes={"note": "The locked tag was deleted and will be recreated."},
            item_updates=[{"itemId": item_id, "rma": "C3209937888"}],
        )["request"]
        self.assertEqual(corrected_after_delete["items"][0]["lifecycle"]["stage"], 4)

    def test_manually_sent_fault_tag_gets_internal_id_and_waits_for_warehouse_email(self) -> None:
        config = self.store.config
        config["email"]["warehouse_sender_domain"] = "@warehouse.example.test"
        self.store.save_config(config)
        request = self.request_at_spare_replaced()
        item = request["items"][0]

        registered = self.service.register_sent_fault_tag(
            [{"itemId": item["item_id"], "condition": "Faulty"}]
        )
        fault_tag_id = registered["faultTagId"]
        self.assertRegex(fault_tag_id, r"^FT-\d{12}$")
        self.assertEqual(registered["faultTag"]["status"], "sent")
        self.assertEqual(registered["faultTag"]["lockedSource"], "manual")
        self.assertIsNone(registered["faultTag"]["export"]["filename"])
        self.assertEqual(
            self.service.spare_request(request["requestId"])["items"][0]["lifecycle"]["stage"],
            4,
        )
        with self.assertRaisesRegex(ValidationError, "warehouse email evidence"):
            self.service.bulk_spare_lifecycle(
                item_ids=[item["item_id"]], action="advance"
            )

        warehouse = {
            "message_key": "manual-tag-warehouse-reply",
            "timestamp": "2026-08-10T09:00:00-05:00",
            "direction": "received",
            "sender": "Warehouse",
            "sender_address": "agent@warehouse.example.test",
            "subject": f"SR4956964 {item['rma']} RT12345678",
            "body": f"SR4956964 {item['rma']} RT12345678",
            "html_body": "",
        }
        with self.store.transaction("manual-tag-warehouse", {}) as staging:
            apply_spare_request_messages(self.store, staging, [warehouse])

        updated_request = self.service.spare_request(request["requestId"])
        self.assertEqual(updated_request["items"][0]["lifecycle"]["stage"], 5)
        updated_tag = self.service.fault_tag(fault_tag_id)
        self.assertEqual(updated_tag["status"], "awaiting_user_confirmation")
        self.assertEqual(
            updated_tag["members"][0]["warehouseEvidenceAt"],
            warehouse["timestamp"],
        )

    def test_partial_fault_tag_stays_active_until_every_member_is_confirmed(self) -> None:
        config = self.store.config
        config["email"]["warehouse_sender_domain"] = "@warehouse.example.test"
        self.store.save_config(config)
        request = self.request_at_spare_replaced(2)
        first, second = request["items"]
        fault_tag_id = self.service.export_spare_return(
            [
                {"itemId": first["item_id"], "condition": "Faulty"},
                {"itemId": second["item_id"], "condition": "New"},
            ]
        )["faultTagId"]

        def warehouse_message(item: dict, key: str) -> dict:
            return {
                "message_key": key,
                "timestamp": "2026-08-10T09:00:00-05:00",
                "direction": "received",
                "sender": "Warehouse",
                "sender_address": "agent@warehouse.example.test",
                "subject": f"SR4956964 {item['rma']} RT12345678",
                "body": f"SR4956964 {item['rma']} RT12345678",
                "html_body": "",
            }

        with self.store.transaction("first-warehouse-evidence", {}) as staging:
            apply_spare_request_messages(
                self.store, staging, [warehouse_message(first, "warehouse-first")]
            )
        active_tag = self.store.read_fault_tag(fault_tag_id, completed=False)
        self.assertEqual(fault_tag_status(active_tag), "partial_warehouse")
        request = self.service.spare_request(request["requestId"])
        stages = {item["item_id"]: item["lifecycle"]["stage"] for item in request["items"]}
        self.assertEqual(stages[first["item_id"]], 5)
        self.assertEqual(stages[second["item_id"]], 4)

        self.service.bulk_spare_lifecycle(
            item_ids=[first["item_id"]], action="advance"
        )
        active_tag = self.store.read_fault_tag(fault_tag_id, completed=False)
        self.assertEqual(fault_tag_status(active_tag), "partial_warehouse")
        self.assertTrue(active_tag["members"][0]["user_confirmed_at"])
        self.assertEqual(
            self.service.reexport_fault_tag(fault_tag_id)["faultTag"]["faultTagId"],
            fault_tag_id,
        )

        with self.store.transaction("second-warehouse-evidence", {}) as staging:
            apply_spare_request_messages(
                self.store, staging, [warehouse_message(second, "warehouse-second")]
            )
        self.service.bulk_spare_lifecycle(
            item_ids=[second["item_id"]], action="advance"
        )
        completed_tag = self.store.read_fault_tag(fault_tag_id, completed=True)
        self.assertEqual(fault_tag_status(completed_tag), "completed")
        self.assertFalse(self.store.spare_request_file(request["requestId"]).exists())

    def test_manual_export_persists_and_reexports_as_revision(self) -> None:
        lines = lines_input(2)
        lines[0].pop("reportDate")
        payload = {
            "source": "manual",
            "ticketId": "39416095",
            "reportDate": "2026-07-01",
            "profile": profile_input(),
            "lines": lines,
        }
        result = self.service.export_spare_request(payload)
        request_id = result["request"]["requestId"]
        self.assertTrue(Path(result["path"]).is_file())
        self.assertEqual(len(result["request"]["items"]), 2)
        self.assertEqual(result["request"]["reportDate"], "2026-07-01")
        self.assertTrue(
            all(
                line["report_date"] == "2026-07-01"
                for line in result["request"]["requestLines"]
            )
        )
        self.assertTrue(result["warnings"])
        revised = self.service.reexport_spare_request(request_id)
        self.assertRegex(revised["filename"], r"-r2\.xlsx$")
        dashboard = spare_requests_dashboard_payload(self.store, view="active")
        self.assertEqual(len(dashboard["spareRequests"]), 2)

    def test_submitted_source_part_stays_reserved_after_archive(self) -> None:
        self.store.write_ticket_bundle(self.store.current, source_ticket())
        payload = {
            "source": "ticket",
            "ticketId": "39416095",
            "reportDate": "2026-07-01",
            "profile": profile_input(),
            "lines": [
                {
                    **lines_input(1)[0],
                    "deviceNumber": 1,
                    "partNumber": 1,
                }
            ],
        }

        before = spare_requests_dashboard_payload(self.store, view="eligible")
        self.assertEqual([row["rowId"] for row in before["eligibleParts"]], ["39416095:1:1"])
        registered = self.service.register_spare_request(payload)["request"]

        self.assertEqual(registered["creationMethod"], "zeus_create")
        self.assertIsNone(registered["spareSr"])
        self.assertIsNone(registered["items"][0]["rma"])
        self.assertIsNone(registered["export"]["request_filename"])
        self.assertEqual(registered["history"][0]["action"], "request-created")
        self.assertFalse(self.exports.exists())
        self.assertEqual(
            spare_requests_dashboard_payload(self.store, view="eligible")["eligibleParts"],
            [],
        )
        with self.assertRaisesRegex(ValidationError, "already submitted"):
            self.service.export_spare_request(payload)

        item_id = registered["items"][0]["item_id"]
        self.service.archive_spare_items(
            item_ids=[item_id],
            reason="cancelled",
            note="The externally prepared request was cancelled.",
        )
        still_reserved = spare_requests_dashboard_payload(self.store, view="eligible")
        self.assertEqual(still_reserved["eligibleParts"], [])
        completed = spare_requests_dashboard_payload(
            self.store,
            view="completed",
            closed_path=self.workbooks / "Closed.xlsx",
        )["spareRequests"][0]
        self.assertEqual(completed["lifecycleStage"], 0)
        self.assertEqual(completed["lifecycleStageLabel"], "Cancelled")

        current = self.store.read_ticket("39416095")
        revision = ticket_revision(current)
        spare_parts = deepcopy(current["local"]["spare_parts"])
        spare_parts[0]["parts"].append(
            {
                "part_number": 2,
                "slot": "1/0/1",
                "part": "Controller board",
                "bom": "02312RCC",
                "notes": "Second physical replacement",
                "new_sn": None,
                "submitted_request_ids": [],
            }
        )
        spare_parts[0]["next_part_number"] = 3
        self.service.edit_ticket(
            "39416095",
            changes={"Spare Parts": spare_parts},
            expected_revision=revision,
        )
        eligible = spare_requests_dashboard_payload(self.store, view="eligible")
        self.assertEqual([row["rowId"] for row in eligible["eligibleParts"]], ["39416095:1:2"])

    def test_delete_unconfirmed_request_releases_source_and_preserves_export(self) -> None:
        self.store.write_ticket_bundle(self.store.current, source_ticket())
        payload = {
            "source": "ticket",
            "ticketId": "39416095",
            "reportDate": "2026-07-01",
            "profile": profile_input(),
            "lines": [{**lines_input(1)[0], "deviceNumber": 1, "partNumber": 1}],
        }
        exported = self.service.export_spare_request(payload)
        request = exported["request"]
        export_path = Path(exported["path"])
        marker = self.store.read_ticket("39416095")["local"]["spare_parts"][0]["parts"][0]
        self.assertEqual(marker["submitted_request_ids"], [request["requestId"]])
        self.assertEqual(request["trackingId"], request["requestId"])
        self.assertTrue(request["trackingIdProvisional"])
        self.assertTrue(request["canDelete"])

        deleted = self.service.delete_unconfirmed_spare_request(
            request["requestId"], expected_revision=request["revision"]
        )

        self.assertEqual(deleted["deleted"], request["requestId"])
        self.assertTrue(export_path.is_file())
        self.assertFalse(self.store.spare_request_file(request["requestId"]).exists())
        released = self.store.read_ticket("39416095")["local"]["spare_parts"][0]["parts"][0]
        self.assertEqual(released["submitted_request_ids"], [])
        eligible = spare_requests_dashboard_payload(self.store, view="eligible")
        self.assertEqual([row["rowId"] for row in eligible["eligibleParts"]], ["39416095:1:1"])

        confirmed = self.service.export_spare_request(payload)["request"]
        confirmed = self.service.edit_spare_request(
            confirmed["requestId"],
            expected_revision=confirmed["revision"],
            changes={"spareSr": "SR4956964", "note": "Confirmed manually"},
            item_updates=[
                {"itemId": confirmed["items"][0]["item_id"], "rma": "C3209937826"}
            ],
        )["request"]
        self.assertFalse(confirmed["canDelete"])
        with self.assertRaisesRegex(ValidationError, "confirmed Spare Request"):
            self.service.delete_unconfirmed_spare_request(
                confirmed["requestId"], expected_revision=confirmed["revision"]
            )

    def test_submitted_part_is_immutable_but_deletable_while_device_stays_owned(self) -> None:
        self.store.write_ticket_bundle(self.store.current, source_ticket())
        payload = {
            "source": "ticket",
            "ticketId": "39416095",
            "reportDate": "2026-07-01",
            "profile": profile_input(),
            "lines": [{**lines_input(1)[0], "deviceNumber": 1, "partNumber": 1}],
        }
        self.service.register_spare_request(payload)
        current = self.store.read_ticket("39416095")
        modified = deepcopy(current["local"]["spare_parts"])
        modified[0]["parts"][0]["bom"] = "DIFFERENT-BOM"
        with self.assertRaisesRegex(ValidationError, "submitted part 1 is immutable"):
            self.service.edit_ticket(
                "39416095",
                changes={"Spare Parts": modified},
                expected_revision=ticket_revision(current),
            )

        current = self.store.read_ticket("39416095")
        without_part = deepcopy(current["local"]["spare_parts"])
        without_part[0]["parts"] = []
        self.service.edit_ticket(
            "39416095",
            changes={"Spare Parts": without_part},
            expected_revision=ticket_revision(current),
        )
        current = self.store.read_ticket("39416095")
        self.assertEqual(current["local"]["spare_parts"][0]["parts"], [])
        self.assertEqual(
            self.service.ticket("39416095")["spareParts"][0]["active_request_ids"],
            [next(self.store.iter_spare_request_ids())],
        )
        with self.assertRaisesRegex(ValidationError, "still belongs to an Active Request"):
            self.service.edit_ticket(
                "39416095",
                changes={"Spare Parts": []},
                expected_revision=ticket_revision(current),
            )

        new_record = deepcopy(current["local"]["spare_parts"])
        new_record[0]["parts"] = [
            {
                "part_number": 2,
                "slot": "1/0/1",
                "part": "Controller board",
                "bom": "02312RCC",
                "notes": "New replacement record",
                "new_sn": None,
                "submitted_request_ids": [],
            }
        ]
        new_record[0]["next_part_number"] = 3
        self.service.edit_ticket(
            "39416095",
            changes={"Spare Parts": new_record},
            expected_revision=ticket_revision(current),
        )
        eligible = spare_requests_dashboard_payload(self.store, view="eligible")
        self.assertEqual([row["rowId"] for row in eligible["eligibleParts"]], ["39416095:1:2"])

    def test_device_removal_requires_a_separate_spare_part_removal_save(self) -> None:
        self.store.write_ticket_bundle(self.store.current, source_ticket())
        current = self.store.read_ticket("39416095")
        with self.assertRaisesRegex(ValidationError, "still has Spare Parts records"):
            self.service.edit_ticket(
                "39416095",
                changes={"Spare Parts": []},
                expected_revision=ticket_revision(current),
            )

        without_parts = deepcopy(current["local"]["spare_parts"])
        without_parts[0]["parts"] = []
        self.service.edit_ticket(
            "39416095",
            changes={"Spare Parts": without_parts},
            expected_revision=ticket_revision(current),
        )
        current = self.store.read_ticket("39416095")
        self.assertEqual(len(current["local"]["spare_parts"]), 1)
        self.assertEqual(current["local"]["spare_parts"][0]["parts"], [])
        self.assertEqual(current["local"]["fields"]["Spare"], "N")
        self.service.edit_ticket(
            "39416095",
            changes={"Spare Parts": []},
            expected_revision=ticket_revision(current),
        )
        self.assertEqual(self.store.read_ticket("39416095")["local"]["spare_parts"], [])

    def test_archiving_a_legacy_active_request_backfills_its_submission_marker(self) -> None:
        self.store.write_ticket_bundle(self.store.current, source_ticket())
        legacy = request_record(1)
        legacy["creation_method"] = "zeus_export"
        legacy["request_lines"][0]["source_device_number"] = 1
        legacy["request_lines"][0]["source_part_number"] = 1
        legacy["items"][0]["source_device_number"] = 1
        legacy["items"][0]["source_part_number"] = 1
        with self.store.transaction("seed-legacy-active-request", {}) as staging:
            self.store.write_spare_request(staging, legacy)
        source = self.store.read_ticket("39416095")["local"]["spare_parts"][0]["parts"][0]
        self.assertEqual(source["submitted_request_ids"], [])

        self.service.archive_spare_items(
            item_ids=[legacy["items"][0]["item_id"]],
            reason="cancelled",
            note="Legacy request cancelled after migration.",
        )

        source = self.store.read_ticket("39416095")["local"]["spare_parts"][0]["parts"][0]
        self.assertEqual(source["submitted_request_ids"], [legacy["request_id"]])

    def test_lifecycle_requires_outbound_and_warehouse_evidence_and_switches_tracking_id(self) -> None:
        request = self.service.export_spare_request(
            {
                "source": "manual",
                "ticketId": "39416095",
                "reportDate": "2026-07-01",
                "profile": profile_input(),
                "lines": lines_input(1),
            }
        )["request"]
        item_id = request["items"][0]["item_id"]
        self.assertEqual(request["items"][0]["lifecycle"]["stage"], 0)
        self.assertEqual(request["trackingId"], request["requestId"])
        self.assertTrue(request["trackingIdProvisional"])
        with self.assertRaisesRegex(ValidationError, "one lifecycle stage"):
            self.service.advance_spare_request_stage(
                request["requestId"],
                item_id=item_id,
                target_stage=2,
                expected_revision=request["revision"],
            )

        with self.assertRaisesRegex(ValidationError, "outbound email evidence"):
            self.service.advance_spare_request_stage(
                request["requestId"],
                item_id=item_id,
                target_stage=1,
                expected_revision=request["revision"],
            )
        self.detect_request_sent(request)
        request = self.service.spare_request(request["requestId"])
        self.assertEqual(request["items"][0]["lifecycle"]["stage"], 1)
        request = self.service.edit_spare_request(
            request["requestId"],
            expected_revision=request["revision"],
            changes={"spareSr": "SR4956964", "note": "Confirmed manually"},
            item_updates=[{"itemId": item_id, "rma": "C3209937826"}],
        )["request"]
        self.assertEqual(request["items"][0]["lifecycle"]["stage"], 1)
        self.assertEqual(request["trackingId"], "SR4956964")
        self.assertFalse(request["trackingIdProvisional"])

        request = self.service.advance_spare_request_stage(
            request["requestId"],
            item_id=item_id,
            target_stage=2,
            expected_revision=request["revision"],
        )["request"]
        self.assertEqual(request["items"][0]["lifecycle"]["stage"], 2)

        for stage in range(3, 5):
            request = self.service.advance_spare_request_stage(
                request["requestId"],
                item_id=item_id,
                target_stage=stage,
                expected_revision=request["revision"],
            )["request"]
            self.assertEqual(request["items"][0]["lifecycle"]["stage"], stage)
            self.assertEqual(
                lifecycle_stage(
                    self.store.read_spare_request(request["requestId"])["items"][0],
                    self.store.read_spare_request(request["requestId"]),
                ),
                stage,
            )
        with self.assertRaisesRegex(ValidationError, "warehouse email evidence"):
            self.service.advance_spare_request_stage(
                request["requestId"],
                item_id=item_id,
                target_stage=5,
                expected_revision=request["revision"],
            )
        with self.store.transaction("warehouse-evidence", {}) as staging:
            persisted = self.store.read_spare_request(request["requestId"], staging)
            item = persisted["items"][0]
            item["warehouse_candidate_at"] = "2026-08-10T09:00:00-05:00"
            item["warehouse_confirmation_source"] = "email"
            item["warehouse_message_key"] = "warehouse-evidence"
            item["warehouse_evidence"] = [{
                "message_key": "warehouse-evidence",
                "timestamp": "2026-08-10T09:00:00-05:00",
                "rt": "RT12345678",
            }]
            self.store.write_spare_request(staging, persisted)
        request = self.service.spare_request(request["requestId"])
        self.assertEqual(request["items"][0]["lifecycle"]["stage"], 5)
        completed = self.service.advance_spare_request_stage(
            request["requestId"],
            item_id=item_id,
            target_stage=6,
            expected_revision=request["revision"],
        )
        self.assertIsNone(completed["request"])
        self.assertEqual(completed["completed"], [item_id])

    def test_rma_correction_reserves_aliases_and_return_archive_requires_confirmation(self) -> None:
        result = self.service.export_spare_request(
            {
                "source": "manual",
                "ticketId": "39416095",
                "profile": profile_input(),
                "lines": lines_input(1),
            }
        )
        detail = result["request"]
        item_id = detail["items"][0]["item_id"]
        saved = self.service.edit_spare_request(
            detail["requestId"],
            expected_revision=detail["revision"],
            changes={"spareSr": "SR4956964", "note": "manual confirmation"},
            item_updates=[
                {
                    "itemId": item_id,
                    "rma": "C3209937826",
                    "deliveredBom": "02540255",
                    "newSn": "NEW-1",
                    "dispatchAt": "2026-08-08T12:00",
                    "attended": True,
                }
            ],
        )["request"]
        with self.assertRaisesRegex(ValidationError, "audit note"):
            self.service.edit_spare_request(
                saved["requestId"],
                expected_revision=saved["revision"],
                changes={},
                item_updates=[{"itemId": item_id, "rma": "C3209937827"}],
            )
        corrected = self.service.edit_spare_request(
            saved["requestId"],
            expected_revision=saved["revision"],
            changes={"note": "The supplier corrected a transposed RMA digit."},
            item_updates=[{"itemId": item_id, "rma": "C3209937827"}],
        )["request"]
        self.assertEqual(corrected["items"][0]["rma"], "C3209937827")
        self.assertEqual(corrected["items"][0]["rma_aliases"], ["C3209937826"])
        self.assertEqual(corrected["history"][-2]["action"], "rma-corrected")
        another = self.service.export_spare_request(
            {
                "source": "manual",
                "ticketId": "39416096",
                "profile": profile_input(),
                "lines": lines_input(1),
            }
        )["request"]
        with self.assertRaisesRegex(ValidationError, "historical alias"):
            self.service.edit_spare_request(
                another["requestId"],
                expected_revision=another["revision"],
                changes={},
                item_updates=[
                    {"itemId": another["items"][0]["item_id"], "rma": "C3209937826"}
                ],
            )
        with self.assertRaisesRegex(ValidationError, "Warehouse email evidence"):
            self.service.archive_spare_items(
                item_ids=[item_id], reason="returned", note="", manual_override=False
            )
        with self.assertRaisesRegex(ValidationError, "cannot replace"):
            self.service.archive_spare_items(
                item_ids=[item_id],
                reason="returned",
                note="Verified by warehouse call",
                manual_override=True,
            )
        archived = self.service.archive_spare_items(
            item_ids=[item_id],
            reason="cancelled",
            note="Request cancelled after the return could not be verified.",
        )
        self.assertEqual(archived["archived"], [item_id])
        self.assertFalse(self.store.spare_request_file(saved["requestId"]).exists())
        completed = spare_requests_dashboard_payload(
            self.store,
            view="completed",
            closed_path=self.workbooks / "Closed.xlsx",
        )
        self.assertEqual(completed["spareRequests"][0]["itemId"], item_id)
        with self.assertRaisesRegex(ValidationError, "completed item"):
            self.service.edit_spare_request(
                another["requestId"],
                expected_revision=another["revision"],
                changes={},
                item_updates=[
                    {"itemId": another["items"][0]["item_id"], "rma": "C3209937826"}
                ],
            )
        backup_dir = self.workbooks / "Zeus Backups" / "Spare Requests"
        backup_dir.mkdir(parents=True, exist_ok=True)
        private_backup = backup_dir / "Closed_20260808_120000_000000.xlsx"
        shutil.copy2(self.workbooks / "Closed.xlsx", private_backup)
        purged = self.service.purge_spare_data(item_ids=[item_id])
        self.assertEqual(purged["removed"], 1)
        self.assertEqual(purged["backupsRemoved"], 1)
        self.assertFalse(private_backup.exists())


if __name__ == "__main__":
    unittest.main()
