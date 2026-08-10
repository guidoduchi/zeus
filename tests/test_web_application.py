from __future__ import annotations

import http.client
import json
import re
import sys
import tempfile
import threading
import time
import unittest
import urllib.error
import urllib.request
from datetime import date
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

import zeus2.main as main_module
from tests.test_zeus2 import pending_row, upstream_row, write_advanced, write_managed
from zeus2.application.edits import (
    TicketRevisionConflictError,
    edit_ticket_through_pendings,
)
from zeus2.application.errors import ValidationError
from zeus2.application.jobs import EventBroker, JobManager
from zeus2.application.serialization import (
    dashboard_payload,
    spare_parts_dashboard_payload,
    ticket_revision,
)
from zeus2.application.service import ApplicationService
from zeus2.config import save_config
from zeus2.excel_export import (
    publish_operational_workbooks,
    recreate_pendings_from_database,
    recover_pendings_recreation,
)
from zeus2.excel_import import WorkbookValidationError, read_pendings
from zeus2.main import _restart_command
from zeus2.reconcile import sync_advanced_search
from zeus2.reference_data import save_user_profile, user_profile_path
from zeus2.startup import StartupResult, run_startup
from zeus2.store import ZeusStore
from zeus2.tickets import PENDING_COLUMNS
from zeus2.utils import sha256_file
from zeus2.web.runtime import InstanceRegistry, process_creation_marker
from zeus2.web.server import ZeusWebServer


class WebFixture(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.home = self.root / "home"
        self.books = self.root / "books"
        self.downloads = self.root / "downloads"
        self.books.mkdir()
        self.downloads.mkdir()
        self.store = ZeusStore(self.home / "data", config_home=self.home)
        self.store.ensure_layout()
        config = self.store.config
        config["paths"]["workbook_directory"] = str(self.books)
        config["paths"]["advanced_search_directory"] = str(self.downloads)
        config["paths"]["outlook_store_path"] = None
        config["advanced_search"]["poll_interval_minutes"] = 0
        save_config(self.home, config)

    def tearDown(self) -> None:
        self.temporary.cleanup()

    def seed_database(self) -> None:
        """Create one discovered ticket plus representative Zeus-owned fields."""

        row = pending_row(
            "12345678",
            summary="Advanced Search builds Zeus",
            **{"Site": "GYE", "Notes": "Initial note"},
        )
        write_managed(
            self.books / "Pendings.xlsx",
            [row],
        )
        source = self.downloads / "Advanced Search(Service Request)20260801010101.xlsx"
        write_advanced(
            source,
            [upstream_row("12345678", summary="Advanced Search builds Zeus")],
        )
        sync_advanced_search(self.store, source)
        ticket = self.store.read_ticket("12345678")
        edit_ticket_through_pendings(
            self.store,
            "12345678",
            {"Site": "GYE", "Notes": "Initial note"},
            expected_revision=ticket_revision(ticket),
        )


class DatabaseFirstSourceTests(WebFixture):
    def test_pendings_alone_is_read_only_output_and_does_not_seed_database(self) -> None:
        write_managed(
            self.books / "Pendings.xlsx",
            [pending_row("12345678", **{"Notes": "must not import"})],
        )
        before_hash = sha256_file(self.books / "Pendings.xlsx")
        result = run_startup(self.store)

        self.assertFalse(result.advanced_search_valid)
        self.assertFalse(
            any("ADVANCED SEARCH WARNING" in warning for warning in result.warnings)
        )
        self.assertTrue(
            any("Advanced Search was skipped" in notice for notice in result.notices)
        )
        self.assertFalse((self.books / "Closed.xlsx").exists())
        self.assertEqual(sha256_file(self.books / "Pendings.xlsx"), before_hash)
        dashboard = dashboard_payload(self.store)
        self.assertEqual(dashboard["stats"]["active"], 0)
        self.assertEqual(dashboard["tickets"], [])

    def test_browser_edit_writes_database_without_touching_pendings(self) -> None:
        self.seed_database()
        run_startup(self.store)
        before = self.store.read_ticket("12345678")
        workbook_hash = sha256_file(self.books / "Pendings.xlsx")

        result = edit_ticket_through_pendings(
            self.store,
            "12345678",
            {"Notes": "Changed from Zeus web", "Done?": "Y"},
            expected_revision=ticket_revision(before),
        )

        self.assertTrue(result["changed"])
        self.assertEqual(sha256_file(self.books / "Pendings.xlsx"), workbook_hash)
        updated = self.store.read_ticket("12345678")
        self.assertEqual(updated["local"]["fields"]["Notes"], "Changed from Zeus web")
        self.assertEqual(updated["local"]["fields"]["Done?"], "Y")
        self.assertFalse((self.books / "Zeus Backups" / "Web edits").exists())

    def test_selected_ticket_drafts_commit_in_one_database_transaction(self) -> None:
        source = self.downloads / "Advanced Search(Service Request)20260808010101.xlsx"
        write_advanced(
            source,
            [
                upstream_row("12345678", summary="First protected draft"),
                upstream_row("87654321", summary="Second protected draft"),
            ],
        )
        sync_advanced_search(self.store, source)
        service = ApplicationService(self.store)
        try:
            first = service.ticket("12345678")
            second = service.ticket("87654321")
            result = service.edit_tickets(
                [
                    {
                        "ticketId": "12345678",
                        "revision": first["revision"],
                        "changes": {"Notes": "First saved draft"},
                    },
                    {
                        "ticketId": "87654321",
                        "revision": second["revision"],
                        "changes": {"Planned Date": "2026-08-20", "Done?": "P"},
                    },
                ]
            )
        finally:
            service.stop()

        self.assertTrue(result["changed"])
        self.assertEqual(result["ticketIds"], ["12345678", "87654321"])
        self.assertEqual(
            self.store.read_ticket("12345678")["local"]["fields"]["Notes"],
            "First saved draft",
        )
        self.assertEqual(
            self.store.read_ticket("87654321")["local"]["fields"]["Done?"],
            "P",
        )
        audit = self.store.audit_file.read_text(encoding="utf-8")
        self.assertEqual(audit.count('"action": "web-local-batch-edit"'), 1)

    def test_conflicted_draft_batch_changes_no_ticket(self) -> None:
        source = self.downloads / "Advanced Search(Service Request)20260808010202.xlsx"
        write_advanced(
            source,
            [
                upstream_row("12345678", summary="First protected draft"),
                upstream_row("87654321", summary="Second protected draft"),
            ],
        )
        sync_advanced_search(self.store, source)
        first = self.store.read_ticket("12345678")
        second = self.store.read_ticket("87654321")
        first_revision = ticket_revision(first)
        stale_second_revision = ticket_revision(second)
        edit_ticket_through_pendings(
            self.store,
            "87654321",
            {"Notes": "Concurrent database change"},
            expected_revision=stale_second_revision,
        )
        service = ApplicationService(self.store)
        try:
            with self.assertRaises(TicketRevisionConflictError):
                service.edit_tickets(
                    [
                        {
                            "ticketId": "12345678",
                            "revision": first_revision,
                            "changes": {"Notes": "Must roll back"},
                        },
                        {
                            "ticketId": "87654321",
                            "revision": stale_second_revision,
                            "changes": {"Notes": "Must conflict"},
                        },
                    ]
                )
        finally:
            service.stop()

        self.assertNotEqual(
            self.store.read_ticket("12345678")["local"]["fields"].get("Notes"),
            "Must roll back",
        )
        self.assertEqual(
            self.store.read_ticket("87654321")["local"]["fields"]["Notes"],
            "Concurrent database change",
        )

    def test_normalized_spare_parts_save_to_database_then_export(self) -> None:
        self.seed_database()
        run_startup(self.store)
        before = self.store.read_ticket("12345678")
        self.assertEqual(before["local"]["fields"]["Spare"], "N")
        before_export_hash = sha256_file(self.books / "Pendings.xlsx")

        spare_parts = [
            {
                "device": "server-a",
                "model": "2288H V5",
                "parts": [
                    {
                        "slot": "Slot 1",
                        "part": "Disk",
                        "bom": "BOM-9000",
                        "faulty_sn": "FAULTY-1",
                        "new_sn": "NEW-1",
                    },
                    {
                        "slot": "Slot 2",
                        "part": "Disk",
                        "bom": "BOM-9001",
                        "faulty_sn": "FAULTY-2",
                        "new_sn": None,
                    },
                ],
            },
            {
                "device": "server-b",
                "model": "CH121 V5",
                "parts": [
                    {
                        "slot": "DIMM 3",
                        "part": "Memory",
                        "bom": "BOM-9002",
                        "faulty_sn": "FAULTY-3",
                        "new_sn": "NEW-3",
                    }
                ],
            },
        ]
        result = edit_ticket_through_pendings(
            self.store,
            "12345678",
            {"Spare Parts": spare_parts, "Planned Date": "2026-08-21"},
            expected_revision=ticket_revision(before),
        )

        self.assertEqual(
            result["changedFields"],
            ["Spare Parts", "Planned Date", "Spare"],
        )
        self.assertEqual(sha256_file(self.books / "Pendings.xlsx"), before_export_hash)
        updated = self.store.read_ticket("12345678")
        self.assertEqual(updated["local"]["fields"]["Planned Date"], "2026-08-21")
        self.assertEqual(updated["local"]["fields"]["Spare"], "Y")
        self.assertEqual(
            updated["local"]["spare_parts"],
            [
                {
                    "device": "server-a",
                    "model": "2288H V5",
                    "faulty_sns": ["FAULTY-1", "FAULTY-2"],
                    "parts": [
                        {
                            "slot": "Slot 1",
                            "part": "Disk",
                            "bom": "BOM-9000",
                            "new_sn": "NEW-1",
                            "notes": None,
                        },
                        {
                            "slot": "Slot 2",
                            "part": "Disk",
                            "bom": "BOM-9001",
                            "new_sn": None,
                            "notes": None,
                        },
                    ],
                },
                {
                    "device": "server-b",
                    "model": "CH121 V5",
                    "faulty_sns": ["FAULTY-3"],
                    "parts": [
                        {
                            "slot": "DIMM 3",
                            "part": "Memory",
                            "bom": "BOM-9002",
                            "new_sn": "NEW-3",
                            "notes": None,
                        }
                    ],
                },
            ],
        )

        publish_operational_workbooks(self.store, self.books, create_missing=True)
        workbook = load_workbook(self.books / "Pendings.xlsx", data_only=False)
        try:
            headers = [cell.value for cell in workbook.active[1]]
            bom = workbook.active.cell(2, headers.index("BOM") + 1)
            planned = workbook.active.cell(2, headers.index("Planned Date") + 1)
            spare = workbook.active.cell(2, headers.index("Spare") + 1)
            self.assertEqual(bom.value, "BOM-9000\nBOM-9001\nBOM-9002")
            self.assertEqual(planned.value.date(), date(2026, 8, 21))
            self.assertEqual(planned.number_format, "yyyy-mm-dd")
            self.assertEqual(spare.value, "Y")
            spare_sheet = workbook["Spare Parts"]
            spare_headers = [cell.value for cell in spare_sheet[1]]
            self.assertEqual(spare_sheet.max_row, 4)
            self.assertEqual(
                spare_sheet.cell(2, spare_headers.index("Device") + 1).value,
                "server-a",
            )
            self.assertEqual(
                spare_sheet.cell(4, spare_headers.index("Device #") + 1).value,
                2,
            )
        finally:
            workbook.close()

        spare_dashboard = spare_parts_dashboard_payload(
            self.store, sort="bom", direction="asc"
        )
        self.assertEqual(spare_dashboard["workspace"], "spare-parts")
        self.assertEqual(spare_dashboard["stats"], {
            "tickets": 1,
            "currentTickets": 1,
            "closedTickets": 0,
            "devices": 2,
            "parts": 3,
            "withBom": 3,
            "missingBom": 0,
            "newSnRecorded": 2,
        })
        self.assertEqual(
            [row["rowId"] for row in spare_dashboard["spareParts"]],
            ["12345678:1:1", "12345678:1:2", "12345678:2:1"],
        )
        self.assertEqual(spare_dashboard["columns"][0]["key"], "ticketId")
        self.assertTrue(
            all(row["ticketId"] == "12345678" for row in spare_dashboard["spareParts"])
        )
        descending_sr = spare_parts_dashboard_payload(
            self.store, sort="sr", direction="desc"
        )
        self.assertEqual(
            [row["rowId"] for row in descending_sr["spareParts"]],
            ["12345678:1:1", "12345678:1:2", "12345678:2:1"],
        )
        self.assertEqual(
            spare_dashboard["spareParts"][2]["part"],
            "Memory",
        )
        filtered = spare_parts_dashboard_payload(self.store, search="FAULTY-2")
        self.assertEqual(
            [row["bom"] for row in filtered["spareParts"]],
            ["BOM-9000", "BOM-9001"],
        )

    def test_spare_parts_export_folder_is_configurable_but_never_queried(self) -> None:
        self.seed_database()
        run_startup(self.store)
        export_directory = self.root / "spare-exports"
        export_directory.mkdir()
        external_workbook = export_directory / "request-draft.xlsx"
        external_workbook.write_bytes(b"not an importable workbook")

        service = ApplicationService(self.store)
        try:
            payload = service.save_settings({
                "paths.spare_parts_export_directory": str(export_directory),
            })
            setting = next(
                item
                for item in payload["settings"]
                if item["key"] == "paths.spare_parts_export_directory"
            )
            self.assertEqual(setting["value"], str(export_directory.resolve()))
            self.assertTrue(setting["status"]["writeOnly"])
            self.assertIn("Export-only", setting["status"]["message"])

            query = run_startup(self.store)
            self.assertEqual(
                external_workbook.read_bytes(),
                b"not an importable workbook",
            )
            self.assertFalse(
                any("request-draft" in message for message in query.warnings)
            )
        finally:
            service.stop()

    def test_finalized_spare_parts_remain_grouped_under_their_sr(self) -> None:
        self.seed_database()
        run_startup(self.store)
        before = self.store.read_ticket("12345678")
        edit_ticket_through_pendings(
            self.store,
            "12345678",
            {
                "Spare Parts": [
                    {
                        "device": "server-closed",
                        "model": "2288H V5",
                        "parts": [
                            {
                                "slot": "Slot 3",
                                "part": "Disk",
                                "bom": "BOM-CLOSED",
                                "faulty_sn": "FAULTY-CLOSED",
                                "new_sn": None,
                            }
                        ],
                    }
                ]
            },
            expected_revision=ticket_revision(before),
        )
        closing = self.store.read_ticket("12345678")
        closing["lifecycle"]["status"] = "closure_pending"
        with self.store.transaction("test-finalized-spare-parts", {}) as staging:
            self.store.write_ticket_bundle(staging, closing)
        publish_operational_workbooks(
            self.store,
            self.books,
            create_missing=True,
        )
        self.assertFalse(self.store.ticket_file("12345678").exists())

        service = ApplicationService(self.store)
        try:
            dashboard = service.dashboard(
                workspace="spare-parts",
                sort="sr",
                direction="desc",
                search="",
            )
            self.assertEqual(dashboard["columns"][0]["key"], "ticketId")
            self.assertEqual(dashboard["stats"]["currentTickets"], 0)
            self.assertEqual(dashboard["stats"]["closedTickets"], 1)
            self.assertEqual(len(dashboard["spareParts"]), 1)
            row = dashboard["spareParts"][0]
            self.assertEqual(row["ticketId"], "12345678")
            self.assertEqual(row["lifecycle"], "closed")
            self.assertTrue(row["readOnly"])
            self.assertEqual(row["source"], "closed")

            detail = service.ticket("12345678")
            self.assertEqual(detail["ticketId"], "12345678")
            self.assertEqual(detail["summary"], "Advanced Search builds Zeus")
            self.assertTrue(detail["readOnly"])
            self.assertEqual(detail["source"], "closed")
            self.assertEqual(detail["spareParts"][0]["parts"][0]["bom"], "BOM-CLOSED")
            with self.assertRaisesRegex(ValidationError, "finalized and read-only"):
                service.edit_ticket(
                    "12345678",
                    changes={"Notes": "must not change"},
                    expected_revision=detail["revision"],
                )
        finally:
            service.stop()

    def test_normalized_spare_part_rejects_a_missing_parent_sr(self) -> None:
        self.seed_database()
        run_startup(self.store)
        before = self.store.read_ticket("12345678")
        edit_ticket_through_pendings(
            self.store,
            "12345678",
            {
                "Spare Parts": [
                    {
                        "device": "server-a",
                        "model": "2288H V5",
                        "parts": [
                            {
                                "slot": "Slot 1",
                                "part": "Disk",
                                "bom": "BOM-1",
                                "faulty_sn": "FAULTY-1",
                                "new_sn": None,
                            }
                        ],
                    }
                ]
            },
            expected_revision=ticket_revision(before),
        )
        publish_operational_workbooks(self.store, self.books, create_missing=True)
        path = self.books / "Pendings.xlsx"
        workbook = load_workbook(path)
        try:
            workbook["Spare Parts"]["A2"] = None
            workbook.save(path)
        finally:
            workbook.close()

        with self.assertRaisesRegex(
            WorkbookValidationError,
            "Spare Parts row 2: ticket ID is blank or invalid",
        ):
            read_pendings(path)

    def test_spare_cannot_be_supplied_by_a_browser_edit(self) -> None:
        self.seed_database()
        run_startup(self.store)
        before = self.store.read_ticket("12345678")
        workbook_hash = sha256_file(self.books / "Pendings.xlsx")

        with self.assertRaisesRegex(
            ValidationError,
            "Spare is an export-only value and cannot be edited directly",
        ):
            edit_ticket_through_pendings(
                self.store,
                "12345678",
                {"Spare": "Y"},
                expected_revision=ticket_revision(before),
            )

        self.assertEqual(sha256_file(self.books / "Pendings.xlsx"), workbook_hash)
        self.assertEqual(self.store.read_ticket("12345678")["local"]["fields"]["Spare"], "N")

    def test_spare_parts_sheet_manual_edits_never_flow_back_to_database(self) -> None:
        self.seed_database()
        run_startup(self.store)
        before = self.store.read_ticket("12345678")
        edit_ticket_through_pendings(
            self.store,
            "12345678",
            {
                "Spare Parts": [
                    {
                        "device": "server-a",
                        "model": "2288H V5",
                        "parts": [
                            {
                                "slot": "Slot 1",
                                "part": "Disk",
                                "bom": "BOM-1",
                                "faulty_sn": "OLD-1",
                                "new_sn": None,
                            }
                        ],
                    }
                ]
            },
            expected_revision=ticket_revision(before),
        )
        publish_operational_workbooks(self.store, self.books, create_missing=True)

        workbook = load_workbook(self.books / "Pendings.xlsx")
        sheet = workbook["Spare Parts"]
        headers = [cell.value for cell in sheet[1]]
        sheet.cell(2, headers.index("New SN") + 1).value = "NEW-1"
        sheet.append(
            [
                "12345678",
                1,
                "server-a",
                "2288H V5",
                2,
                "Slot 2",
                "Disk",
                "BOM-2",
                "OLD-2",
                "NEW-2",
            ]
        )
        workbook.save(self.books / "Pendings.xlsx")
        workbook.close()

        run_startup(self.store)
        authoritative = self.store.read_ticket("12345678")
        self.assertEqual(len(authoritative["local"]["spare_parts"]), 1)
        self.assertEqual(len(authoritative["local"]["spare_parts"][0]["parts"]), 1)
        self.assertIsNone(
            authoritative["local"]["spare_parts"][0]["parts"][0]["new_sn"]
        )

        publish_operational_workbooks(self.store, self.books, create_missing=True)
        generated = read_pendings(self.books / "Pendings.xlsx")
        self.assertEqual(
            len(generated.records["12345678"]["local"]["spare_parts"][0]["parts"]),
            1,
        )

    def test_normalized_workbook_rejects_deleted_spare_parts_sheet(self) -> None:
        self.seed_database()
        run_startup(self.store)
        before = self.store.read_ticket("12345678")
        edit_ticket_through_pendings(
            self.store,
            "12345678",
            {
                "Spare Parts": [
                    {
                        "device": "server-a",
                        "model": "2288H V5",
                        "parts": [
                            {
                                "slot": "Slot 1",
                                "part": "Disk",
                                "bom": "BOM-1",
                                "faulty_sn": "OLD-1",
                                "new_sn": None,
                            }
                        ],
                    }
                ]
            },
            expected_revision=ticket_revision(before),
        )
        publish_operational_workbooks(self.store, self.books, create_missing=True)

        path = self.books / "Pendings.xlsx"
        workbook = load_workbook(path)
        workbook.remove(workbook["Spare Parts"])
        workbook.save(path)
        workbook.close()

        with self.assertRaisesRegex(
            WorkbookValidationError,
            "Spare Parts is missing from a normalized Zeus workbook",
        ):
            read_pendings(path)

    def test_dashboard_combines_email_age_and_count_in_last_email_column(self) -> None:
        self.seed_database()
        run_startup(self.store)
        ticket = self.store.read_ticket("12345678")
        ticket["email"]["total_received"] = 5
        ticket["email"]["total_sent"] = 3
        self.store.write_ticket_bundle(self.store.current, ticket)

        dashboard = dashboard_payload(self.store)
        self.assertEqual(dashboard["tickets"][0]["emailCount"], 8)
        email_column = next(
            column for column in dashboard["columns"] if column["key"] == "emailLabel"
        )
        mw_column = next(
            column for column in dashboard["columns"] if column["key"] == "done"
        )
        self.assertEqual(email_column["label"], "Last Email")
        self.assertTrue(email_column["default"])
        self.assertEqual(mw_column, {
            "key": "done",
            "label": "MW",
            "width": 104,
            "default": True,
        })
        self.assertEqual(dashboard["tickets"][0]["emailLabel"], "No email")
        self.assertEqual(dashboard["tickets"][0]["customerContact"], "Customer")
        self.assertNotIn("emailCount", {column["key"] for column in dashboard["columns"]})

    def test_external_excel_change_is_ignored_by_database_edit(self) -> None:
        self.seed_database()
        run_startup(self.store)
        ticket = self.store.read_ticket("12345678")
        workbook = load_workbook(self.books / "Pendings.xlsx")
        headers = [cell.value for cell in workbook.active[1]]
        workbook.active.cell(2, headers.index("Notes") + 1).value = "Unsynchronized Excel edit"
        workbook.save(self.books / "Pendings.xlsx")
        workbook.close()
        changed_hash = sha256_file(self.books / "Pendings.xlsx")

        edit_ticket_through_pendings(
            self.store,
            "12345678",
            {"Notes": "Database value wins at export"},
            expected_revision=ticket_revision(ticket),
        )

        self.assertEqual(sha256_file(self.books / "Pendings.xlsx"), changed_hash)
        self.assertEqual(
            self.store.read_ticket("12345678")["local"]["fields"]["Notes"],
            "Database value wins at export",
        )

    def test_query_never_recreates_or_restores_deleted_pendings(self) -> None:
        self.seed_database()
        run_startup(self.store)
        backup_directory = self.books / "Zeus Backups"
        backup_directory.mkdir(exist_ok=True)
        write_managed(
            backup_directory / "Pendings_20990101_000000_000000.xlsx",
            [pending_row("12345678", **{"Notes": "Value from the latest backup"})],
        )
        (self.books / "Pendings.xlsx").unlink()

        ordinary_startup = run_startup(self.store)
        self.assertFalse((self.books / "Pendings.xlsx").exists())
        self.assertFalse(any("Pendings" in warning for warning in ordinary_startup.warnings))

        query = run_startup(self.store, recreate_missing_pendings=True)

        self.assertNotIn("pendings_recreation", query.operations)
        self.assertFalse((self.books / "Pendings.xlsx").exists())
        self.assertEqual(
            self.store.read_ticket("12345678")["local"]["fields"]["Notes"],
            "Initial note",
        )
        self.assertFalse((self.books / "Closed.xlsx").exists())

    def test_save_succeeds_without_pendings_and_waits_for_export(self) -> None:
        self.seed_database()
        run_startup(self.store)
        before = self.store.read_ticket("12345678")
        (self.books / "Pendings.xlsx").unlink()

        result = edit_ticket_through_pendings(
            self.store,
            "12345678",
            {"Notes": "Saved after deletion"},
            expected_revision=ticket_revision(before),
        )

        self.assertTrue(result["changed"])
        self.assertNotIn("pendingsRecreated", result)
        self.assertFalse((self.books / "Pendings.xlsx").exists())
        self.assertEqual(
            self.store.read_ticket("12345678")["local"]["fields"]["Notes"],
            "Saved after deletion",
        )
        self.assertFalse((self.books / "Closed.xlsx").exists())

    def test_closure_pending_is_deleted_only_after_successful_explicit_export(self) -> None:
        self.seed_database()
        run_startup(self.store)
        before_export_hash = sha256_file(self.books / "Pendings.xlsx")
        newer = self.downloads / "Advanced Search(Service Request)20260802010101.xlsx"
        write_advanced(newer, [upstream_row("87654321", summary="New source ticket")])

        query = run_startup(self.store)

        self.assertEqual(
            query.operations["advanced_search"]["closure_pending_ids"],
            ["12345678"],
        )
        self.assertEqual(sha256_file(self.books / "Pendings.xlsx"), before_export_hash)
        self.assertEqual(
            self.store.read_ticket("12345678")["lifecycle"]["status"],
            "closure_pending",
        )
        self.assertTrue(self.store.ticket_file("87654321").exists())

        # Rechecking the same source keeps the marker and never finalizes it.
        run_startup(self.store)
        self.assertTrue(self.store.ticket_file("12345678").exists())

        publish_operational_workbooks(self.store, self.books, create_missing=True)

        self.assertFalse(self.store.ticket_file("12345678").exists())
        self.assertEqual(
            set(read_pendings(self.books / "Pendings.xlsx").records),
            {"87654321"},
        )
        closed = load_workbook(self.books / "Closed.xlsx", read_only=True)
        try:
            self.assertEqual(str(closed.active["A2"].value), "12345678")
        finally:
            closed.close()

    def test_interrupted_recreation_is_finalized_from_its_journal(self) -> None:
        self.seed_database()
        run_startup(self.store)
        (self.books / "Pendings.xlsx").unlink()

        with patch(
            "zeus2.excel_export._finalize_pendings_recreation",
            side_effect=RuntimeError("simulated interruption"),
        ):
            with self.assertRaisesRegex(RuntimeError, "simulated interruption"):
                recreate_pendings_from_database(self.store, self.books)

        self.assertTrue((self.books / "Pendings.xlsx").is_file())
        self.assertTrue(
            (self.store.current / "pendings_recreation_journal.json").is_file()
        )
        recovered = recover_pendings_recreation(self.store, self.books)
        self.assertEqual(recovered["recovered"], "finalized")
        self.assertFalse(
            (self.store.current / "pendings_recreation_journal.json").exists()
        )

    def test_dashboard_supports_true_sr_ascending_and_descending_order(self) -> None:
        write_advanced(
            self.downloads / "Advanced Search(Service Request)20260801010101.xlsx",
            [upstream_row("12345678"), upstream_row("12345680"), upstream_row("12345679")],
        )
        run_startup(self.store)

        ascending = dashboard_payload(self.store, sort="sr", direction="asc")
        descending = dashboard_payload(self.store, sort="sr", direction="desc")
        default_order = dashboard_payload(self.store, sort="sr")

        self.assertEqual(
            [ticket["ticketId"] for ticket in ascending["tickets"]],
            ["12345678", "12345679", "12345680"],
        )
        self.assertEqual(
            [ticket["ticketId"] for ticket in descending["tickets"]],
            ["12345680", "12345679", "12345678"],
        )
        self.assertEqual(ascending["direction"], "asc")
        self.assertEqual(descending["direction"], "desc")
        self.assertEqual(default_order["direction"], "desc")
        self.assertEqual(default_order["tickets"], descending["tickets"])


class JobManagerTests(unittest.TestCase):
    def test_jobs_are_serialized_and_emit_visible_progress(self) -> None:
        broker = EventBroker()
        manager = JobManager(broker)
        order: list[str] = []

        def first(context: object) -> dict[str, bool]:
            order.append("first-start")
            context.report("query", "Checking Advanced Search")  # type: ignore[attr-defined]
            time.sleep(0.03)
            order.append("first-end")
            return {"ok": True}

        def second(context: object) -> dict[str, bool]:
            order.append("second")
            return {"ok": True}

        one = manager.submit("query", "Query", first)
        two = manager.submit("doctor", "Doctor", second)
        deadline = time.monotonic() + 3
        while time.monotonic() < deadline:
            snapshots = manager.snapshots()
            if all(job["status"] in {"succeeded", "failed"} for job in snapshots):
                break
            time.sleep(0.01)
        events = broker.wait_after(0, timeout=0)
        manager.stop()

        self.assertEqual(order, ["first-start", "first-end", "second"])
        self.assertTrue(any(event["payload"].get("message") == "Checking Advanced Search" for event in events))
        self.assertEqual(manager.get(one["id"])["status"], "succeeded")
        self.assertEqual(manager.get(two["id"])["status"], "succeeded")


class ApplicationServiceContractTests(WebFixture):
    def test_dashboard_projection_is_cached_until_the_dataset_changes(self) -> None:
        self.seed_database()
        service = ApplicationService(self.store)
        try:
            with patch(
                "zeus2.application.service.dashboard_payload",
                wraps=dashboard_payload,
            ) as serializer:
                first = service.dashboard(sort="report", search="")
                first["tickets"].clear()
                second = service.dashboard(sort="report", search="")
                self.assertEqual(serializer.call_count, 1)
                self.assertEqual(len(second["tickets"]), 1)

                service._touch_data("cache-test")
                service.dashboard(sort="report", search="")
                self.assertEqual(serializer.call_count, 2)
        finally:
            service.stop()

    def test_spare_prefill_maps_customer_and_multislot_device_data(self) -> None:
        self.seed_database()
        ticket = self.store.read_ticket("12345678")
        ticket["upstream"]["fields"].update(
            {
                "Customer Organization": "Customer Org from Advanced Search",
                "Customer Contact": "Customer Contact from Advanced Search",
                "Contact Email": "customer@example.com",
            }
        )
        ticket["local"]["spare_parts"] = [
            {
                "device": "server-a",
                "model": "Server Model",
                "faulty_sns": ["SERVER-SN-01", "MEMORY-SN-02"],
                "parts": [
                    {
                        "slot": "DIMM101\nDIMM203\nDIMM103",
                        "part": "DIMM",
                        "bom": "BOM-9000",
                        "notes": "Memory diagnostics completed.",
                        "new_sn": None,
                    }
                ],
            }
        ]
        self.store.write_ticket_bundle(self.store.current, ticket)
        service = ApplicationService(self.store)
        try:
            prefill = service.spare_request_prefill("12345678")
        finally:
            service.stop()

        profile = prefill["profile"]
        self.assertEqual(profile["customerOrganization"], "Customer Org from Advanced Search")
        self.assertEqual(profile["customerName"], "Customer Contact from Advanced Search")
        self.assertEqual(profile["contact"]["email"], "customer@example.com")
        self.assertEqual(prefill["reportDate"], "2026-07-01 10:00:00")
        self.assertEqual(
            prefill["lines"],
            [
                {
                    "bom": "BOM-9000",
                    "amount": 3,
                    "description": "DIMM",
                    "part": "DIMM",
                    "model": "Server Model",
                    "device": "server-a",
                    "slot": "DIMM101\nDIMM203\nDIMM103",
                    "slots": ["DIMM101", "DIMM203", "DIMM103"],
                    "faultySn": "SERVER-SN-01\nMEMORY-SN-02",
                    "notes": "Memory diagnostics completed.",
                    "deviceNumber": 1,
                    "partNumber": 1,
                }
            ],
        )

    def test_customer_import_uses_the_validated_form_values(self) -> None:
        self.seed_database()
        service = ApplicationService(self.store)
        profile = {
            "customerOrganization": "Consorcio Ecuatoriano de Telecomunicaciones",
            "customerName": "Angel Guerrero",
            "contact": {
                "name": "Angel Guerrero",
                "email": "angel@example.com",
                "phone": "+593980000001",
            },
        }
        try:
            with self.assertRaisesRegex(ValidationError, "email and phone"):
                service.import_customer_from_ticket(
                    "12345678",
                    {**profile, "contact": {**profile["contact"], "phone": ""}},
                )
            imported = service.import_customer_from_ticket("12345678", profile)
        finally:
            service.stop()

        organization = next(
            row
            for row in imported["data"]["organizations"]
            if row["name"] == profile["customerOrganization"]
        )
        customer = next(
            row
            for row in imported["data"]["customers"]
            if row["name"] == profile["customerName"]
        )
        self.assertEqual(customer["organizationId"], organization["id"])
        self.assertEqual(customer["email"], "angel@example.com")
        self.assertEqual(customer["phone"], "+593980000001")

    def test_legacy_email_threads_are_compacted_without_resynchronizing(self) -> None:
        self.seed_database()
        run_startup(self.store)
        ticket = self.store.read_ticket("12345678")
        raw_body = (
            "Estimada Ingri,\n\n"
            "De acuerdo, proceder con la suspensión.\n\n"
            "Regards/Saludos cordiales,\n\n"
            "De: Ingri Yoselin Herrera Vazquez <ingri@example.com>\n"
            "Enviado el: lunes, 27 de julio de 2026 18:30\n"
            "Para: TIC Karen Narvaez <karen@example.com>\n"
            "CC: Cloud Support <cloud@example.com>\n"
            "Asunto: RE: [SR 12345678] DIMM MCE error\n\n"
            "Este es el historial anterior que debe permanecer preservado."
        )
        ticket["email"]["messages"] = [
            {
                "message_key": "legacy-reply-chain",
                "timestamp": "2026-07-27T18:30:14Z",
                "direction": "received",
                "subject": "RE: [SR 12345678] DIMM MCE error",
                "sender": "ingri@example.com",
                "body": raw_body,
            }
        ]
        self.store.write_ticket_bundle(self.store.current, ticket)

        service = ApplicationService(self.store)
        try:
            message = service.ticket("12345678")["email"]["messages"][0]
        finally:
            service.stop()

        self.assertEqual(
            message["latestReplyBody"],
            "Estimada Ingri,\n\nDe acuerdo, proceder con la suspensión.\n\n"
            "Regards/Saludos cordiales,",
        )
        self.assertTrue(message["quotedHistoryHidden"])
        self.assertGreaterEqual(message["quotedHistoryLines"], 6)
        self.assertIn("historial anterior", message["body"])

    def test_web_edit_returns_a_complete_immediately_renderable_ticket(self) -> None:
        self.seed_database()
        run_startup(self.store)
        service = ApplicationService(self.store)
        try:
            before = service.ticket("12345678")
            result = service.edit_ticket(
                "12345678",
                changes={"Notes": "Saved without unmounting React"},
                expected_revision=before["revision"],
            )

            detail = result["ticket"]
            self.assertEqual(detail["localFields"]["Notes"], "Saved without unmounting React")
            self.assertIsInstance(detail["history"], list)
            self.assertIsInstance(detail["mops"], list)
            self.assertIn("email", detail)
            self.assertIn("upstreamFields", detail)
        finally:
            service.stop()

    def test_scheduled_refresh_uses_the_advanced_search_query(self) -> None:
        service = ApplicationService(self.store)
        try:
            with patch.object(
                service,
                "submit_job",
                return_value={"kind": "query", "status": "queued"},
            ) as submit:
                result = service._submit_scheduled_query()
            submit.assert_called_once_with("query", {"scheduled": True})
            self.assertEqual(result["kind"], "query")
        finally:
            service.stop()

    def test_manual_query_only_reconciles_advanced_search(self) -> None:
        service = ApplicationService(self.store)
        try:
            with (
                patch("zeus2.application.service.run_startup") as full_startup,
                patch(
                    "zeus2.application.service.reconcile_advanced_and_new_mail",
                    return_value=StartupResult(advanced_search_valid=True),
                ) as advanced_only,
            ):
                job = service.submit_job("query", {})
                deadline = time.monotonic() + 3
                snapshot = service.jobs.get(job["id"])
                while snapshot and snapshot["status"] in {"queued", "running"}:
                    if time.monotonic() >= deadline:
                        self.fail("Advanced Search-only query did not finish")
                    time.sleep(0.01)
                    snapshot = service.jobs.get(job["id"])

            self.assertEqual(snapshot["status"], "succeeded")
            full_startup.assert_not_called()
            self.assertFalse(advanced_only.call_args.kwargs["fetch_new"])
        finally:
            service.stop()

    def test_query_leaves_database_empty_when_only_pendings_exists(self) -> None:
        write_managed(
            self.books / "Pendings.xlsx",
            [pending_row("12345678", **{"Notes": "must not import"})],
        )
        service = ApplicationService(self.store)
        try:
            job = service.submit_job("query", {})
            deadline = time.monotonic() + 3
            snapshot = service.jobs.get(job["id"])
            while snapshot and snapshot["status"] in {"queued", "running"}:
                if time.monotonic() >= deadline:
                    self.fail("Advanced Search query did not finish")
                time.sleep(0.01)
                snapshot = service.jobs.get(job["id"])

            self.assertIsNotNone(snapshot)
            self.assertEqual(snapshot["status"], "succeeded")
            self.assertTrue(
                any(
                    "Advanced Search was skipped" in notice
                    for notice in snapshot["result"]["notices"]
                )
            )
            self.assertEqual(service.dashboard(sort="report", search="")["stats"]["active"], 0)
        finally:
            service.stop()

    def test_query_job_does_not_recreate_a_deleted_pendings_workbook(self) -> None:
        self.seed_database()
        run_startup(self.store)
        (self.books / "Pendings.xlsx").unlink()
        service = ApplicationService(self.store)
        try:
            job = service.submit_job("query", {})
            deadline = time.monotonic() + 3
            snapshot = service.jobs.get(job["id"])
            while snapshot and snapshot["status"] in {"queued", "running"}:
                if time.monotonic() >= deadline:
                    self.fail("Deleted-Pendings query did not finish")
                time.sleep(0.01)
                snapshot = service.jobs.get(job["id"])

            self.assertIsNotNone(snapshot)
            self.assertEqual(snapshot["status"], "succeeded")
            self.assertFalse((self.books / "Pendings.xlsx").exists())
            self.assertNotIn("pendings_recreation", snapshot["result"]["operations"])
        finally:
            service.stop()

    def test_save_response_is_database_first_when_pendings_is_missing(self) -> None:
        self.seed_database()
        startup = run_startup(self.store)
        (self.books / "Pendings.xlsx").unlink()
        service = ApplicationService(self.store)
        service.latest_startup = startup
        try:
            before = service.ticket("12345678")
            result = service.edit_ticket(
                "12345678",
                changes={"Notes": "Saved directly by the endpoint"},
                expected_revision=before["revision"],
            )

            self.assertNotIn("pendingsRecreated", result)
            self.assertFalse((self.books / "Pendings.xlsx").exists())
            self.assertEqual(result["ticket"]["localFields"]["Notes"], "Saved directly by the endpoint")
        finally:
            service.stop()

    def test_restart_command_handles_python_and_frozen_windows_builds(self) -> None:
        with patch("zeus2.main.sys.executable", "C:/Python/pythonw.exe"):
            self.assertEqual(
                _restart_command(),
                ["C:/Python/pythonw.exe", "-m", "zeus2", "serve"],
            )
        with (
            patch("zeus2.main.sys.executable", "C:/Zeus/zeus.exe"),
            patch.object(main_module.sys, "frozen", True, create=True),
        ):
            self.assertEqual(_restart_command(), ["C:/Zeus/zeus.exe", "serve"])

    def test_every_direct_email_job_fails_closed_when_outlook_is_disabled(self) -> None:
        service = ApplicationService(self.store)
        try:
            for kind in ("email-fetch", "email-sync", "email-rebuild"):
                job = service.submit_job(kind, {})
                deadline = time.monotonic() + 3
                snapshot = service.jobs.get(job["id"])
                while snapshot and snapshot["status"] in {"queued", "running"}:
                    if time.monotonic() >= deadline:
                        self.fail(f"{kind} did not finish")
                    time.sleep(0.01)
                    snapshot = service.jobs.get(job["id"])
                self.assertIsNotNone(snapshot)
                self.assertEqual(snapshot["status"], "failed")
                self.assertIn("Outlook email is disabled", snapshot["message"])
        finally:
            service.stop()


class WebServerTests(WebFixture):
    def setUp(self) -> None:
        super().setUp()
        save_user_profile(
            self.store.root,
            {
                "name": "Zeus Test User",
                "email": "zeus.user@example.com",
                "phone": "+593 99 000 0000",
                "username": "zeus-user",
            },
        )
        self.seed_database()
        run_startup(self.store)
        self.service = ApplicationService(self.store)
        static = Path(__file__).resolve().parents[1] / "zeus2" / "web" / "static"
        self.server = ZeusWebServer(
            ("127.0.0.1", 0),
            self.service,
            static_root=static,
            instance_id="test-instance",
            control_token="control-token",
        )
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()
        self.url = self.server.url

    def tearDown(self) -> None:
        self.server.shutdown()
        self.thread.join(timeout=3)
        self.service.stop()
        self.server.server_close()
        super().tearDown()

    def read_json(self, path: str) -> tuple[int, dict[str, object], dict[str, str]]:
        try:
            response = urllib.request.urlopen(self.url + path, timeout=3)
        except urllib.error.HTTPError as exc:
            response = exc
        with response:
            return response.status, json.loads(response.read()), dict(response.headers.items())

    def test_browser_refresh_reads_state_without_querying_sources(self) -> None:
        audit_before = self.store.audit_file.read_text(encoding="utf-8")
        status, dashboard, _ = self.read_json("/api/dashboard?sort=report&search=")
        _, index = self.read_json("/api/bootstrap")[:2]
        audit_after = self.store.audit_file.read_text(encoding="utf-8")

        self.assertEqual(status, 200)
        self.assertEqual(dashboard["stats"]["active"], 1)  # type: ignore[index]
        self.assertEqual(index["instanceId"], "test-instance")
        self.assertEqual(index["appearance"], {"fontScale": "standard"})
        self.assertEqual(audit_before, audit_after)

    def test_first_run_requires_profile_then_unlocks_the_workbench(self) -> None:
        user_profile_path(self.store.root).unlink()
        status, bootstrap, _ = self.read_json("/api/bootstrap")
        self.assertEqual(status, 200)
        self.assertTrue(bootstrap["onboarding"]["required"])

        status, blocked, _ = self.read_json("/api/dashboard?sort=report&search=")
        self.assertEqual(status, 428)
        self.assertEqual(blocked["error"]["code"], "setup_required")

        request = urllib.request.Request(
            self.url + "/api/profile",
            data=json.dumps(
                {
                    "profile": {
                        "name": "First Run User",
                        "email": "first.run@example.com",
                        "phone": "+593 98 765 4321",
                        "username": "first-run",
                        "photoDataUrl": None,
                    }
                }
            ).encode("utf-8"),
            method="PATCH",
            headers={
                "Content-Type": "application/json",
                "X-Zeus-CSRF": str(bootstrap["csrfToken"]),
            },
        )
        with urllib.request.urlopen(request, timeout=3) as response:
            saved = json.loads(response.read())
        self.assertTrue(saved["complete"])
        self.assertEqual(saved["profile"]["name"], "First Run User")
        self.assertNotIn("password", saved["profile"])

        status, dashboard, _ = self.read_json("/api/dashboard?sort=report&search=")
        self.assertEqual(status, 200)
        self.assertEqual(dashboard["stats"]["active"], 1)

    def test_spare_requests_workspace_is_available_without_mutating_sources(self) -> None:
        audit_before = self.store.audit_file.read_text(encoding="utf-8")
        status, payload, _ = self.read_json(
            "/api/dashboard?workspace=spare-requests&view=eligible&sort=tt&direction=desc&search="
        )
        audit_after = self.store.audit_file.read_text(encoding="utf-8")

        self.assertEqual(status, 200)
        self.assertEqual(payload["workspace"], "spare-requests")
        self.assertEqual(payload["view"], "eligible")
        self.assertIn("eligibleParts", payload)
        self.assertEqual(payload["columns"][0]["label"], "TT")
        self.assertNotIn("tickets", payload)
        self.assertEqual(audit_before, audit_after)

    def test_manual_spare_registration_route_needs_no_export_configuration(self) -> None:
        _, bootstrap, _ = self.read_json("/api/bootstrap")
        payload = {
            "source": "manual",
            "ticketId": "12345678",
            "reportDate": "2026-07-01",
            "profile": {
                "customerName": "Customer Network Team",
                "siteCode": "GYE",
                "siteAddress": "Guayaquil operations center",
                "cloud": "Cloud",
                "requester": {
                    "name": "Zeus Test User",
                    "email": "zeus.user@example.com",
                    "phone": "+593990000000",
                },
                "contact": {
                    "name": "Customer Contact",
                    "email": "customer@example.com",
                    "phone": "+593980000000",
                },
            },
            "lines": [
                {
                    "bom": "BOM-1",
                    "amount": 1,
                    "description": "Disk",
                    "part": "Disk",
                    "reportDate": "2026-07-01",
                }
            ],
        }
        request = urllib.request.Request(
            self.url + "/api/spare-requests/register-manual",
            data=json.dumps(payload).encode("utf-8"),
            method="POST",
            headers={
                "Content-Type": "application/json",
                "X-Zeus-CSRF": str(bootstrap["csrfToken"]),
            },
        )
        with urllib.request.urlopen(request, timeout=3) as response:
            registered = json.loads(response.read())
            self.assertEqual(response.status, 201)

        self.assertEqual(registered["request"]["creationMethod"], "manual_confirmation")
        self.assertIsNone(registered["request"]["export"]["request_filename"])
        self.assertEqual(
            registered["request"]["history"][0]["action"],
            "request-registered-manually",
        )

    def test_ticket_patch_returns_the_complete_detail_contract(self) -> None:
        _, bootstrap, _ = self.read_json("/api/bootstrap")
        _, detail, _ = self.read_json("/api/tickets/12345678")
        body = json.dumps(
            {
                "revision": detail["revision"],
                "changes": {"Notes": "Complete API response"},
            }
        ).encode("utf-8")
        request = urllib.request.Request(
            self.url + "/api/tickets/12345678/local",
            data=body,
            method="PATCH",
            headers={
                "Content-Type": "application/json",
                "If-Match": str(detail["revision"]),
                "X-Zeus-CSRF": str(bootstrap["csrfToken"]),
            },
        )
        with urllib.request.urlopen(request, timeout=3) as response:
            payload = json.loads(response.read())

        edited = payload["ticket"]
        self.assertEqual(edited["localFields"]["Notes"], "Complete API response")
        self.assertIsInstance(edited["history"], list)
        self.assertIsInstance(edited["mops"], list)

    def test_bulk_draft_patch_returns_complete_saved_tickets(self) -> None:
        _, bootstrap, _ = self.read_json("/api/bootstrap")
        _, detail, _ = self.read_json("/api/tickets/12345678")
        body = json.dumps(
            {
                "edits": [
                    {
                        "ticketId": "12345678",
                        "revision": detail["revision"],
                        "changes": {"Notes": "Saved from protected drafts manager"},
                    }
                ]
            }
        ).encode("utf-8")
        request = urllib.request.Request(
            self.url + "/api/tickets/bulk-local",
            data=body,
            method="PATCH",
            headers={
                "Content-Type": "application/json",
                "X-Zeus-CSRF": str(bootstrap["csrfToken"]),
            },
        )
        with urllib.request.urlopen(request, timeout=3) as response:
            payload = json.loads(response.read())

        self.assertEqual(payload["ticketIds"], ["12345678"])
        saved = payload["tickets"]["12345678"]
        self.assertEqual(saved["localFields"]["Notes"], "Saved from protected drafts manager")
        self.assertIsInstance(saved["history"], list)

    def test_bundled_react_entrypoint_and_hashed_assets_are_served_locally(self) -> None:
        with urllib.request.urlopen(self.url + "/", timeout=3) as response:
            html = response.read().decode("utf-8")
            headers = dict(response.headers.items())
        self.assertIn('<div id="root"></div>', html)
        self.assertEqual(headers["Cache-Control"], "no-store")
        asset_match = re.search(r'src="(/[^"]+\.js)"', html)
        self.assertIsNotNone(asset_match)
        with urllib.request.urlopen(self.url + asset_match.group(1), timeout=3) as response:
            self.assertGreater(len(response.read()), 100_000)
            self.assertIn("immutable", response.headers["Cache-Control"])

    def test_restore_preview_reports_changes_without_mutating_the_ticket(self) -> None:
        backup_directory = self.books / "Zeus Backups"
        backup_directory.mkdir()
        backup = backup_directory / "Pendings_preview.xlsx"
        write_managed(
            backup,
            [
                pending_row(
                    "12345678",
                    summary="Pendings alone builds Zeus",
                    **{"Site": "GYE", "Notes": "Backup note"},
                )
            ],
        )
        before = self.store.read_ticket("12345678")["local"]["fields"]["Notes"]
        status, payload, _ = self.read_json(
            "/api/backups/pendings/preview?name=Pendings_preview.xlsx"
        )
        self.assertEqual(status, 200)
        self.assertEqual(payload["preview"]["changed_ids"], ["12345678"])
        self.assertEqual(
            self.store.read_ticket("12345678")["local"]["fields"]["Notes"],
            before,
        )

    def test_csrf_host_and_security_headers_are_enforced(self) -> None:
        status, bootstrap, headers = self.read_json("/api/bootstrap")
        self.assertEqual(status, 200)
        self.assertIn("default-src 'self'", headers["Content-Security-Policy"])
        self.assertIn("script-src 'self';", headers["Content-Security-Policy"])
        self.assertIn(
            "style-src 'self' 'unsafe-inline';",
            headers["Content-Security-Policy"],
        )

        request = urllib.request.Request(
            self.url + "/api/jobs/query",
            data=b"{}",
            method="POST",
            headers={"Content-Type": "application/json"},
        )
        with self.assertRaises(urllib.error.HTTPError) as missing:
            urllib.request.urlopen(request, timeout=3)
        self.assertEqual(missing.exception.code, 403)

        request.add_header("X-Zeus-CSRF", str(bootstrap["csrfToken"]))
        with urllib.request.urlopen(request, timeout=3) as response:
            self.assertEqual(response.status, 202)

        connection = http.client.HTTPConnection("127.0.0.1", self.server.server_port, timeout=3)
        connection.request("GET", "/api/health", headers={"Host": "malicious.example"})
        rejected = connection.getresponse()
        self.assertEqual(rejected.status, 422)
        rejected.read()
        connection.close()

    def test_control_token_stops_only_the_registered_server(self) -> None:
        request = urllib.request.Request(
            self.url + "/api/system/shutdown",
            data=b"{}",
            method="POST",
            headers={"Content-Type": "application/json", "X-Zeus-Control": "wrong"},
        )
        with self.assertRaises(urllib.error.HTTPError) as rejected:
            urllib.request.urlopen(request, timeout=3)
        self.assertEqual(rejected.exception.code, 403)


class InstanceRegistryTests(WebFixture):
    def test_current_process_has_a_stable_birth_marker(self) -> None:
        import os

        marker = process_creation_marker(os.getpid())
        if marker is None and sys.platform != "win32":
            self.skipTest("This container does not expose process birth metadata")
        self.assertIsNotNone(marker)
        self.assertEqual(process_creation_marker(os.getpid()), marker)

    def test_stale_records_are_removed_without_touching_unrelated_processes(self) -> None:
        registry = InstanceRegistry(self.home)
        registry.ensure()
        stale = registry.instances / "stale.json"
        stale.write_text(
            json.dumps(
                {
                    "instanceId": "stale",
                    "pid": 99999999,
                    "processCreationMarker": "not-real",
                    "port": 65534,
                    "url": "http://127.0.0.1:65534",
                    "controlToken": "x",
                }
            ),
            encoding="utf-8",
        )
        self.assertEqual(registry.live_instances(), [])
        self.assertFalse(stale.exists())


if __name__ == "__main__":
    unittest.main()
